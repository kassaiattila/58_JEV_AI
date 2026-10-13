"""136 (backlog F-revolut-csv, E4): a bank's tabular statement export read by code, without AI.

The owner's decisions: a general tabular statement reader, not code for one bank (DECISIONS 132; the company card's
export comes later); processed as a work package with one statement per account, and only the accounts a person
chooses go into the store (DECISIONS 132, 136). The 2026 Revolut consolidated statement, an Excel workbook, holds the
summaries of every account first and then one transaction table per account; nothing in it needs a model.

A profile (`configs/statement_tables.json`) describes a kind of export as data:

- `columns`: the header aliases of each role (date, description, category, amount, balance); a row holding the aliases
  of every `required` role is a table header. A repeated header (an amount in the account's currency, then converted)
  takes its first column;
- `account_title`: a pattern for the row naming an account and its currency code in brackets; the summary block and
  the transaction table of an account follow such a row. Two accounts may share a title, so the n-th summary of a
  title belongs to its n-th table;
- `summary`: the labels of the account number and of the opening and closing balances, each followed by its value;
- `total_row`: the first cell of the row that totals a table and ends it;
- `period_from_filename`: the statement period, when only the file name states it; otherwise the first and the last
  booking date.

Reading takes the rows of cell texts of a workbook's first sheet (read only, values) or of a CSV file, within the
limits of the configuration. Three repairs an export may need, each recorded in `Reading.repairs`:

- `csv_in_first_columns`: a CSV export opened as a workbook keeps each record in its first one or two cells, so such
  rows are joined and split as CSV again;
- `mojibake`: UTF-8 text decoded as cp1250 is turned back, only when every text of the file survives the round trip
  (a byte cp1250 leaves undefined comes through as its own code point and goes back as that byte);
- `html_entities`: an entity left in a text (an ampersand written as `&amp;`) is unescaped.

The statement of an account has the fields of the bank statement type packs (`configs/types/statement_cib.json`):
statement type, account number, currency, period, opening and closing balance, debit and credit totals and the
transactions (booking date, direction, a positive amount, the running balance, the counterparty's name as printed in
the description column, and the category as the description). Without a summary, the opening balance is the first
line's balance less its amount and the closing balance the last line's. Code checks the running balance line by line,
the total row against the lines and the opening balance plus the lines against the closing balance (`Account.checks`);
a line whose date or amount cannot be read is left out and reported (`Account.problems`).

137 (DECISIONS 137, "XML + PDF by code"): a **ledger** profile (`layout: ledger`) reads an export that lists every line
with its own account and statement number, such as the Erste Excel 2003 XML export (SpreadsheetML, read with
defusedxml): one statement per account and statement number, the period from the file name's month. Such an export has
no balance, and a card line names no merchant, so the profile's **companion** is the same month's PDF statement, found
by its name next to the export or in a `pdf` subfolder. Its text layer (the isolated PDF reader) is read with the
profile's patterns: the period, the account, the opening and closing balances and every line with its running balance,
the merchant and a foreign purchase's original amount and currency. The export's lines are matched to the PDF's by
booking date and amount, in order; the export stays the list of lines, the PDF adds the balances and the details, and a
line on one side only is a finding (`Account.companion_problems`). Without a PDF the balances cannot be checked
(`Account.balances` is `none`). A folder is read as a whole: every export in it, each with its own source and
companion, and the files no profile knows are listed (`Reading.skipped`).

138 (Q-table-reader-isolation, DECISIONS 138): the third-party parsing of an export (`parse_rows`) runs in the isolated
PDF reader's helper process, within its read time limit and memory limit, as the local service reads a file a person
picks; our own reading of the rows stays in the calling process.

141 (DECISIONS 141, the sports club's statements read by code first): a **text** profile (`layout: text`) reads a statement printed as text:
a PDF's text layer (the isolated PDF reader) or the bank's fixed-width text export (`.txt`, in one of the profile's
`encodings`), such as the CIB statement. A profile knows its statement by `signature` (every pattern found in the text);
`header` gives the statement number, period, currency, account, IBAN, opening and closing balances and the printed
totals; the lines stand between `lines_start` and `lines_end`, without the page breaks (`skip_blocks`). A line is its
first row (booking date, kind, signed amount, running balance), then its value date row, whose text is the
counterparty's account (the next unlabelled row is then the counterparty's name), a foreign transfer's original amount
(`original`) or the bank's own description of a fee; the labelled rows (`labels`) give the counterparty's name and
account and the memo. The statement checks itself: the running balance line by line, the printed totals and the closing
balance. A PDF no text profile knows is left alone when a folder is read: it may be another export's companion or a
statement to extract (`jav/statement_package.py`).

141 (DECISIONS 141, the Erste account's source): a **camt053** profile (`layout: camt053`) reads a bank's ISO 20022
camt.053 statement (an XML file in the `urn:iso:std:iso:20022:tech:xsd:camt.053.` namespace, parsed with defusedxml in
the isolated reader's helper): per statement its account (IBAN), currency, number, period, the booked opening and
closing balances and the printed totals, and per entry its booking and value dates, signed amount, running balance
(`POSTBALANCE`), counterparty (the creditor of a debit, the debtor of a credit) with its account, and memo. A bank,
known by its BIC (`banks`), adds the texts of its own transaction codes (`kinds`, as its Excel export prints them) and
the patterns of its card details in the entry's additional information (`details`: the merchant, a foreign purchase's
original amount and currency); another bank's statement keeps its codes. The statement checks itself like an export.
"""

from __future__ import annotations

import calendar
import csv
import hashlib
import html
import io
import json
import re
import zipfile
from dataclasses import dataclass, field, replace
from decimal import Decimal
from pathlib import Path
from typing import Any

from jav import card_numbers, cfg, isolated_pdf
from jav.dates import read_date
from jav.numbers import read_number

CONFIG = "statement_tables"
SUFFIXES = (".xlsx", ".csv", ".xml")
TEXT_SUFFIXES = (".pdf", ".txt")  # 141: a statement printed as text
CENT = Decimal("0.01")


class StatementTableError(ValueError):
    """The file cannot be read as a statement table: over the limits, unreadable, or no profile knows it."""


@dataclass(frozen=True)
class Account:
    key: str  # the account number without spaces and the currency (one number may hold several currencies), or the
    #           printed title with its occurrence when no number is printed
    title: str  # the account's name without its currency
    currency: str
    occurrence: int  # the n-th account of the same printed title
    statement: dict[str, Any]
    checks: dict[str, Any]
    problems: list[str] = field(default_factory=list)
    source: dict[str, Any] | None = None  # 137: the export file (and its companion) this account was read from
    balances: str = "export"  # 137: where the balances come from: export (141: or the text statement itself) |
    #                           companion | none (cannot be checked)
    companion_problems: list[str] = field(default_factory=list)  # 137: lines on one side only, another account


@dataclass(frozen=True)
class Reading:
    profile: str
    institution: str
    file_name: str
    sha256: str
    repairs: list[str]
    accounts: list[Account]
    skipped: list[dict[str, str]] = field(default_factory=list)  # 137: the files of a folder no profile reads


def _conf() -> dict[str, Any]:
    return cfg.load(CONFIG)


def config_hash() -> str:
    return cfg.config_hash(CONFIG)


# --- rows of cell texts ------------------------------------------------------------------------------------------------


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _workbook_rows(data: bytes, limits: dict[str, Any]) -> list[list[str]]:
    import openpyxl

    from jav.readers.contracts import ReadLimits
    from jav.readers.limits import DEFAULT_LIMITS, ReadFailure, inspect_package

    archive = limits["archive"]
    try:  # the shared ZIP inspection: no unsafe member, no expansion beyond the bounds, external parts inert
        inspect_package(data, ReadLimits(**{**DEFAULT_LIMITS.model_dump(), **archive, "input_bytes": limits["max_bytes"],
                                            "visited_cells": limits["max_cells"]}))
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (ReadFailure, OSError, KeyError, ValueError, zipfile.BadZipFile) as exc:  # 138: a broken archive too
        raise StatementTableError(f"the workbook cannot be read: {exc}") from exc
    try:
        rows, cells = [], 0
        for row in wb.worksheets[0].iter_rows(values_only=True):
            texts = [_text(v) for v in row]
            while texts and not texts[-1]:
                texts.pop()
            cells += len(texts)
            if len(rows) >= limits["max_rows"] or cells > limits["max_cells"]:
                raise StatementTableError("the workbook is over the row limit or the cell limit")
            rows.append(texts)
        return rows
    finally:
        wb.close()


_SS = "{urn:schemas-microsoft-com:office:spreadsheet}"


def _spreadsheetml_rows(data: bytes, limits: dict[str, Any]) -> list[list[str]]:
    """137: the first sheet of an Excel 2003 XML workbook (SpreadsheetML). defusedxml refuses a document type
    declaration, entities and external references; a cell's `ss:Index` skips the empty cells before it; a date-time
    keeps its date."""
    from defusedxml import ElementTree

    try:
        root = ElementTree.fromstring(data)
    except Exception as exc:  # defusedxml raises its own and ElementTree's errors (as in jav/fx.py)
        raise StatementTableError(f"the XML workbook cannot be read: {type(exc).__name__}") from exc
    sheet = root.find(_SS + "Worksheet") if root.tag == _SS + "Workbook" else None
    table = sheet.find(_SS + "Table") if sheet is not None else None
    if table is None:
        raise StatementTableError("not an Excel 2003 XML workbook with a table")
    rows, cells = [], 0
    for row in table.iter(_SS + "Row"):
        texts: list[str] = []
        for cell in row.findall(_SS + "Cell"):
            if (index := cell.get(_SS + "Index")) and index.isdigit():
                texts += [""] * max(0, int(index) - 1 - len(texts))
            data_el = cell.find(_SS + "Data")
            value = _text(data_el.text) if data_el is not None else ""
            if data_el is not None and data_el.get(_SS + "Type") == "DateTime":
                value = value[:10]
            texts.append(value)
        while texts and not texts[-1]:
            texts.pop()
        cells += len(texts)
        if len(rows) >= limits["max_rows"] or cells > limits["max_cells"]:
            raise StatementTableError("the workbook is over the row limit or the cell limit")
        rows.append(texts)
    return rows


def _csv_rows(data: bytes, limits: dict[str, Any]) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1250")
        except UnicodeDecodeError as exc:
            raise StatementTableError("the CSV file is neither UTF-8 nor cp1250 text") from exc
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") > first.count(",") else ","
    rows: list[list[str]] = []
    cells = 0
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):  # 153 (review C3): the cell limit while reading
        cells += len(row)
        if len(rows) >= limits["max_rows"] or cells > limits["max_cells"]:
            raise StatementTableError("the file is over the row limit or the cell limit")
        rows.append([_text(v) for v in row])
    return rows


def _joined_csv(rows: list[list[str]], conf: dict[str, Any]) -> list[list[str]] | None:
    """The records of a CSV export opened as a workbook: most rows fill at most `max_filled_cells` cells and hold a
    comma; each row's filled cells are joined with a comma and split as CSV again."""
    filled = [[c for c in r if c] for r in rows]
    content = [f for f in filled if f]
    narrow = [f for f in content if len(f) <= conf["max_filled_cells"]]
    if not content or len(narrow) < conf["min_share"] * len(content) or not any("," in f[0] for f in content):
        return None
    return [next(csv.reader(io.StringIO(",".join(f))), []) if f else [] for f in filled]


def _to_bytes(text: str, encoding: str) -> bytes:
    out = bytearray()
    for ch in text:
        try:
            out += ch.encode(encoding)
        except UnicodeEncodeError:
            if ord(ch) > 0xFF:
                raise
            out.append(ord(ch))  # a byte the encoding leaves undefined came through as its own code point
    return bytes(out)


def _unmangled(rows: list[list[str]], conf: dict[str, Any]) -> list[list[str]] | None:
    """UTF-8 read as `read_as`, turned back when a marker shows it and every text survives the round trip."""
    if not any(m in c for r in rows for c in r for m in conf["markers"]):
        return None
    try:
        return [[_to_bytes(c, conf["read_as"]).decode("utf-8") if not c.isascii() else c for c in r] for r in rows]
    except (UnicodeEncodeError, UnicodeDecodeError):
        return None


def parse_rows(suffix: str, data: bytes, limits: dict[str, Any]) -> list[list[str]]:
    """The rows of cell texts of an export's bytes: the third-party parsing (openpyxl, defusedxml, the CSV reader). 138
    (Q-table-reader-isolation): it runs in the isolated reader's helper process (`jav/isolated_pdf.py`), within its time
    and memory limits, every setting a parameter."""
    if suffix == ".csv":
        return _csv_rows(data, limits)
    if suffix == ".xml":
        return _spreadsheetml_rows(data, limits)
    return _workbook_rows(data, limits)


def _parsed(suffix: str, data: bytes, limits: dict[str, Any]) -> list[list[str]]:
    """`parse_rows` in the isolated helper; its error keeps its message, a limit is a named refusal."""
    try:
        return isolated_pdf.run(parse_rows, kind="read", suffix=suffix, data=data, limits=limits)
    except isolated_pdf.PdfReaderError as exc:
        name, _, message = str(exc).partition(": ")
        raise StatementTableError(message if name == StatementTableError.__name__ else str(exc)) from exc
    except isolated_pdf.PdfReaderLimit as exc:
        raise StatementTableError(f"the file could not be read within the reader's limits ({exc.reason}): {exc}") from exc


def _rows(path: Path, data: bytes, conf: dict[str, Any]) -> tuple[list[list[str]], list[str]]:
    limits, repairs = conf["limits"], []
    if len(data) > limits["max_bytes"]:
        raise StatementTableError("the file is over the size limit")
    rows = _parsed(path.suffix.lower(), data, limits)
    if (fixed := _unmangled(rows, conf["repairs"]["mojibake"])) is not None:
        rows, repairs = fixed, [*repairs, "mojibake"]
    if path.suffix.lower() != ".csv" and (joined := _joined_csv(rows, conf["repairs"]["csv_in_first_columns"])) is not None:
        rows, repairs = joined, [*repairs, "csv_in_first_columns"]
    if any("&" in c and html.unescape(c) != c for r in rows for c in r):
        rows, repairs = [[html.unescape(c) for c in r] for r in rows], [*repairs, "html_entities"]
    return rows, repairs


# --- sections ------------------------------------------------------------------------------------------------------------


def _norm(text: str) -> str:
    return " ".join(text.casefold().split())


def _header(row: list[str], profile: dict[str, Any]) -> dict[str, int] | None:
    """The column of each role (its first occurrence) when the row holds every required role's alias."""
    cells = [_norm(c) for c in row]
    found: dict[str, int] = {}
    for role, aliases in profile["columns"].items():
        wanted = {_norm(a) for a in aliases}
        index = next((i for i, c in enumerate(cells) if c in wanted), None)
        if index is not None:
            found[role] = index
    return found if all(r in found for r in profile["required"]) else None


def _value_after(row: list[str], aliases: list[str]) -> str | None:
    wanted = {_norm(a) for a in aliases}
    for i, c in enumerate(row):
        if _norm(c) in wanted:
            return next((v for v in row[i + 1:] if v), None)
    return None


@dataclass
class _Block:
    title: str  # as printed, with the currency
    name: str
    currency: str
    summary: dict[str, str] = field(default_factory=dict)
    columns: dict[str, int] | None = None
    lines: list[list[str]] = field(default_factory=list)
    total: list[str] | None = None


def _blocks(rows: list[list[str]], profile: dict[str, Any]) -> list[_Block]:
    title_re = re.compile(profile["account_title"])
    totals = {_norm(t) for t in profile["total_row"]}
    blocks: list[_Block] = []
    in_table = False
    for row in rows:
        filled = [c for c in row if c]
        if not filled:
            continue
        if len(filled) == 1:
            in_table = False
            if m := title_re.match(filled[0]):
                blocks.append(_Block(title=filled[0], name=m["name"].strip(), currency=m["currency"]))
            continue
        if not blocks:
            continue
        block = blocks[-1]
        if block.columns is None and (columns := _header(row, profile)) is not None:
            block.columns, in_table = columns, True
            continue
        if in_table:
            if _norm(row[0]) in totals:
                block.total, in_table = row, False
            else:
                block.lines.append(row)
            continue
        for key, aliases in profile["summary"].items():
            if key not in block.summary and (value := _value_after(row, aliases)) is not None:
                block.summary[key] = value
    return blocks


# --- the statement of an account ---------------------------------------------------------------------------------------


def _cents(value: Decimal) -> Decimal | None:
    """153 (review A6): an amount in cents, only when no non-zero digit stands beyond them ("10.000" is 10.00); "10.005"
    is None, so the line is unreadable and said so, instead of rounded without a word."""
    return value.quantize(CENT) if value == value.quantize(CENT) else None


def _money(text: str | None) -> Decimal | None:
    if not text:
        return None
    read = read_number(text)
    return None if read.value is None or read.ambiguous else _cents(Decimal(read.value))


def _cell(row: list[str], columns: dict[str, int], role: str) -> str:
    i = columns.get(role)
    return row[i] if i is not None and i < len(row) else ""


def _plain(value: Decimal | None) -> str | None:
    return None if value is None else f"{value:.2f}"


def _period(file_name: str, profile: dict[str, Any], dates: list[str]) -> tuple[str | None, str | None]:
    pattern = profile.get("period_from_filename")
    if pattern and (m := re.search(pattern, file_name)):
        return m["start"], m["end"]
    return (min(dates), max(dates)) if dates else (None, None)


def _account(block: _Block, summary: dict[str, str], occurrence: int, file_name: str, profile: dict[str, Any]) -> Account:
    columns = block.columns or {}
    lines, amounts, balances, problems = [], [], [], []
    for n, row in enumerate(block.lines, start=1):
        booked = read_date(_cell(row, columns, "date"))
        amount, balance = _money(_cell(row, columns, "amount")), _money(_cell(row, columns, "balance"))
        if booked.value is None or booked.ambiguous:
            problems.append(f"line {n}: unreadable date")
            continue
        if amount is None:
            problems.append(f"line {n}: unreadable amount")
            continue
        amounts.append(amount)
        balances.append(balance)
        lines.append({"booking_date": booked.value.isoformat(), "value_date": None,
                      "direction": "debit" if amount < 0 else "credit", "amount": _plain(abs(amount)),
                      "running_balance": _plain(balance), "description": _cell(row, columns, "category") or None,
                      "counterparty_name": _cell(row, columns, "description") or None, "counterparty_account": None, "memo": None})
        if (fee := _money(_cell(row, columns, "fee"))) is not None and fee != 0:  # 144: a fee column (Revolut)
            _with(lines[-1], fee=f"{_plain(abs(fee))} {block.currency}")
    iban = (summary.get("account_iban") or "").replace(" ", "") or None
    opening, closing = _money(summary.get("opening_balance")), _money(summary.get("closing_balance"))
    if opening is None and balances and balances[0] is not None:
        opening = balances[0] - amounts[0]
    if closing is None and balances and balances[-1] is not None:
        closing = balances[-1]
    start, end = _period(file_name, profile, [t["booking_date"] for t in lines])
    statement = {"statement_type": profile["statement_type"], "account_no": None, "account_iban": iban,
                 "period_start": start, "period_end": end, "currency": block.currency,
                 "opening_balance": _plain(opening), "closing_balance": _plain(closing),
                 "total_debit": _plain(sum((-a for a in amounts if a < 0), Decimal(0))),
                 "total_credit": _plain(sum((a for a in amounts if a > 0), Decimal(0))), "transactions": lines}
    running, previous = 0, opening
    for amount, balance in zip(amounts, balances):
        running += previous is not None and balance is not None and previous + amount == balance
        previous = balance
    total = _money(_cell(block.total, columns, "amount")) if block.total else None
    checks = {"running_balance": {"ok": running, "lines": len(lines)},
              "total_row": None if block.total is None else total == sum(amounts, Decimal(0)),
              "closing_balance": None if opening is None or closing is None else opening + sum(amounts, Decimal(0)) == closing}
    key = f"{iban}:{block.currency}" if iban else f"{block.title}#{occurrence}"
    return Account(key=key, title=block.name, currency=block.currency, occurrence=occurrence, statement=statement,
                   checks=checks, problems=problems)


def _accounts(blocks: list[_Block], file_name: str, profile: dict[str, Any]) -> list[Account]:
    summaries: dict[str, list[dict[str, str]]] = {}
    for b in blocks:
        if b.columns is None:
            summaries.setdefault(b.title, []).append(b.summary)
    seen: dict[str, int] = {}
    out = []
    for b in blocks:
        if b.columns is None:
            continue
        n = seen[b.title] = seen.get(b.title, 0) + 1
        own = summaries.get(b.title, [])
        summary = b.summary or (own[n - 1] if n <= len(own) else {})
        out.append(_account(b, summary, n, file_name, profile))
    return out


# --- 137: a ledger export (a line per row, with its account and statement number) and its PDF companion ----------------


def _ledger_amount(text: str, profile: dict[str, Any]) -> Decimal | None:
    """A machine number ("-7291.00") when the profile says so, otherwise the shared number reader."""
    if profile.get("amounts") == "machine":
        return _cents(Decimal(text)) if re.fullmatch(r"-?\d+(?:\.\d+)?", text) else None
    return _money(text)


def _ledgers(rows: list[list[str]], profile: dict[str, Any]
             ) -> tuple[dict[str, int] | None, dict[tuple[str, str], list[list[str]]]]:
    """The header's columns and the rows after it, grouped by account and statement number, in order."""
    columns: dict[str, int] | None = None
    groups: dict[tuple[str, str], list[list[str]]] = {}
    for row in rows:
        if not any(row):
            continue
        if columns is None:
            columns = _header(row, profile)
            continue
        groups.setdefault((_cell(row, columns, "account"), _cell(row, columns, "statement")), []).append(row)
    return columns, groups


# 153 (review C1): what a statement's patterns read stays short. A crafted file with a 20,000-digit amount or a line of
# thousands of spaces kept the service busy for minutes (the patterns backtrack, and the reading after the isolated
# helper's parse runs in the service): a field or a line over LINE_MAX characters is cut or left out, and a run of three
# or more spaces becomes two (the text profiles separate their columns by two or more).
LINE_MAX = 2000
_SPACE_RUN = re.compile(r"[ \t]{3,}")


def _bounded(lines: list[str]) -> tuple[list[str], list[str]]:
    """The lines with their space runs shortened, then any line still over LINE_MAX characters left out, and the
    repairs that say how many were left out."""
    out, cut = [], 0
    for ln in lines:
        ln = _SPACE_RUN.sub("  ", ln)
        if len(ln) > LINE_MAX:
            cut += 1
            continue
        out.append(ln)
    return out, ([f"{cut} line(s) over {LINE_MAX} characters left out"] if cut else [])


def _within_size(path: Path, limits: dict[str, Any]) -> None:
    """153 (review C3): the size limit is checked before a file is read (a 1 GB file raised the service's memory to
    1.27 GB before it was refused)."""
    if path.stat().st_size > limits["max_bytes"]:
        raise StatementTableError("the file is over the size limit")


def _pdf_lines(path: Path) -> list[str]:
    """The PDF's text layer as lines (the isolated PDF reader, jav/pdf.py); the tests replace it."""
    from jav import pdf

    return pdf.read_pdf(path).lines


def _hu_amount(text: str) -> Decimal | None:
    read = read_number(text)
    return None if read.value is None or read.ambiguous else _cents(Decimal(read.value))


def _dot_amount(text: str) -> Decimal | None:
    """An amount with a decimal point and optional thousands commas ("1,234.56", "20.00")."""
    plain = text.replace(",", "")
    return _cents(Decimal(plain)) if re.fullmatch(r"\d+(?:\.\d+)?", plain) else None


def _without_page_breaks(lines: list[str], blocks: list[list[str]]) -> list[str]:
    """The lines without the repeated page footer and header: from a block's first pattern to its last, inclusive."""
    out: list[str] = []
    skipping: str | None = None
    for line in lines:
        if skipping is None:
            skipping = next((end for start, end in blocks if re.search(start, line)), None)
            if skipping is None:
                out.append(line)
        elif re.search(skipping, line):
            skipping = None
    return out


def parse_companion(lines: list[str], conf: dict[str, Any]) -> dict[str, Any]:
    """The PDF statement's period, account, currency, balances and lines (each with its details), from its text lines
    and the profile's patterns (`companion`). Pure: no file is read."""
    lines = _without_page_breaks(lines, conf.get("skip_blocks", []))
    text = "\n".join(lines)

    def first(key: str) -> re.Match[str] | None:
        return re.search(conf[key], text, re.MULTILINE) if conf.get(key) else None

    period, opening, closing = first("period"), first("opening"), first("closing")
    line_re, start_re, end_re = re.compile(conf["line"]), re.compile(conf["lines_start"]), re.compile(conf["lines_end"])
    details = {k: re.compile(v) for k, v in conf.get("details", {}).items()}
    started = False
    out: list[dict[str, Any]] = []
    problems: list[str] = []
    for line in lines:
        if not started:
            started = bool(start_re.search(line))
            continue
        if end_re.search(line):
            break
        if m := line_re.match(line):
            booked, amount, balance = read_date(m["booked"]), _hu_amount(m["amount"]), _hu_amount(m["balance"])
            if booked.value is None or amount is None or balance is None:
                problems.append(f"PDF line {len(out) + 1}: unreadable")
            out.append({"booked": booked.value.isoformat() if booked.value else None, "kind": m["kind"].strip(),
                        "amount": amount, "balance": balance, "details": {}})
            continue
        if out:
            for key, rx in details.items():
                if key not in out[-1]["details"] and (d := rx.search(line)):
                    out[-1]["details"][key] = {k: v.strip() for k, v in d.groupdict().items() if v}
    start = read_date(period["start"]).value if period else None
    end = read_date(period["end"]).value if period else None
    account, currency = first("account"), first("currency")
    return {"period": (start.isoformat() if start else None, end.isoformat() if end else None),
            "account": account["value"] if account else None, "currency": currency["value"] if currency else None,
            "opening": _hu_amount(opening["value"]) if opening else None,
            "closing": _hu_amount(closing["value"]) if closing else None, "lines": out, "problems": problems}


def _companion_path(path: Path, conf: dict[str, Any] | None) -> Path | None:
    if not conf or not (m := re.match(conf["name_from"], path.name)):
        return None
    name = conf["name"].format(**m.groupdict())
    return next((p for folder in conf.get("folders", ["."]) if (p := path.parent / folder / name).is_file()), None)


def _month(file_name: str, pattern: str | None) -> tuple[str | None, str | None]:
    """The calendar month the file name names (`year`, `month`), as a period."""
    if not pattern or not (m := re.search(pattern, file_name)):
        return None, None
    year, month = int(m["year"]), int(m["month"])
    return f"{year:04d}-{month:02d}-01", f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"


def _ledger_lines(rows: list[list[str]], columns: dict[str, int], profile: dict[str, Any]
                  ) -> tuple[list[dict[str, Any]], list[Decimal], list[str]]:
    placeholders = {_norm(x) for x in profile.get("placeholder_counterparty", [])}
    lines, amounts, problems = [], [], []
    for n, row in enumerate(rows, start=1):
        booked = read_date(_cell(row, columns, "date"))
        valued = read_date(_cell(row, columns, "value_date"))
        amount = _ledger_amount(_cell(row, columns, "amount"), profile)
        if booked.value is None or booked.ambiguous:
            problems.append(f"line {n}: unreadable date")
            continue
        if amount is None:
            problems.append(f"line {n}: unreadable amount")
            continue
        party = _cell(row, columns, "counterparty")
        amounts.append(amount)
        lines.append({"booking_date": booked.value.isoformat(),
                      "value_date": valued.value.isoformat() if valued.value and not valued.ambiguous else None,
                      "direction": "debit" if amount < 0 else "credit", "amount": _plain(abs(amount)),
                      "running_balance": None, "description": _cell(row, columns, "kind") or None,
                      "counterparty_name": None if _norm(party) in placeholders else party or None,
                      "counterparty_account": _cell(row, columns, "counterparty_account") or None,
                      "memo": _cell(row, columns, "memo") or None, "original_amount": None, "original_currency": None})
    return lines, amounts, problems


LINE_DETAILS = ("bank_reference", "order_id", "end_to_end_id", "card_number", "purchase_date", "transaction_time", "terminal",
                "authorisation_code", "exchange_rate", "fee")  # 144: the line's facts beyond the money (type pack 1.2.0)


def _with(line: dict[str, Any], **values: str | None) -> None:
    """144: the line's details that have a value (a line keeps no empty detail keys)."""
    line.update({k: v for k, v in values.items() if k in LINE_DETAILS and v not in (None, "")})


def _clock(text: str | None) -> str | None:
    """A time of day as HH:MM or HH:MM:SS ("10:00", "101500" → "10:15:00", "2026-06-04T10:15:30" → "10:15:30")."""
    if not text:
        return None
    m = re.search(r"(?:T|^)(\d{1,2}):?(\d{2})(?::?(\d{2}))?(?:\.\d+)?(?:[+-]\d{2}:?\d{2}|Z)?$", text.strip())
    if m is None or int(m[1]) > 23 or int(m[2]) > 59 or (m[3] and int(m[3]) > 59):
        return None
    return f"{int(m[1]):02d}:{m[2]}" + (f":{m[3]}" if m[3] else "")


def _card_details(found: dict[str, dict[str, str]]) -> dict[str, str | None]:
    """144: a line's card and charge details from the matches of a profile's `details` patterns (the camt.053 entry's
    additional information, the Erste PDF statement's card block): the masked card number, the purchase day and time,
    the terminal and the authorisation code, the exchange rate and the fee as the bank prints them."""
    card, terminal = found.get("card") or {}, found.get("terminal") or {}
    day = read_date(card["date"]).value if card.get("date") else None
    return {"card_number": card_numbers.masked(card.get("card")), "purchase_date": day.isoformat() if day else None,
            "transaction_time": _clock(card.get("time")), "terminal": terminal.get("terminal"),
            "authorisation_code": terminal.get("auth"),
            "exchange_rate": " ".join(found["rate"]["rate"].split()) if (found.get("rate") or {}).get("rate") else None,
            "fee": " ".join(found["fee"]["fee"].split()) if (found.get("fee") or {}).get("fee") else None}


def _merge_companion(lines: list[dict[str, Any]], amounts: list[Decimal], pdf: dict[str, Any], currency: str
                     ) -> tuple[list[Decimal | None], list[str]]:
    """Each export line takes the first unused PDF line of its booking date and amount: its running balance, the
    merchant of a card line without a partner and a foreign purchase's original amount and currency."""
    used: set[int] = set()
    balances: list[Decimal | None] = []
    problems: list[str] = []
    for n, (line, amount) in enumerate(zip(lines, amounts), start=1):
        j = next((k for k, c in enumerate(pdf["lines"]) if k not in used and c["booked"] == line["booking_date"]
                  and c["amount"] == amount), None)
        if j is None:
            balances.append(None)
            problems.append(f"line {n}: not on the PDF statement")
            continue
        used.add(j)
        c = pdf["lines"][j]
        balances.append(c["balance"])
        line["running_balance"] = _plain(c["balance"])
        if not line["counterparty_name"] and (merchant := c["details"].get("merchant")):
            line["counterparty_name"] = merchant.get("value")
        original = c["details"].get("original") or {}
        value = _dot_amount(original.get("amount", ""))
        if original.get("currency") and original["currency"] != currency and value is not None:
            line["original_amount"], line["original_currency"] = _plain(value), original["currency"]
        _with(line, **_card_details(c["details"]))  # 144
    problems += [f"PDF line {k + 1}: not in the export" for k in range(len(pdf["lines"])) if k not in used]
    return balances, problems


def _ledger_account(account_no: str, number: str, rows: list[list[str]], columns: dict[str, int],
                    profile: dict[str, Any], path: Path, file_name: str, source: dict[str, Any]) -> Account:
    lines, amounts, problems = _ledger_lines(rows, columns, profile)
    currency = profile["currency"]
    start, end = _month(file_name, profile.get("period_month_from_filename"))
    companion_path = _companion_path(path, profile.get("companion"))
    opening = closing = None
    balances: list[Decimal | None] = [None] * len(lines)
    companion_problems: list[str] = []
    if companion_path is not None:
        _within_size(companion_path, _conf()["limits"])
        data = companion_path.read_bytes()
        pdf = parse_companion(_bounded(_pdf_lines(companion_path))[0], profile["companion"])
        source = {**source, "companion": {"file": companion_path.name, "sha256": hashlib.sha256(data).hexdigest()}}
        if pdf["account"] and re.sub(r"\D", "", pdf["account"]) != re.sub(r"\D", "", account_no):
            companion_problems.append("the PDF statement is of another account")
        else:
            if pdf["currency"] and pdf["currency"] != currency:
                companion_problems.append(f"the PDF statement's currency is {pdf['currency']}")
            balances, merged = _merge_companion(lines, amounts, pdf, currency)
            companion_problems += pdf["problems"] + merged
            opening, closing = pdf["opening"], pdf["closing"]
            start, end = pdf["period"][0] or start, pdf["period"][1] or end
    if start is None and lines:
        start, end = min(t["booking_date"] for t in lines), max(t["booking_date"] for t in lines)
    statement = {"statement_type": profile["statement_type"], "account_no": account_no, "account_iban": None,
                 "period_start": start, "period_end": end, "currency": currency,
                 "opening_balance": _plain(opening), "closing_balance": _plain(closing),
                 "total_debit": _plain(sum((-a for a in amounts if a < 0), Decimal(0))),
                 "total_credit": _plain(sum((a for a in amounts if a > 0), Decimal(0))), "transactions": lines}
    checked = opening is not None and closing is not None
    running = None
    if checked:
        ok, previous = 0, opening
        for amount, balance in zip(amounts, balances):
            ok += previous is not None and balance is not None and previous + amount == balance
            previous = balance
        running = {"ok": ok, "lines": len(lines)}
    checks = {"running_balance": running, "total_row": None,
              "closing_balance": opening + sum(amounts, Decimal(0)) == closing if checked else None}
    digits = re.sub(r"\D", "", account_no)
    return Account(key=f"{digits}:{currency}:{start}", title=f"…{digits[-4:]} {number}".strip(), currency=currency,
                   occurrence=1, statement=statement, checks=checks, problems=problems, source=source,
                   balances="companion" if checked else "none", companion_problems=companion_problems)


def _ledger_accounts(rows: list[list[str]], profile: dict[str, Any], path: Path, file_name: str,
                     source: dict[str, Any]) -> list[Account]:
    columns, groups = _ledgers(rows, profile)
    if columns is None:
        return []
    return [_ledger_account(account_no, number, group, columns, profile, path, file_name, source)
            for (account_no, number), group in groups.items() if account_no]


# --- 141: a statement printed as text (a PDF's text layer or the bank's text export) -----------------------------------


def _signed(text: str | None) -> Decimal | None:
    """A Hungarian amount whose minus sign may stand apart ("- 2.000.000,00")."""
    return _hu_amount(re.sub(r"^-\s+", "-", text.strip())) if text else None


def _comma_amount(text: str) -> Decimal | None:
    """An amount with a decimal comma and no thousands separator ("1000,000"), when it is a whole number of cents."""
    if not re.fullmatch(r"\d+(?:,\d+)?", text):
        return None
    value = Decimal(text.replace(",", "."))
    return value.quantize(CENT) if value == value.quantize(CENT) else None


def _text_line(m: re.Match[str], *, expect_name: bool, reference: re.Pattern[str] | None = None) -> dict[str, Any]:
    """A line's first row; a card statement's row prints the purchase day (`value`) and no running balance. 144: the
    bank's reference after the kind's semicolon, when it has the profile's `reference` shape."""
    kind, groups = " ".join(m["kind"].split()), m.groupdict()
    value = read_date(groups["value"]).value if groups.get("value") else None
    after = kind.split(";", 1)[1].strip() if ";" in kind else ""
    ref = reference.fullmatch(after) if reference is not None and after else None
    return {"booked": read_date(m["booked"]).value, "kind": kind.split(";", 1)[0].strip() or kind,
            "amount": _hu_amount(m["amount"]), "balance": _hu_amount(groups["balance"]) if groups.get("balance") else None,
            "has_balance": "balance" in groups, "value_date": value, "counterparty_name": None, "counterparty_account": None,
            "memo": None, "note": None, "original_amount": None, "original_currency": None, "expect_name": expect_name,
            "details": {"bank_reference": ref["value"] if ref else None}}


def _card_row(current: dict[str, Any], m: re.Match[str]) -> None:
    """144: a card charge's detail row (`1234 5678 9012 3456 20260105 101500; …`): the masked card number, the purchase
    day and the time, on the line it belongs to."""
    groups = m.groupdict()
    day = read_date(re.sub(r"^(\d{4})(\d{2})(\d{2})$", r"\1-\2-\3", groups["date"])).value if groups.get("date") else None
    current["details"].update({"card_number": card_numbers.masked(groups["card"]),
                               "purchase_date": day.isoformat() if day else None, "transaction_time": _clock(groups.get("time"))})


_DOUBLED = re.compile(r"(?<!\S)(?=\S*\d)(?=\S*(?:\.\.|,,|--))(?:([0-9.,\-])\1)+(?!\S)")


def _undoubled(line: str) -> str:
    """141: a number the PDF's text layer prints with every character twice ("22002266..0022..0055.." for
    "2026.02.05.", the CIB statement's first row after a page break) is read once. Only a whole token of digits and
    separators whose every character is doubled, with a doubled separator, so no real number is touched."""
    return _DOUBLED.sub(lambda m: m.group(0)[::2], line)


def parse_text_statement(lines: list[str], profile: dict[str, Any]) -> dict[str, Any]:
    """The statement's header fields and lines from its text lines and the profile's patterns. Pure: no file is read.
    A text keeps single spaces only, as a fixed-width export pads inside a description where the PDF does not."""
    lines = [_undoubled(ln.strip()) for ln in lines]
    text = "\n".join(lines)
    header = {k: re.search(v, text, re.MULTILINE) for k, v in profile["header"].items()}
    start_re, end_re = re.compile(profile["lines_start"]), re.compile(profile["lines_end"])
    first = next((i for i, ln in enumerate(lines) if start_re.search(ln)), len(lines))
    last = next((i for i in range(first + 1, len(lines)) if end_re.search(lines[i])), len(lines))
    drop = [re.compile(x) for x in profile.get("drop_rows", [])]

    def optional(key: str) -> re.Pattern[str] | None:
        return re.compile(profile[key]) if profile.get(key) else None

    # 144: a card charge's detail row (the card, the day, the time) is kept even where its digits fit a dropped row
    card_row, card_section, reference = optional("card_row"), optional("card_section"), optional("reference")
    body = [ln for ln in _without_page_breaks(lines[first + 1:last], profile.get("skip_blocks", []))
            if ln.strip("-= ") and ((card_row is not None and card_row.match(ln)) or not any(rx.search(ln) for rx in drop))]
    line_re, value_re = re.compile(profile["line"]), optional("value_line")
    account_re, original_re, original_row = optional("counterparty_account"), optional("original"), optional("original_row")
    named_first = profile.get("name_from") == "first_unlabelled"  # a card purchase's merchant follows its line
    labels = {k: re.compile(v) for k, v in profile["labels"].items()}
    labelled, dated = re.compile(profile["labelled"]), re.compile(r"^\d{4}\.\d{2}\.\d{2}\.")
    out: list[dict[str, Any]] = []
    problems: list[str] = []
    expect_value = False
    section_card: str | None = None  # 144: a credit card statement heads each card's purchases with its number
    for ln in body:
        if card_section is not None and (c := card_section.match(ln)):
            section_card = card_numbers.masked(c["card"])
            expect_value = False
            if out:
                out[-1]["expect_name"] = False
            continue
        if m := line_re.match(ln):
            out.append(_text_line(m, expect_name=named_first, reference=reference))
            if section_card:
                out[-1]["details"]["card_number"] = section_card
            if out[-1]["booked"] is None or out[-1]["amount"] is None or (out[-1]["has_balance"] and out[-1]["balance"] is None):
                problems.append(f"line {len(out)}: unreadable")
            expect_value = value_re is not None
            continue
        current = out[-1] if out else None
        if expect_value and current is not None and value_re is not None and (v := value_re.match(ln)):
            expect_value = False
            current["value_date"] = read_date(v["value_date"]).value
            rest = (v["text"] or "").strip()
            if account_re is not None and account_re.fullmatch(rest):
                current["counterparty_account"] = rest if "-" in rest else rest.replace(" ", "")
                current["expect_name"] = True
            elif rest:
                current["note"] = " ".join(rest.split())
                if original_re and (o := original_re.match(rest)) and (amount := _comma_amount(o["amount"])) is not None:
                    current["original_amount"], current["original_currency"] = amount, o["currency"]
                if card_row is not None and (c := card_row.match(current["note"])):
                    _card_row(current, c)
            continue
        expect_value = False
        if current is None:
            continue
        if dated.match(ln):
            problems.append(f"line {len(out)}: an unreadable row after it")
            continue
        if card_row is not None and (c := card_row.match(ln)):
            _card_row(current, c)
            continue
        if original_row is not None and (o := original_row.match(ln)):  # a card purchase's amount in its own currency
            if current["original_amount"] is None and (amount := _dot_amount(o["amount"])) is not None:
                current["original_amount"], current["original_currency"] = amount, o["currency"]
            continue
        found = False
        for key, rx in labels.items():
            if d := rx.match(ln):
                found = True
                value = " ".join(d["value"].split())
                if key in LINE_DETAILS:  # 144: an identifier the statement labels (the parties' end-to-end id, the order id)
                    current["details"][key] = current["details"].get(key) or value
                elif current[key] is None:
                    current[key] = value
        if not found and current["expect_name"] and not labelled.match(ln):
            current["counterparty_name"] = " ".join(ln.split())
        current["expect_name"] = False
    period = header.get("period")
    totals = header.get("totals")
    return {"number": header["number"]["value"] if header.get("number") else None,
            "period": (read_date(period["start"]).value, read_date(period["end"]).value) if period else (None, None),
            "currency": header["currency"]["value"] if header.get("currency") else None,
            "account": header["account"]["value"] if header.get("account") else None,
            "iban": header["iban"]["value"].replace(" ", "") if header.get("iban") else None,
            "opening": _signed(header["opening"]["value"]) if header.get("opening") else None,
            "closing": _signed(header["closing"]["value"]) if header.get("closing") else None,
            "totals": (_signed(totals["debit"]), _signed(totals["credit"])) if totals else None,
            "lines": out, "problems": problems}


def _text_account(parsed: dict[str, Any], profile: dict[str, Any], source: dict[str, Any]) -> Account:
    """One account's statement from a parsed text statement, with the checks of its own balances and totals."""
    lines = [t for t in parsed["lines"] if t["booked"] is not None and t["amount"] is not None]
    amounts = [t["amount"] for t in lines]
    currency = parsed["currency"] or ""
    for t in lines:  # an amount in the statement's own currency is no original amount
        if t["original_currency"] == currency:
            t["original_amount"] = t["original_currency"] = None
    transactions = [{"booking_date": t["booked"].isoformat(),
                     "value_date": t["value_date"].isoformat() if t["value_date"] else None,
                     "direction": "debit" if t["amount"] < 0 else "credit", "amount": _plain(abs(t["amount"])),
                     "running_balance": _plain(t["balance"]), "description": t["kind"] or None,
                     "counterparty_name": t["counterparty_name"], "counterparty_account": t["counterparty_account"],
                     "memo": t["memo"] or t["note"], "original_amount": _plain(t["original_amount"]),
                     "original_currency": t["original_currency"]} for t in lines]
    for line, t in zip(transactions, lines):
        _with(line, **t["details"])
    opening, closing = parsed["opening"], parsed["closing"]
    start, end = (d.isoformat() if d else None for d in parsed["period"])
    if start is None and transactions:
        start, end = min(t["booking_date"] for t in transactions), max(t["booking_date"] for t in transactions)
    debits = sum((-a for a in amounts if a < 0), Decimal(0))
    credits = sum((a for a in amounts if a > 0), Decimal(0))
    statement = {"statement_type": profile["statement_type"], "account_no": parsed["account"],
                 "account_iban": parsed["iban"], "period_start": start, "period_end": end, "currency": currency or None,
                 "opening_balance": _plain(opening), "closing_balance": _plain(closing), "total_debit": _plain(debits),
                 "total_credit": _plain(credits), "transactions": transactions}
    running = None
    if opening is not None and lines and all(t["has_balance"] for t in lines):  # a card statement prints none
        ok, previous = 0, opening
        for t in lines:
            ok += previous is not None and t["balance"] is not None and previous + t["amount"] == t["balance"]
            previous = t["balance"]
        running = {"ok": ok, "lines": len(lines)}
    totals = parsed["totals"]
    checks = {"running_balance": running,
              "total_row": None if totals is None or None in totals else (abs(totals[0]), totals[1]) == (debits, credits),
              "closing_balance": None if opening is None or closing is None else opening + sum(amounts, Decimal(0)) == closing}
    digits = re.sub(r"\D", "", parsed["account"] or (parsed["iban"] or "")[4:])
    number = parsed["number"]
    title = f"…{digits[-4:]} {int(number)}" if number else f"…{digits[-4:]}"
    return Account(key=f"{digits}:{currency}:{start}", title=title, currency=currency, occurrence=1, statement=statement,
                   checks=checks, problems=list(parsed["problems"]), source=source)


def _text_lines(path: Path, data: bytes, profiles: list[dict[str, Any]]) -> list[str]:
    """The lines of a PDF's text layer (the isolated PDF reader) or of a text export in one of the profiles' encodings."""
    if path.suffix.lower() == ".pdf":
        try:
            return _pdf_lines(path)
        except (ValueError, OSError) as exc:  # the PDF reader's refusals (size, pages, a parser error) are ValueErrors
            raise StatementTableError(f"the PDF cannot be read: {exc}") from exc
    for encoding in dict.fromkeys(e for p in profiles for e in p.get("encodings", ["utf-8"])):
        try:
            return data.decode(encoding).splitlines()
        except UnicodeDecodeError:
            continue
    raise StatementTableError("the text file is in none of the profiles' encodings")


def _read_text(path: Path, *, name: str | None = None) -> Reading:
    conf = _conf()
    _within_size(path, conf["limits"])
    data = path.read_bytes()
    profiles = [p for p in conf["profiles"] if p.get("layout") == "text" and path.suffix.lower() in p["suffixes"]]
    lines, repairs = _bounded([ln.strip() for ln in _text_lines(path, data, profiles)] if profiles else [])
    text = "\n".join(lines)
    file_name = name or path.name
    sha = hashlib.sha256(data).hexdigest()
    for profile in profiles:
        if all(re.search(s, text, re.MULTILINE) for s in profile["signature"]):
            account = _text_account(parse_text_statement(lines, profile), profile,
                                    {"file": file_name, "sha256": sha, "repairs": repairs})
            return Reading(profile=profile["name"], institution=profile["institution"], file_name=file_name,
                           sha256=sha, repairs=repairs, accounts=[account])
    raise StatementTableError("no statement of a known profile in the file")


# --- 141: a bank's ISO 20022 camt.053 statement ------------------------------------------------------------------------

CAMT = "urn:iso:std:iso:20022:tech:xsd:camt.053."


_ROOT = re.compile(rb"<(?![?!])(?:[\w.\-]+:)?([\w.\-]+)([^>]*)>")


def _is_camt(data: bytes) -> bool:
    """Whether the file's root element is a camt.053 document. 144: the root decides, not any mention of the namespace
    in the head: the Erste Excel 2003 XML export declares a `camt` prefix on its `Workbook` root, and since 141 every
    such export went to the camt reader, which refused it."""
    root = _ROOT.search(re.sub(rb"<!--.*?-->", b"", data[:4096], flags=re.DOTALL))
    return root is not None and root.group(1) == b"Document" and CAMT.encode("ascii") in root.group(2)


def parse_camt(data: bytes, limits: dict[str, Any]) -> list[dict[str, Any]]:
    """The statements of a camt.053 file as plain texts (defusedxml refuses a document type declaration, entities and
    external references). It runs in the isolated reader's helper process, as the export parsing does (138)."""
    from defusedxml import ElementTree

    try:
        root = ElementTree.fromstring(data)
    except Exception as exc:  # defusedxml raises its own and ElementTree's errors (as in _spreadsheetml_rows)
        raise StatementTableError(f"the XML statement cannot be read: {type(exc).__name__}") from exc
    ns = root.tag[1:].split("}")[0] if root.tag.startswith("{") else ""
    if not ns.startswith(CAMT):
        raise StatementTableError("not an ISO 20022 camt.053 statement")

    def q(path: str) -> str:
        return "/".join(f"{{{ns}}}{part}" for part in path.split("/"))

    def text(el: Any, path: str) -> str | None:
        found = el.find(q(path))
        return found.text.strip() if found is not None and found.text and found.text.strip() else None

    out, count = [], 0
    for stmt in root.iter(q("Stmt")):
        balances = {}
        for bal in stmt.findall(q("Bal")):
            amount = text(bal, "Amt")
            if amount is not None:
                balances[text(bal, "Tp/CdOrPrtry/Cd")] = ("-" if text(bal, "CdtDbtInd") == "DBIT" else "") + amount
        entries = []
        for n in stmt.findall(q("Ntry")):
            count += 1
            if count > limits["max_rows"]:
                raise StatementTableError("the statement is over the row limit")
            debit = text(n, "CdtDbtInd") == "DBIT"
            role = "Cdtr" if debit else "Dbtr"
            post = next((text(a, "Amt") for a in n.iter(q("PrtryAmt")) if text(a, "Tp") == "POSTBALANCE"), None)
            party = f"NtryDtls/TxDtls/RltdPties/{role}"
            charges = [(text(c, "Amt"), c.find(q("Amt")).get("Ccy") if c.find(q("Amt")) is not None else None)
                       for c in n.findall(q("Chrgs"))]
            entries.append({"amount": text(n, "Amt"), "debit": debit, "booked": text(n, "BookgDt/Dt"), "value": text(n, "ValDt/Dt"),
                            "code": text(n, "BkTxCd/Prtry/Cd"), "name": text(n, f"{party}/Nm"),
                            "account": text(n, f"{party}Acct/Id/IBAN") or text(n, f"{party}Acct/Id/Othr/Id"),
                            "memo": " ".join(u.text.strip() for u in n.iter(q("Ustrd")) if u.text and u.text.strip()) or None,
                            "info": text(n, "NtryDtls/TxDtls/AddtlTxInf") or text(n, "AddtlNtryInf"), "balance": post,
                            # 144: the bank's reference, the parties' end-to-end id, when it was accepted, its charges
                            "reference": text(n, "NtryDtls/TxDtls/Refs/TxId") or text(n, "NtryDtls/TxDtls/Refs/AcctSvcrRef")
                            or text(n, "AcctSvcrRef"),
                            "end_to_end": text(n, "NtryDtls/TxDtls/Refs/EndToEndId"),
                            "accepted": text(n, "NtryDtls/TxDtls/RltdDts/AccptncDtTm"), "charges": charges})
        out.append({"id": text(stmt, "Id"), "iban": text(stmt, "Acct/Id/IBAN"), "other": text(stmt, "Acct/Id/Othr/Id"),
                    "currency": text(stmt, "Acct/Ccy"), "bic": text(stmt, "Acct/Svcr/FinInstnId/BIC"),
                    "bank": text(stmt, "Acct/Svcr/FinInstnId/Nm"), "from": (text(stmt, "FrToDt/FrDtTm") or "")[:10] or None,
                    "to": (text(stmt, "FrToDt/ToDtTm") or "")[:10] or None,
                    "opening": balances.get("OPBD") or balances.get("PRCD"), "closing": balances.get("CLBD"),
                    "credit_total": text(stmt, "TxsSummry/TtlCdtNtries/Sum"), "debit_total": text(stmt, "TxsSummry/TtlDbtNtries/Sum"),
                    "entries": entries})
    if not out:
        raise StatementTableError("no statement in the camt.053 file")
    return out


def _xml_amount(value: str | None) -> Decimal | None:
    """An XML decimal ("7291.00", ".00", "-12.5"). 153 (review C1): at most 40 characters, by a pattern that does not
    backtrack (a 20,000-digit amount took seconds)."""
    if not value or len(value) > 40 or not re.fullmatch(r"-?(?:\d+(?:\.\d+)?|\.\d+)", value):
        return None
    return _cents(Decimal(value))


def _camt_account(s: dict[str, Any], profile: dict[str, Any], source: dict[str, Any]) -> Account:
    bank = profile.get("banks", {}).get(s["bic"] or "", {})
    kinds, details = bank.get("kinds", {}), {k: re.compile(v) for k, v in bank.get("details", {}).items()}
    currency = s["currency"] or ""
    lines, amounts, balances, problems = [], [], [], []
    for n, e in enumerate(s["entries"], start=1):
        amount, booked = _xml_amount(e["amount"]), read_date(e["booked"]).value
        if amount is None or booked is None:
            problems.append(f"line {n}: unreadable")
            continue
        amount = -amount if e["debit"] else amount
        info = (e["info"] or "")[:LINE_MAX]  # 153 (review C1): the bank's patterns read a bounded text
        name = " ".join(e["name"].split()) if e["name"] else None  # single spaces, as the text statements keep them
        if not name and "merchant" in details and (m := details["merchant"].search(info)):
            name = " ".join(m["value"].split())
        original = details["original"].search(info) if "original" in details else None
        foreign = original is not None and original["currency"] != currency and _dot_amount(original["amount"]) is not None
        balance = _xml_amount(e["balance"])
        amounts.append(amount)
        balances.append(balance)
        value = read_date(e["value"]).value
        lines.append({"booking_date": booked.isoformat(), "value_date": value.isoformat() if value else None,
                      "direction": "debit" if amount < 0 else "credit", "amount": _plain(abs(amount)),
                      "running_balance": _plain(balance), "description": kinds.get(e["code"] or "", e["code"]),
                      "counterparty_name": name, "counterparty_account": e["account"],
                      "memo": " ".join(e["memo"].split()) if e["memo"] else None,
                      "original_amount": _plain(_dot_amount(original["amount"])) if foreign else None,
                      "original_currency": original["currency"] if foreign else None})
        # 144: the entry's references and its card and charge details (the bank's own text first, else its charges)
        found = {k: m.groupdict() for k, rx in details.items() if (m := rx.search(info))}
        card = _card_details(found)
        charged = [(a, c) for a, c in e.get("charges", []) if (v := _xml_amount(a)) is not None and v != 0]
        card["fee"] = card["fee"] or (" + ".join(f"{_plain(_xml_amount(a))} {c or currency}" for a, c in charged) or None)
        card["transaction_time"] = card["transaction_time"] or _clock(e.get("accepted"))
        end_to_end = e.get("end_to_end")
        _with(lines[-1], bank_reference=e.get("reference"), **card,
              end_to_end_id=None if not end_to_end or end_to_end.upper() == "NOTPROVIDED" else end_to_end)
    opening, closing = _xml_amount(s["opening"]), _xml_amount(s["closing"])
    start = read_date(s["from"]).value if s["from"] else None
    end = read_date(s["to"]).value if s["to"] else None
    if start is None and lines:
        start, end = (read_date(x).value for x in (min(t["booking_date"] for t in lines), max(t["booking_date"] for t in lines)))
    debits = sum((-a for a in amounts if a < 0), Decimal(0))
    credits = sum((a for a in amounts if a > 0), Decimal(0))
    iban = (s["iban"] or "").replace(" ", "") or None
    statement = {"statement_type": profile["statement_type"], "account_no": None if iban else s["other"], "account_iban": iban,
                 "period_start": start.isoformat() if start else None, "period_end": end.isoformat() if end else None,
                 "currency": currency or None, "opening_balance": _plain(opening), "closing_balance": _plain(closing),
                 "total_debit": _plain(debits), "total_credit": _plain(credits), "transactions": lines}
    running = None
    if opening is not None and lines and all(b is not None for b in balances):
        ok, previous = 0, opening
        for amount, balance in zip(amounts, balances):
            ok += previous + amount == balance
            previous = balance
        running = {"ok": ok, "lines": len(lines)}
    printed = (_xml_amount(s["debit_total"]), _xml_amount(s["credit_total"]))
    checks = {"running_balance": running,
              "total_row": None if None in printed else (abs(printed[0]), printed[1]) == (debits, credits),
              "closing_balance": None if opening is None or closing is None else opening + sum(amounts, Decimal(0)) == closing}
    digits = re.sub(r"\D", "", iban[4:]) if iban else re.sub(r"\D", "", s["other"] or "")
    return Account(key=f"{digits}:{currency}:{statement['period_start']}", title=f"…{digits[-4:]} {s['id'] or ''}".strip(),
                   currency=currency, occurrence=1, statement=statement, checks=checks, problems=problems, source=source)


def _read_camt(data: bytes, profile: dict[str, Any], limits: dict[str, Any], file_name: str, sha: str) -> Reading:
    try:
        statements = isolated_pdf.run(parse_camt, kind="read", data=data, limits=limits)
    except isolated_pdf.PdfReaderError as exc:
        name, _, message = str(exc).partition(": ")
        raise StatementTableError(message if name == StatementTableError.__name__ else str(exc)) from exc
    except isolated_pdf.PdfReaderLimit as exc:
        raise StatementTableError(f"the file could not be read within the reader's limits ({exc.reason}): {exc}") from exc
    source = {"file": file_name, "sha256": sha, "repairs": []}
    accounts = [_camt_account(s, profile, source) for s in statements]
    bank = profile.get("banks", {}).get(statements[0]["bic"] or "", {})
    return Reading(profile=profile["name"], institution=bank.get("institution") or statements[0]["bank"] or profile["institution"],
                   file_name=file_name, sha256=sha, repairs=[], accounts=accounts)


# --- reading a file or a folder ----------------------------------------------------------------------------------------


def _read_file(path: Path, *, name: str | None = None) -> Reading:
    if path.suffix.lower() in TEXT_SUFFIXES:
        return _read_text(path, name=name)
    if path.suffix.lower() not in SUFFIXES:
        raise StatementTableError(f"a statement table is a {' or '.join(SUFFIXES + TEXT_SUFFIXES)} file")
    conf = _conf()
    _within_size(path, conf["limits"])
    data = path.read_bytes()
    file_name = name or path.name
    sha = hashlib.sha256(data).hexdigest()
    camt = next((p for p in conf["profiles"] if p.get("layout") == "camt053"), None)
    if camt is not None and path.suffix.lower() == ".xml" and _is_camt(data):  # 141: a bank's ISO 20022 statement
        if len(data) > conf["limits"]["max_bytes"]:
            raise StatementTableError("the file is over the size limit")
        return _read_camt(data, camt, conf["limits"], file_name, sha)
    rows, repairs = _rows(path, data, conf)
    source = {"file": file_name, "sha256": sha, "repairs": repairs}
    for profile in conf["profiles"]:
        if profile.get("layout") in ("text", "camt053"):  # 141: read from a PDF, a text export or a camt.053 file only
            continue
        if profile.get("layout") == "ledger":
            accounts = _ledger_accounts(rows, profile, path, file_name, source)
        else:
            accounts = [replace(a, source=source) for a in _accounts(_blocks(rows, profile), file_name, profile)]
        if accounts:
            return Reading(profile=profile["name"], institution=profile["institution"], file_name=file_name,
                           sha256=sha, repairs=repairs, accounts=accounts)
    raise StatementTableError("no statement table of a known profile in the file")


def inside(path: Path, folder: Path) -> bool:
    """153 (review C8): a file of the folder that stays in it once its links are followed (`work.folder_documents` does
    the same); a link to a file elsewhere is passed over, as the other folder readers do."""
    try:
        return path.resolve().is_relative_to(folder.resolve())
    except (OSError, ValueError):
        return False


def _read_folder(folder: Path) -> Reading:
    """137: every export directly in the folder; a file no profile reads is listed, and a statement read twice (the
    same statement in two files) keeps its first file. 141: and every statement printed as text; a PDF or a text file
    no profile reads is not listed (a PDF may be an export's companion or a statement to extract, a text file a note)."""
    limit = _conf()["limits"]["max_files"]
    files = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in SUFFIXES + TEXT_SUFFIXES
                   and inside(p, folder))
    if len(files) > limit:
        raise StatementTableError(f"the folder holds more than {limit} statement files")
    readings, skipped = [], []
    for f in files:
        try:
            readings.append(_read_file(f))
        except StatementTableError as exc:  # listed for the person, not dropped
            if f.suffix.lower() not in TEXT_SUFFIXES:
                skipped.append({"file": f.name, "reason": str(exc)})
    accounts: dict[str, Account] = {}
    for r in readings:
        for a in r.accounts:
            if a.key in accounts:
                skipped.append({"file": r.file_name, "reason": f"the statement {a.title} is already in another file"})
            else:
                accounts[a.key] = a
    if not accounts:
        raise StatementTableError("no statement export of a known profile in the folder")
    parts = sorted({f"{a.source['file']}:{a.source['sha256']}:{(a.source.get('companion') or {}).get('sha256', '')}"
                    for a in accounts.values() if a.source})
    return Reading(profile=",".join(dict.fromkeys(r.profile for r in readings)), institution=readings[0].institution,
                   file_name=folder.name, sha256=hashlib.sha256("\n".join(parts).encode("utf-8")).hexdigest(),
                   repairs=sorted({x for r in readings for x in r.repairs}), accounts=list(accounts.values()),
                   skipped=skipped)


def read(path: Path, *, name: str | None = None) -> Reading:
    """Reads a statement table with the first profile that finds a table in it; 137: or every export of a folder.
    `name`: the file name to take the period from, when the path is a copy (the source instance) of the original."""
    path = Path(path)
    return _masked(_read_folder(path) if path.is_dir() else _read_file(path, name=name))


def _masked(reading: Reading) -> Reading:
    """144 (DECISIONS 144): no full card number leaves a reader; the texts keep only its first six and last four digits."""
    for account in reading.accounts:
        for line in account.statement.get("transactions") or []:
            for key in ("memo", "description", "counterparty_name"):
                line[key] = card_numbers.mask_text(line.get(key))
    return reading


def checks_ok(account: Account) -> bool:
    """Every line's running balance checks out (141: or, on a statement printing none, such as a card statement, the
    closing balance does), and no check failed and nothing was unreadable."""
    c = account.checks
    running = c["running_balance"]
    balances = running["ok"] == running["lines"] if running is not None else c["closing_balance"] is True
    return (balances and c["total_row"] is not False and c["closing_balance"] is not False and not account.problems
            and not account.companion_problems)


def survey(reading: Reading) -> dict[str, Any]:
    """The accounts found, without their lines and balances: what a person chooses from (the account number only by
    its last four characters). 137: whether the balances could be checked, and the companion PDF's name."""
    out = []
    for a in reading.accounts:
        dates = [t["booking_date"] for t in a.statement["transactions"]]
        iban = a.statement["account_iban"]
        companion = (a.source or {}).get("companion")
        out.append({"key": a.key, "account": "…" + iban[-4:] if iban else None, "title": a.title, "occurrence": a.occurrence,
                    "currency": a.currency, "lines": len(dates), "first": min(dates) if dates else None,
                    "last": max(dates) if dates else None, "period_start": a.statement["period_start"],
                    "period_end": a.statement["period_end"], "checks_ok": checks_ok(a),
                    "problems": len(a.problems) + len(a.companion_problems),
                    "balance_checked": a.balances != "none", "companion": companion["file"] if companion else None})
    return {"profile": reading.profile, "institution": reading.institution, "file": reading.file_name,
            "repairs": reading.repairs, "accounts": out, "skipped": len(reading.skipped)}


# --- one file per chosen account ------------------------------------------------------------------------------------------

FORMAT = "jav.statement_table"
FORMAT_VERSION = 1
_UNSAFE = re.compile(r'[<>:"/\|?*\x00-\x1f]+')


def derived_dir() -> Path:
    """Where the accounts' statement files are written: next to the store in use (tests get their own folder)."""
    from jav import store

    return store.current_path().parent / "statement_tables"


def _file_name(reading: Reading, account: Account) -> str:
    s = account.statement
    name = f"{reading.institution} - {account.title} ({account.currency}) {s['period_start'] or ''}_{s['period_end'] or ''}"
    if account.occurrence > 1:
        name += f" {account.occurrence}"
    return _UNSAFE.sub("_", name).strip(" .") + ".json"


def derive(reading: Reading, keys: list[str], out_dir: Path | None = None) -> list[Path]:
    """Writes one statement file per chosen account (`keys`, as `Account.key`), the work package item of that
    account: the statement, the export's name and fingerprint, the account and the reader's findings. The content is
    canonical, so the same account of the same export always gives the same file and the same item. An unknown key is
    refused, so nothing is written for a wrong choice."""
    by_key = {a.key: a for a in reading.accounts}
    unknown = [k for k in keys if k not in by_key]
    if unknown or not keys:
        raise StatementTableError(f"choose at least one account of the file; unknown: {len(unknown)}")
    folder = (out_dir or derived_dir()) / reading.sha256[:16]
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for key in dict.fromkeys(keys):
        a = by_key[key]
        path = folder / _file_name(reading, a)
        path.write_bytes(derived_bytes(reading, a))
        paths.append(path)
    return paths


def derived_bytes(reading: Reading, a: Account) -> bytes:
    """The content of an account's statement file (`derive`), canonical. 140: its fingerprint is the id of the item
    and of the document, so a survey knows it before anything is written (`derived_id`)."""
    source = a.source or {"file": reading.file_name, "sha256": reading.sha256, "repairs": reading.repairs}
    content = {"format": FORMAT, "version": FORMAT_VERSION, "profile": reading.profile, "institution": reading.institution,
               "source": {k: source[k] for k in ("file", "sha256", "repairs")},
               "account": {"key": a.key, "title": a.title, "currency": a.currency, "occurrence": a.occurrence},
               "checks": a.checks, "problems": a.problems, "statement": a.statement}
    if a.source is not None and a.balances != "export":  # 137: a ledger export; its companion PDF and findings
        content |= {"companion": source.get("companion"), "balances": a.balances, "companion_problems": a.companion_problems}
    return (json.dumps(content, ensure_ascii=False, sort_keys=True, indent=1) + "\n").encode("utf-8")


def derived_id(reading: Reading, a: Account) -> str:
    """The id an account's statement gets as an item and as a document: its statement file's fingerprint."""
    return hashlib.sha256(derived_bytes(reading, a)).hexdigest()


def load_derived(data: bytes) -> dict[str, Any]:
    """An account's statement file, checked for its format."""
    try:
        content = json.loads(data.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise StatementTableError("not a statement file") from exc
    if not isinstance(content, dict) or content.get("format") != FORMAT or content.get("version") != FORMAT_VERSION \
            or not isinstance(content.get("statement"), dict):
        raise StatementTableError("not a statement file of this version")
    return content

