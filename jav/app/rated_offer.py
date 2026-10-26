"""A rated offer as a Word document (231, stage D, D-f7): the insurance quote's document, built from a person's approved
premium decision.

Plain-language summary: once a person accepted a case's premium, the offer a client reads is printed from what the case
kept, never computed again: the premium decision's approved version names the premium record it accepted
(`evaluation_steps`), that record names the rule evaluation it was computed after, and the evaluation names the checked
input a person froze and the product version it was evaluated for. From these the output contract's `rated_offer`
(`configs/outputs.json`) gets its values - the policyholder and its tax number, the product, the day cover starts, the
risk's verdict, the modules' subtotal, the total (whether it was rounded, whether the policy's minimum applies), the
condition documents - and two tables: each cover module's premium, and each factor of the tariff on the total (its
title as the tariff names it), so a client sees how the total follows from the modules. The built-in template is
Hungarian (`configs/offer_documents.json`, document vocabulary); a person's own Word template is filled the same way
(`jav/app/outputs/docx.py`). The shared priced document (`priced_document@1`) prints it and stores it with the
decision's version, so the download of the version never changes.

- `OUTPUT`; `data(content, case_id=, package=, version=, approved_by=, approved_at=, output=)`: the output's data;
  `RatedOfferError` when the decision names no premium the case kept;
- `builtin_template(data)`: the built-in template for it; `file_name(content, package, version)`;
- `register()`: the output's built-in template, its labels and how a version becomes its data.

Layer: application; it names no process and no business (the product, its documents and its tariff are data).
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from jav.app import evaluation_steps, facts, naming, package_release, product_rules
from jav.app.outputs import contract
from jav.core import cfg
from jav.core.clock import utc_stamp as _now

OUTPUT = "rated_offer"
TABLE, ADJUSTMENTS = "modules", "adjustments"
COLUMNS = ("name", "premium")
ADJUSTMENT_COLUMNS = ("name", "factor")
SETTINGS = "offer_documents"
# the values the built-in template prints as lines, in their order; the rest stand elsewhere on the page
_LINES = ("policyholder", "tax_number", "product", "cover_start", "verdict")


class RatedOfferError(ValueError):
    """The decision names no premium the case kept (or the premium names no evaluation, the evaluation no input)."""


def settings() -> dict[str, Any]:
    return cfg.load(SETTINGS)


def _value_labels() -> dict[str, str]:
    return dict(settings()["labels"])


def _heads() -> dict[str, dict[str, str]]:
    return {table: dict(heads) for table, heads in settings()["heads"].items()}


def sample_heads() -> dict[str, dict[str, str]]:
    """The headers of the output's tables, for a template's marker list and preview."""
    return _heads()


def sample_values() -> dict[str, str]:
    return {k: v for k, v in _value_labels().items() if k in contract.spec(OUTPUT)["values"]}


def _kept(case_id: str, content: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any], facts.CheckedInput]:
    """The premium record the decision accepted, the evaluation it was computed after, and the checked input."""
    result = content.get("result") or {}
    try:
        premium = evaluation_steps.record(case_id, str(content["about"]), int(result["version"]))
        evaluation = evaluation_steps.record(case_id, premium["input_step"], premium["input_version"])
        frozen = package_release.own_version(case_id, evaluation["input_version"], step=evaluation["input_step"])
    except (KeyError, TypeError, ValueError) as exc:
        raise RatedOfferError(f"the decision names no premium the case kept: {exc}") from None
    if premium["kind"] != evaluation_steps.RATING_KIND or evaluation["kind"] != evaluation_steps.EVALUATION_KIND:
        raise RatedOfferError(f"step {content['about']} keeps no premium")
    return premium, evaluation, facts.checked_input(frozen["content"])


def data(content: dict[str, Any], *, case_id: str, package: str, version: int | None = None, approved_by: str | None = None,
         approved_at: str | None = None, output: str = OUTPUT) -> contract.OutputData:
    """The offer's output data from the approved premium decision (`content`, its version as approved; `version` None:
    marked a draft). `RatedOfferError` when the decision names no premium the case kept or the premium has no total."""
    premium, evaluation, checked = _kept(case_id, content)
    computed = premium["content"]
    if not computed.get("total"):
        raise RatedOfferError("the premium has no total")
    product = product_rules.product_at(evaluation["content"]["product"])
    factors = product_rules.resolve("tariff", computed["tariff_ref"])["content"].get("factors") or {}
    steps = computed.get("total_steps") or []
    before = steps[-1]["after"]["amount"] if steps else (computed.get("subtotal") or {}).get("amount")
    total = computed["total"]["amount"]
    texts = settings()
    top = next((u for u in checked.units if u.parent_id is None), None)
    holder = top.facts if top is not None else {}
    verdict = evaluation["content"].get("verdict")
    minimum = bool(computed.get("policy_minimum_applied"))
    values: dict[str, Any] = {
        "package": package, "version": version, "state": None if version else texts["texts"]["draft"],
        "offer_date": (approved_at or _now())[:10], "approved_by": approved_by,
        "policyholder": holder.get("company_name"), "tax_number": holder.get("tax_number"),
        "product": product.product.name, "cover_start": checked.cover_start.isoformat(),
        "verdict": texts["verdicts"].get(verdict, verdict) if verdict else None,
        "currency": computed["total"]["currency"], "subtotal": (computed.get("subtotal") or {}).get("amount"),
        "total": total,
        # rounded: the total is not the last factor's amount (the policy's minimum is named on its own)
        "rounded": texts["texts"]["rounded"] if before is not None and not minimum and Decimal(before) != Decimal(total) else None,
        "policy_minimum": texts["texts"]["minimum"] if minimum else None,
        "conditions": "; ".join(d.content.title for d in product.documents) or None}
    labels = _value_labels()
    names = {m.code: m.name for m in product.product.modules}
    totals = computed.get("module_totals") or {}
    rows = [[names.get(code, code), (totals.get(code) or {}).get("total", {}).get("amount")] for code in computed["modules"]]
    adjusted = [[(factors.get(s["name"]) or {}).get("title") or s["name"], s["factor"]] for s in steps]
    tables = tuple(_table(output, key, columns, found) for key, columns, found in
                   ((TABLE, COLUMNS, rows), (ADJUSTMENTS, ADJUSTMENT_COLUMNS, adjusted)))
    return contract.OutputData(output, tuple(contract.Value(k, labels.get(k, k), v) for k, v in values.items()), tables)


def _table(output: str, key: str, columns: tuple[str, ...], rows: list[list[Any]]) -> contract.Table:
    spec = contract.spec(output)["tables"][key]["columns"]
    heads = _heads()[key]
    return contract.table(key, key, columns, tuple(heads[c] for c in columns), rows, tuple(spec[c] for c in columns))


def file_name(_content: dict[str, Any], package: str, version: int | None) -> str:
    texts = settings()["texts"]
    slug = naming.safe_part(package, limit=40) or "offer"
    return f"{texts['file']}-{slug}-{f'v{version}' if version else texts['draft'].lower()}.docx"


# --- the built-in template --------------------------------------------------------------------------------------------


def _page_field(paragraph: Any, around: list[str]) -> None:
    """The page number field with its text before and after (`1. oldal`)."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    paragraph.add_run(around[0])
    paragraph._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    paragraph.add_run(around[1])


def _line(doc: Any, label: str, marker: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = 0
    p.add_run(f"{label}: ").bold = True
    p.add_run(marker)


def _right(doc: Any, label: str, marker: str, *, bold: bool = False) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = 0
    p.add_run(f"{label}: ").bold = bold
    p.add_run(marker).bold = bold


def _grid(doc: Any, key: str, columns: tuple[str, ...], heads: dict[str, str], widths: dict[str, float], right: str) -> None:
    """A table of the template: its header row, and one row of markers the filling repeats for each of the data's rows."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    table = doc.add_table(rows=2, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    for k, column in enumerate(columns):
        head, cell = table.rows[0].cells[k], table.rows[1].cells[k]
        table.columns[k].width = head.width = cell.width = Cm(widths[column])
        head.text = heads[column]
        head.paragraphs[0].runs[0].bold = True
        cell.text = "{{" + f"{key}.{column}" + "}}"
        for c in (head, cell):
            c.paragraphs[0].runs[0].font.size = Pt(10)
            if column == right:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def builtin_template(d: contract.OutputData) -> bytes:
    """The built-in Word template for the offer's data: a line or the factors' table is left out when the data has none,
    so the template depends on the data; it is filled like an own template."""
    from docx import Document
    from docx.shared import Cm

    has = {v.key for v in d.values if v.value not in (None, "")}
    currency = next((v.value for v in d.values if v.key == "currency"), None)
    adjusted = any(t.rows for t in d.tables if t.key == ADJUSTMENTS)
    s = settings()
    tx, labels, heads = s["texts"], s["labels"], _heads()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(s["page"]["width_cm"]), Cm(s["page"]["height_cm"])
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(s["page"]["margin_cm"]))
    doc.add_heading(tx["title"] + (" – {{state}}" if "state" in has else ""), level=1)
    _line(doc, labels["offer_date"], "{{offer_date}}")
    doc.add_paragraph(tx["intro"])
    for key in _LINES:
        if key in has:
            _line(doc, labels[key], "{{" + key + "}}")
    doc.add_paragraph()
    # a repeated row fills only its own columns, so the currency stands in the premium's header
    module_heads = {**heads[TABLE], "premium": heads[TABLE]["premium"] + (f" ({currency})" if currency else "")}
    _grid(doc, TABLE, COLUMNS, module_heads, s["widths_cm"][TABLE], "premium")
    if adjusted:  # the factors on the total, between the modules' subtotal and the total
        _right(doc, labels["subtotal"], "{{subtotal}} {{currency}}")
        doc.add_paragraph()
        _grid(doc, ADJUSTMENTS, ADJUSTMENT_COLUMNS, heads[ADJUSTMENTS], s["widths_cm"][ADJUSTMENTS], "factor")
    doc.add_paragraph()
    _right(doc, labels["total"], "{{total}} {{currency}}" + (" ({{rounded}})" if "rounded" in has else ""), bold=True)
    if "policy_minimum" in has:  # a note under the total: the policy's minimum stands in for the computed premium
        note = doc.add_paragraph()
        note.alignment = doc.paragraphs[-1].alignment
        note.add_run("({{policy_minimum}})")
    if "conditions" in has:
        doc.add_paragraph()
        _line(doc, labels["conditions"], "{{conditions}}")
    if "approved_by" in has:
        _line(doc, labels["approved_by"], "{{approved_by}}")
    footer = section.footer.paragraphs[0]
    footer.text = "{{package}}" + (" · " + tx["version"] if "version" in has else "") + " · "
    _page_field(footer, tx["page"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def register() -> None:
    from jav.app.outputs import builtin, parts

    builtin.register_word(OUTPUT, builtin_template)
    parts.register_output_labels(parts.OutputLabels(OUTPUT, __name__, sample_heads, sample_values))
    parts.register_version_document(parts.VersionDocument(OUTPUT, __name__, data, reads_case=True, name=file_name))


__all__ = ["ADJUSTMENTS", "ADJUSTMENT_COLUMNS", "COLUMNS", "OUTPUT", "RatedOfferError", "builtin_template", "data", "file_name",
           "register", "sample_heads", "sample_values", "settings"]
