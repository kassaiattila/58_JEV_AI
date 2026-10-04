"""Run result export (054 K4): CSV / XLSX / JSON from the effective data of the selected run.

Source: the machine data of the run's items merged with the human corrections (`corrections`), with each field's
effective source location (page, source text), whether it was corrected, and the open to-dos, so every number can be
traced back to its source.

Tables:
- `documents`: one row per document, with the type's fields (wide);
- `datapoints`: one row per document × field: value, page, source text, location status, corrected, to-do on the field;
- `line_items`: the rows of the itemised lists, with the page of each row.
On top of these, the XLSX has a separate sheet per type and the utility cost report (`jav/report_utility.py`).

Reuse: the legacy project's `orchestrator/framework/tabular.py` helpers (sheet-name cleaning, header formatting,
BOM + CRLF CSV) and the formula guard of `ui/src/utils/csv.ts` are ported; the old database loader is not (we read
from SQLite).
Formula guard: in the CSV, text starting with `=`, `+`, `-`, `@`, TAB or CR gets a leading `'` (a number is not text
and stays as it is); in the XLSX every text is written as text, never as a formula.
"""

from __future__ import annotations

import csv
import io
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from jav import cfg, corrections, store, typepack, work
from jav.native_contracts import NATIVE_SUFFIXES

NUMBER = re.compile(r"^-?\d[\d\s.,]*$")
SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


# --- the run's records ---------------------------------------------------------------------------------------


def run_records(run_id: str) -> list[dict[str, Any]]:
    """Effective data of the run's documents with source locations. Email items are skipped (they have no extracted
    data); for an email attachment, `source_email` is the email's subject (058 K5.2: the attachment's origin)."""
    with store.session():  # 061: one connection for the queries of all items
        return _run_records(run_id)


def _run_records(run_id: str) -> list[dict[str, Any]]:
    run = work.get_run(run_id)
    subjects = _email_subjects(run["input"]["items"])
    out = []
    for item in run["input"]["items"]:
        if item.get("kind") == "email":
            continue
        dp = corrections.datapoints_row(run_id, item["item_id"])
        if dp is None:
            continue
        corr = corrections.current(run_id, item["item_id"])
        effective = {**(dp.get("datapoints") or {}), **corr["fields"]}
        prov = corrections.effective_provenance(dp, corr, corrections.layer_for(dp))
        pack = typepack.get(dp["doc_type"])
        reasons = [r["reason"] for r in work.item_reasons(run_id, item["item_id"], item)["run"]]
        out.append({
            "item_id": item["item_id"], "file": Path(item["source_path"]).name, "doc_type": dp["doc_type"], "arm": dp.get("arm"),
            "final_status": dp.get("final_status"),
            "fields": {f: effective.get(f) for f in pack.header_fields},
            "kinds": {f: pack.kind(f) for f in pack.fields},
            "lists": {f: effective.get(f) or [] for f in pack.fields if pack.kind(f) == "list"},
            "provenance": {f: {k: v for k, v in (p or {}).items() if k in ("status", "page", "quote", "method", "multiple")}
                           for f, p in prov.items() if (p or {}).get("status") != "list"},
            "row_pages": {f: [r.get("page") for r in (p.get("rows") or [])] for f, p in prov.items() if (p or {}).get("status") == "list"},
            "pages": {f: (p or {}).get("page") for f, p in prov.items()},
            "corrected": sorted(corr["fields"]),
            "open_reasons": reasons,
            "source_email": subjects.get(item.get("parent_item_id") or ""),
        })
    return out


def native_records(run_id: str) -> list[dict[str, Any]]:
    """Native review data, including empty or unsuccessful outcomes and original machine proposals."""
    return [{"file": Path(item["source_path"]).name, **corrections.item_result(run_id, item["item_id"])}
            for item in work.get_run(run_id)["input"]["items"]
            if item.get("kind") == "document" and Path(item["source_path"]).suffix.lower() in NATIVE_SUFFIXES]


NATIVE_COLUMNS = ("file", "item_id", "fact_id", "entity", "property", "effective_value", "unit", "role",
                  "related_entity", "machine_state", "machine_grounding", "corrected", "confirmed", "source_citations",
                  "reading_status", "reading_gaps", "interpretation_status", "correctness", "result_version", "proposal")
NATIVE_HEAD = ("Document", "Item identifier", "Fact identifier", "Entity", "Property", "Value", "Unit", "Role",
               "Related entity", "Machine claim", "Machine grounding", "Corrected", "Confirmed", "Source citations",
               "Reading status", "Reading gaps", "Interpretation status", "Correctness", "Result version", "Original proposal")


def native_facts_table(records: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """One native fact per row; text and null remain distinct in the full JSON export."""
    rows = []
    for record in records:
        gaps = json.dumps(record["reading"]["results"], ensure_ascii=False, sort_keys=True)
        for fact in record["native_facts"]:
            proposal = fact["proposal"]
            rows.append([record["file"], record["item_id"], fact["fact_id"], proposal["entity"], proposal["property"],
                         fact["effective_value"], proposal["unit"], proposal["role"], proposal["related_entity"],
                         proposal["state"], fact["grounding"], fact["fact_id"] in record["correction"]["fields"],
                         fact["confirmed"], json.dumps(fact["native_citations"], ensure_ascii=False, sort_keys=True),
                         record["reading"]["status"], gaps, record["interpretation_outcome"]["status"],
                         (record["interpretation"] or {}).get("correctness"), record["result_version"],
                         json.dumps(proposal, ensure_ascii=False, sort_keys=True)])
    return list(NATIVE_HEAD), rows


def _email_subjects(items: list[dict[str, Any]]) -> dict[str, str]:
    """Subjects of the attachments' parent emails (only if the run has attachments)."""
    from jav.emails import message_version

    parents = {i.get("parent_item_id") for i in items if i.get("parent_item_id")}
    out = {}
    for i in items:
        if i["item_id"] in parents:
            msg, status = message_version(i)  # 086 N04: the version the run processed
            # data (a cell of the download), not a label
            out[i["item_id"]] = ("(a levél futáskori változata nem található)" if status == "changed"
                                 else msg.get("subject") or ("(tárgy nélkül)" if msg else "(a levél nem olvasható)"))
    return out


def email_records(run_id: str) -> list[dict[str, Any]]:
    """058 K5.1: the run's email items: the email's data, the recognised intent with the manual correction, the
    suggested next step (computed in code from the corrected intent), the recognition of the attachments, and how much
    of the email text the recognition saw. For an old run without its own result row, the email's latest result
    (`from_this_run=False`)."""
    from jav import emails, mailbox

    out = []
    for item in work.get_run(run_id)["input"]["items"]:
        if item.get("kind") != "email":
            continue
        path = Path(item["source_path"])
        msg, source_status = emails.message_version(item)  # 086 N04: the version the run processed, not the latest
        raw, own = mailbox.email_result_for(work.flow_run_id(run_id, item["item_id"]), path.parent.name)
        corr = corrections.current(run_id, item["item_id"])
        res = mailbox.effective_email_result(raw, corr["fields"].get("intent")) if raw is not None else {}
        body = (raw or {}).get("body") if own else None
        if not body:
            body = emails.body_coverage(emails.EmailMessage(message_id=path.parent.name, body=msg.get("body") or ""))
        sender = msg.get("sender_name") or ""
        if msg.get("sender"):
            sender = f"{sender} <{msg['sender']}>" if sender else msg["sender"]
        out.append({"item_id": item["item_id"], "message_id": path.parent.name, "subject": msg.get("subject") or "",
                    "sender": sender, "received_at": msg.get("received_at"), "mailbox": msg.get("mailbox"),
                    "intent": res.get("intent"), "machine_intent": res.get("machine_intent"), "confidence": res.get("confidence"),
                    "corrected": bool(res.get("corrected")), "next_flow": res.get("next_flow"),
                    "attachments": res.get("attachments") or [], "from_this_run": own, "body": body,
                    "open_reasons": [r["reason"] for r in work.item_reasons(run_id, item["item_id"], item)["run"]],
                    "tasks": mailbox.task_view(work.flow_run_id(run_id, item["item_id"]), raw) if own else None,
                    "source_status": source_status})
    return out


def task_rows(mails: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """058 K5.3: the task proposals, one per row (with the email, the evidence and the human decision)."""
    out = []
    for m in mails:
        for t in (m.get("tasks") or {}).get("tasks") or []:
            d = t.get("decision") or {}
            out.append({"item_id": m["item_id"], "subject": m["subject"], "index": t["index"], "action": t["action"], "title": t["title"],
                        "due_date": t.get("due_date"), "assignee_hint": t.get("assignee_hint"),
                        "evidence": " | ".join(e["quote"] for e in t.get("evidence") or []),
                        "decision": d.get("decision"), "decided_by": d.get("actor"), "decided_at": d.get("decided_at"),
                        "done_by": d.get("done_by"), "done_at": d.get("done_at")})  # 062: manual "done" mark
    return out


EMAIL_HEAD = ["Tárgy", "Feladó", "Érkezett", "Postafiók", "Szándék", "Gépi szándék", "Valószínűség", "Javítva",
              "Javasolt következő lépés", "Csatolmányok", "Levél szövege", "Levél hossza (karakter)", "Nyitott teendők", "Tétel-azonosító"]


def emails_table(records: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """The email table with Hungarian captions (intent, next step, the part of the text seen, attachment type)."""
    from jav import mailbox

    flows = mailbox.next_flow_labels()
    body = cfg.load("datasets")["labels"]["body_status"]
    doc_types = cfg.load("field_labels")["doc_types"]
    att = lambda a: f"{a.get('filename')} ({doc_types.get(a.get('doc_type'), a.get('doc_type')) or a.get('status') or '–'})"  # noqa: E731
    rows = [[r["subject"], r["sender"], r["received_at"], r["mailbox"], mailbox.intent_name(r["intent"]),
             mailbox.intent_name(r["machine_intent"]), r["confidence"], "igen" if r["corrected"] else "",
             flows.get(r["next_flow"] or "", r["next_flow"]), ", ".join(att(a) for a in r["attachments"]),
             body.get(r["body"]["status"], r["body"]["status"]), r["body"]["chars"], "; ".join(r["open_reasons"]), r["item_id"]]
            for r in records]
    return EMAIL_HEAD, rows


def _field_reason(reasons: list[str], field: str) -> bool:
    return any(len(r.split(":")) > 2 and r.split(":")[2] == field for r in reasons)


# --- tables ----------------------------------------------------------------------------------------------------

DOC_HEAD = ["Irat", "Tétel-azonosító", "Típus", "Út", "Állapot", "Nyitott teendők", "Forrás-levél"]


def _doc_row(r: dict[str, Any]) -> list[Any]:
    return [r["file"], r["item_id"], r["doc_type"], r["arm"], r["final_status"], "; ".join(r["open_reasons"]),
            r.get("source_email") or ""]


def documents_table(records: list[dict[str, Any]], doc_type: str | None = None) -> tuple[list[str], list[list[Any]]]:
    """One row per document; the fields are the union of the types' fields (in order of first appearance)."""
    recs = [r for r in records if doc_type is None or r["doc_type"] == doc_type]
    fields: list[str] = []
    for r in recs:
        fields += [f for f in r["fields"] if f not in fields]
    return DOC_HEAD + fields, [_doc_row(r) + [r["fields"].get(f) for f in fields] for r in recs]


def datapoints_table(records: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    head = ["Irat", "Tétel-azonosító", "Típus", "Mező", "Érték", "Oldal", "Forrásszöveg", "Hely", "Javítva", "Teendő a mezőn"]
    rows = []
    for r in records:
        for f, v in r["fields"].items():
            p = r["provenance"].get(f) or {}
            rows.append([r["file"], r["item_id"], r["doc_type"], f, v, p.get("page"), p.get("quote"), p.get("status"),
                         "igen" if f in r["corrected"] else "", "igen" if _field_reason(r["open_reasons"], f) else ""])
    return head, rows


def line_items_table(records: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    cols: list[str] = []
    for r in records:
        for rows in r["lists"].values():
            for row in rows:
                cols += [c for c in (row if isinstance(row, dict) else {"érték": row}) if c not in cols]
    head = ["Irat", "Tétel-azonosító", "Típus", "Lista", "Sor", "Oldal"] + cols
    out = []
    for r in records:
        for lst, rows in r["lists"].items():
            pages = r["row_pages"].get(lst) or []
            for i, row in enumerate(rows):
                row = row if isinstance(row, dict) else {"érték": row}
                out.append([r["file"], r["item_id"], r["doc_type"], lst, i + 1, pages[i] if i < len(pages) else None]
                           + [row.get(c) for c in cols])
    return head, out


TABLES = {"documents": documents_table, "datapoints": datapoints_table, "line_items": line_items_table}


# --- CSV -------------------------------------------------------------------------------------------------------


def _safe_text(v: Any) -> str:
    if v is None:
        return ""
    s = str(v)
    prefixes = tuple(cfg.load("reports")["export"]["formula_prefixes"])
    if s.startswith(prefixes) and not NUMBER.match(s):
        return "'" + s  # guards against formula injection (OWASP CSV injection; the rule of the legacy csv.ts)
    return s


def csv_bytes(head: list[str], rows: list[list[Any]]) -> bytes:
    """With a UTF-8 BOM, `;` as separator and CRLF line endings: Hungarian Excel then opens it split into columns."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=cfg.load("reports")["export"]["csv_delimiter"], lineterminator="\r\n")
    w.writerow(head)
    for row in rows:
        w.writerow([_safe_text(v) for v in row])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


# --- XLSX ------------------------------------------------------------------------------------------------------


def _safe_sheet(name: str, used: set[str]) -> str:
    base = SHEET_BAD.sub("_", name)[:31] or "Lap"
    out, n = base, 2
    while out.casefold() in used:
        suffix = f" ({n})"
        out, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(out.casefold())
    return out


def _cell_value(v: Any, kind: str | None) -> Any:
    if v is None:
        return None
    if kind in ("money", "number") and not isinstance(v, (dict, list)):
        try:
            return Decimal(str(v))
        except InvalidOperation:
            return str(v)
    return v if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) else str(v)


def _write_sheet(ws, head: list[str], rows: list[list[Any]], kinds: list[str | None] | None = None) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(head)
    for row in rows:
        ws.append([_cell_value(v, kinds[i] if kinds and i < len(kinds) else None) for i, v in enumerate(row)])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"  # text is never a formula
            elif isinstance(c.value, Decimal):
                c.number_format = "#,##0.00"
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEEEF2")
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))


STATUS_FILL = {"missing": "F8D7DA", "overlap": "FFF3CD", "partial": "E2E3F3"}
STATUS_TEXT = {"missing": "hiányzik", "overlap": "átfedés", "partial": "részleges", "ok": ""}


def _utility_sheets(wb, report: dict[str, Any], used: set[str]) -> None:
    from openpyxl.styles import PatternFill

    months = report["months"]
    ws = wb.create_sheet(_safe_sheet("Közmű-költség", used))
    head = ["Fogyasztási hely", "Közmű", "Szolgáltató", "Csak tájékoztató", "Összesen (bruttó)"] + months
    rows = []
    for s in report["series"]:
        rows.append([s["address"], s["utility"], ", ".join(s["suppliers"]), "igen" if s["summary_only"] else "", s["total"]]
                    + [(s["cells"].get(m) or {}).get("amount") or STATUS_TEXT.get((s["cells"].get(m) or {}).get("status", ""), "")
                       for m in months])
    rows.append(["Mindösszesen (a tájékoztató sorok nélkül)", "", "", "", report["grand_total"]] + [""] * len(months))
    _write_sheet(ws, head, rows, [None, None, None, None, "money"] + ["money"] * len(months))
    for r_i, s in enumerate(report["series"], start=2):
        for m_i, m in enumerate(months, start=len(head) - len(months) + 1):
            st = (s["cells"].get(m) or {}).get("status")
            if st in STATUS_FILL:
                ws.cell(row=r_i, column=m_i).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
    src = wb.create_sheet(_safe_sheet("Közmű-források", used))
    srows = [[s["address"], s["utility"], m, "elszámolás" if x.get("settlement") else c["status"], x["amount"], x["file"],
              x["item_id"], x["page"], x["field"], "igen" if x["corrected"] else "", x["open_reasons"] or ""]
             for s in report["series"] for m, c in s["cells"].items() for x in c["sources"]]
    srows += [[None, None, None, "nem vetíthető: " + u["reason"], None, u["file"], u["item_id"], None, None, "", ""] for u in report["unplaced"]]
    srows += [[None, None, None, f"ismétlődő (azonos: {d['same_as_file']})", None, d["file"], d["item_id"], None, None, "", ""]
              for d in report.get("duplicates", [])]
    _write_sheet(src, ["Fogyasztási hely", "Közmű", "Hónap", "Állapot", "Összeg (hónapra jutó)", "Irat", "Tétel-azonosító",
                       "Oldal", "Mező", "Javítva", "Nyitott teendők"], srows, [None] * 4 + ["money"])


def xlsx_bytes(run_id: str, records: list[dict[str, Any]] | None = None) -> bytes:
    from openpyxl import Workbook

    from jav import report_utility

    records = run_records(run_id) if records is None else records
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    mails = email_records(run_id)
    native = native_records(run_id)
    if native:
        head, rows = native_facts_table(native)
        _write_sheet(wb.create_sheet(_safe_sheet("Native facts", used)), head, rows)
        _write_sheet(wb.create_sheet(_safe_sheet("Native readings", used)),
                     ["Document", "Item identifier", "Reading status", "Interpretation status", "Result version", "Reading details"],
                     [[record["file"], record["item_id"], record["reading"]["status"],
                       record["interpretation_outcome"]["status"], record["result_version"],
                       json.dumps(record["reading"], ensure_ascii=False, sort_keys=True)] for record in native])
    if mails:  # 058 K5.1: the email items on their own sheet
        head, rows = emails_table(mails)
        _write_sheet(wb.create_sheet(_safe_sheet("Levelek", used)), head, rows)
        tasks = task_rows(mails)
        if tasks:  # 058 K5.3: the task proposals on a separate sheet, with the decision
            actions = cfg.load("email_tasks")["actions"]
            decisions = {"accepted": "elfogadva", "rejected": "elvetve"}
            _write_sheet(wb.create_sheet(_safe_sheet("Feladatok", used)),
                         ["Levél", "Akció", "Feladat", "Határidő", "Felelős", "Bizonyíték", "Döntés", "Döntött", "Elvégezve", "Elvégezte",
                          "Tétel-azonosító"],
                         [[t["subject"], actions.get(t["action"], t["action"]), t["title"], t["due_date"], t["assignee_hint"], t["evidence"],
                           decisions.get(t["decision"] or "", "vár döntésre"), t["decided_by"], t["done_at"], t["done_by"], t["item_id"]]
                          for t in tasks])
        if not records:  # emails only: no empty document sheet
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()
    for doc_type in sorted({r["doc_type"] for r in records}):
        head, rows = documents_table(records, doc_type)
        kinds = next(r["kinds"] for r in records if r["doc_type"] == doc_type)
        _write_sheet(wb.create_sheet(_safe_sheet(doc_type, used)), head, rows, [None] * len(DOC_HEAD) + [kinds.get(h) for h in head[len(DOC_HEAD):]])
    head, rows = datapoints_table(records)
    _write_sheet(wb.create_sheet(_safe_sheet("Adatpontok", used)), head, rows)
    head, rows = line_items_table(records)
    list_kinds: dict[str, str] = {}
    for r in records:
        pack = typepack.get(r["doc_type"])
        for lf in pack.list_fields.values():
            for c, k in lf.items():
                list_kinds.setdefault(c, k)
    _write_sheet(wb.create_sheet(_safe_sheet("Tételsorok", used)), head, rows, [None] * 6 + [list_kinds.get(h) for h in head[6:]])
    _utility_sheets(wb, report_utility.build(records), used)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- JSON ------------------------------------------------------------------------------------------------------


def _json_export(run_id: str, records: list[dict[str, Any]] | None = None) -> tuple[bytes, int]:
    from jav import datasets, report_utility

    # Capture before every data read: an interleaved edit must make approval fail,
    # never give older exported values the version of an unseen correction.
    review_version = corrections.review_version(run_id)
    run = work.get_run(run_id)
    native_run = any(item.get("kind") == "document" and Path(item["source_path"]).suffix.lower() in NATIVE_SUFFIXES
                     for item in run["input"]["items"])
    # Caller-supplied rows have no bound version. A native or mixed export must
    # reload them after the version capture, using the versioned dataset cache.
    records = datasets.run_records(run_id) if records is None or native_run else records
    body = {"run_id": run_id, "recipe_id": run["recipe_id"], "recipe_version": run["recipe_version"], "params": run["params"],
            "mode": run["mode"], "approval": run["approval"], "created_at": run["created_at"],
            "documents": [{k: v for k, v in r.items() if k != "kinds"} for r in records],
            "emails": email_records(run_id),
            "utility_cost": report_utility.build(records)}
    native = native_records(run_id)
    if native:
        body["native_items"] = native
        body["review_version"] = review_version
    return json.dumps(body, ensure_ascii=False, indent=1, default=str).encode("utf-8"), len(records)


def json_bytes(run_id: str, records: list[dict[str, Any]] | None = None) -> bytes:
    return _json_export(run_id, records)[0]


def render(run_id: str, fmt: str, table: str = "documents") -> tuple[bytes, str, str, int]:
    """(content, media type, file name, row count): what the service's download endpoint returns."""
    from jav import datasets

    if fmt == "json":
        content, count = _json_export(run_id)
        return content, "application/json", f"{run_id}.json", count
    records = datasets.run_records(run_id)  # 056: cache shared with the tables and the utility report
    if fmt == "csv":
        head, rows = native_facts_table(native_records(run_id)) if table == "native_facts" else TABLES[table](records)
        return csv_bytes(head, rows), "text/csv; charset=utf-8", f"{run_id}-{table}.csv", len(rows)
    if fmt == "xlsx":
        return (xlsx_bytes(run_id, records), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"{run_id}.xlsx", len(records))
    raise ValueError(f"unknown export format: {fmt}")


__all__ = ["TABLES", "csv_bytes", "datapoints_table", "documents_table", "email_records", "emails_table", "json_bytes",
           "line_items_table", "render", "run_records", "xlsx_bytes"]
