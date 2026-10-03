"""Native library adapters used only inside the bounded reader worker.

The library opening/iteration/closing patterns follow the legacy provider seam;
these implementations add structural positions and explicit content inventories.
"""
from __future__ import annotations

import base64
import csv
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
import hashlib
from html.parser import HTMLParser
from io import BytesIO, StringIO
import json
from pathlib import PurePosixPath

from .contracts import (
    CellContent, CellLocator, ImageLocator, Issue, NativeValue, ReadLimits,
    SheetLocator, SourceElement, TextLocator, WordLocator,
)
from .limits import ReadFailure, inspect_package, limited, xml


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


class Reading:
    def __init__(self, name: str, mime: str, limits: ReadLimits):
        self.name, self.mime, self.limits = name, mime, limits
        self.elements: list[SourceElement] = []
        self.issues: list[Issue] = []
        self.children: list[dict] = []
        self.texts: dict[str, str] = {}
        self.rasters: list[dict] = []
        self.raster_pixels = 0
        self.source_layers: dict[str, dict] = {}
        self.status: str | None = None

    def issue(self, code: str, message: str, element_id: str | None = None):
        limited(len(self.issues) >= 900, "Issue count limit exceeded")
        self.issues.append(Issue(stage="reading", code=code, message=message[:512], element_id=element_id))

    def add(self, kind, locator, *, parent=None, **kwargs):
        limited(len(self.elements) >= min(4500, self.limits.visited_cells), "Element count limit exceeded")
        identity = f"e{len(self.elements)}"
        element = SourceElement(element_id=identity, parent_id=parent, order=len(self.elements),
                                kind=kind, locator=locator, **kwargs)
        self.elements.append(element)
        return identity

    def text(self, text: str, *, parent=None):
        data = text.encode("utf-8")
        limited(len(text) > 100_000, "Text element limit exceeded")
        digest = sha(data)
        self.texts[digest] = text
        return self.add("text", TextLocator(text_sha256=digest, start=0, end=len(text)),
                        parent=parent, text=text, availability="read" if text else "empty")

    def child(self, name: str, data: bytes, role="embedded", mime="application/octet-stream"):
        limited(len(self.children) >= min(900, self.limits.archive_entries), "Child count limit exceeded")
        limited(sum(c["byte_size"] for c in self.children) + len(data) > self.limits.expanded_bytes,
                "Child byte limit exceeded")
        self.children.append({"name": name[:512] or "unnamed", "data": base64.b64encode(data).decode(),
                              "byte_size": len(data), "sha256": sha(data), "role": role, "mime": mime})

    def output(self):
        if not self.elements and not self.issues:
            self.issue("unread_content", "No readable content")
        return {"parser": self.name, "mime": self.mime,
                "status": self.status or ("partial" if self.issues else "complete"),
                "elements": [e.model_dump(mode="json") for e in self.elements],
                "issues": [i.model_dump(mode="json") for i in self.issues],
                "children": self.children, "texts": self.texts,
                **({"rasters": self.rasters} if self.rasters else {}),
                **({"source_layers": self.source_layers} if self.source_layers else {})}


def native_value(value) -> NativeValue:
    if value is None:
        return NativeValue(kind="empty", lexical=None)
    if isinstance(value, bool):
        return NativeValue(kind="boolean", lexical=str(value).lower())
    if isinstance(value, int):
        return NativeValue(kind="integer", lexical=str(value))
    if isinstance(value, float):
        return NativeValue(kind="decimal", lexical=str(value))
    if isinstance(value, (date, datetime)):
        return NativeValue(kind="date", lexical=value.isoformat())
    return NativeValue(kind="text", lexical=str(value))


def excel(data: bytes, parts: dict, reading: Reading):
    import openpyxl
    # Both views are opened over the same frozen bytes. No recalculation occurs.
    book = openpyxl.load_workbook(BytesIO(data), data_only=False, keep_links=False)
    cached = None
    try:
        cached = openpyxl.load_workbook(BytesIO(data), data_only=True, keep_links=False)
        visited = 0
        for sheet in book.worksheets:
            hidden_sheet = sheet.sheet_state != "visible"
            parent = reading.add("sheet", SheetLocator(sheet=sheet.title), hidden=hidden_sheet)
            visited += sheet.max_row * sheet.max_column
            limited(visited > reading.limits.visited_cells, "Worksheet rectangle exceeds cell limit")
            merges = tuple(sheet.merged_cells.ranges)
            for row in sheet.iter_rows():
                for cell in row:
                    address = cell.coordinate
                    formula = cell.value if cell.data_type == "f" else None
                    if formula is not None and not isinstance(formula, str):
                        raise ReadFailure("unsupported", "unread_content", "Array or special formula requires an adapter")
                    cached_value = cached[sheet.title][address].value if formula else None
                    hidden = hidden_sheet or bool(sheet.row_dimensions[cell.row].hidden) or any(
                        d.hidden and d.min <= cell.column <= d.max for d in sheet.column_dimensions.values())
                    merged = next((str(r) for r in merges if address in r), None)
                    value = native_value(cell.value if not formula else None)
                    content = CellContent(value=value, formula=formula,
                        cached_value=native_value(cached_value) if cached_value is not None else None,
                        cached_state=("unverified" if cached_value is not None else "missing") if formula else "not_applicable",
                        hidden=hidden, merged_range=merged)
                    element_id = reading.add("cell", CellLocator(sheet=sheet.title, cell=address,
                        row=cell.row, column=cell.column), parent=parent, cell=content, hidden=hidden,
                        text=formula or value.lexical, availability="empty" if value.kind == "empty" and not formula else "read")
                    if formula and cached_value is None:
                        reading.issue("formula_cache_missing", "Formula has no saved result; it was not evaluated", element_id)
                    if cell.number_format != "General" and value.kind != "empty":
                        reading.issue("unread_content", f"Display formatting retained only in original: {cell.number_format}", element_id)
        for name, content in parts.items():
            if "/media/" in name:
                reading.child(name, content)
                reading.issue("unread_content", f"Embedded media inventoried; worksheet anchor not yet mapped: {name}")
            elif name.startswith(("xl/charts/", "xl/comments", "xl/pivot", "xl/drawings/")) and name.endswith(".xml"):
                reading.issue("unread_content", f"Office part retained in original, not interpreted: {name}")
    finally:
        book.close()
        if cached is not None:
            cached.close()


def word(data: bytes, parts: dict, reading: Reading):
    from docx import Document
    from docx.table import Table
    document = Document(BytesIO(data))
    seen_images = set()
    seen_parts = set()

    def blocks(container, part_name, prefix="body", parent=None, depth=0):
        limited(depth > reading.limits.source_tree_depth, "Nested Word table depth exceeded")
        for index, block in enumerate(container.iter_inner_content()):
            location = WordLocator(part=part_name, structural_path=f"{prefix}/{index}", block_index=index)
            if isinstance(block, Table):
                table_id = reading.add("table", location, parent=parent)
                for row_index, row in enumerate(block.rows):
                    for col_index, cell in enumerate(row.cells):
                        blocks(cell, part_name, f"{prefix}/{index}/r{row_index}/c{col_index}", table_id, depth + 1)
            else:
                element_id = reading.add("paragraph", location, parent=parent, text=block.text,
                                         availability="read" if block.text else "empty")
                # Relationships give image bytes and the paragraph supplies its host.
                for blip in block._p.xpath(".//a:blip"):
                    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if not rid:
                        reading.issue("unread_content", "Image relationship is missing", element_id)
                        continue
                    image_part = block.part.related_parts[rid]
                    content = image_part.blob
                    name = str(image_part.partname).lstrip("/")
                    if name not in seen_images:
                        reading.child(name, content, mime=image_part.content_type)
                        seen_images.add(name)
                    image_id = reading.add("image", ImageLocator(image_sha256=sha(content), host=location),
                                           parent=element_id, availability="unreadable")
                    reading.issue("needs_ocr", "Image bytes and host retained; OCR has not run", image_id)

    blocks(document, "word/document.xml")
    for section in document.sections:
        for container in (section.header, section.first_page_header, section.even_page_header,
                          section.footer, section.first_page_footer, section.even_page_footer):
            if container.is_linked_to_previous:
                continue
            part_name = str(container.part.partname).lstrip("/")
            if part_name not in seen_parts:
                blocks(container, part_name, prefix=part_name)
                seen_parts.add(part_name)
    for name, content in parts.items():
        if "/media/" in name and name not in seen_images:
            reading.child(name, content)
            reading.issue("unread_content", f"Media without a mapped host: {name}")
        if name.startswith("word/") and name.endswith(".xml"):
            tags = {node.tag.rsplit("}", 1)[-1] for node in xml(content).iter()}
            gaps = tags & {"ins", "del", "sdt", "txbxContent", "altChunk", "object", "footnote", "endnote", "comment",
                           "gridSpan", "vMerge", "fldSimple", "instrText"}
            if gaps:
                reading.issue("unread_content", f"Uninterpreted Word structures in {name}: {', '.join(sorted(gaps))}")


def mail(data: bytes, reading: Reading):
    message = BytesParser(policy=policy.default).parsebytes(data)
    for key in ("Subject", "From", "To", "Date", "Message-ID"):
        if message.get(key):
            reading.text(f"{key}: {message[key]}")

    def visit(part, position="0", depth=0):
        limited(depth > reading.limits.source_tree_depth, "MIME depth limit exceeded")
        if part.defects:
            reading.issue("unread_content", "Malformed MIME structure; inventory may be incomplete")
        attachment = part.get_content_disposition() in ("attachment", "inline") or part.get_filename() is not None
        if part.get_content_type() == "message/rfc822":
            nested = part.get_payload()
            if isinstance(nested, list):
                for index, item in enumerate(nested):
                    reading.child(part.get_filename() or f"message-{position}-{index}.eml",
                                  item.as_bytes(policy=policy.default), role="attachment", mime="message/rfc822")
                reading.issue("unread_content", "Nested message is a library serialization, retained with its original parent")
            else:
                reading.issue("unread_content", "Nested message payload could not be decoded")
            return
        if part.is_multipart() and not attachment:
            for index, child in enumerate(part.iter_parts()):
                visit(child, f"{position}.{index}", depth + 1)
            return
        payload = part.get_payload(decode=True)
        if attachment or part.get_content_maintype() != "text":
            if payload is None:
                reading.children.append({"name": part.get_filename() or f"part-{position}", "missing": True,
                                         "byte_size": 0, "role": "attachment", "mime": part.get_content_type()})
                reading.issue("unread_content", "Attachment payload unavailable")
            else:
                reading.child(part.get_filename() or f"part-{position}", payload,
                              role="inline" if part.get_content_disposition() == "inline" else "attachment",
                              mime=part.get_content_type())
        else:
            try:
                text = (payload or b"").decode(part.get_content_charset() or "utf-8", errors="strict")
            except (UnicodeError, LookupError):
                reading.issue("unread_content", "Email body charset could not be decoded without loss")
                return
            reading.text(text)
            if part.get_content_type() == "text/html":
                reading.issue("unread_content", "HTML body retained as inert text; visual layout is not interpreted")
        if part.defects:
            reading.issue("unread_content", "MIME decoding reported a defect")
    visit(message)


def image(data: bytes, reading: Reading, *, recognise: bool = False):
    from PIL import Image
    from .visual import raster
    with Image.open(BytesIO(data)) as picture:
        frames = getattr(picture, "n_frames", 1)
        total = 0
        for index in range(frames):
            picture.seek(index)
            total += picture.width * picture.height
            limited(total > reading.limits.image_pixels, "Image pixel limit exceeded")
            picture.load()
            identity = reading.add("image", ImageLocator(image_sha256=sha(data), frame=index + 1,
                                   region=(0, 0, picture.width, picture.height)), availability="unreadable")
            reading.issue("needs_ocr", "Original image frame retained; OCR has not run", identity)
            if recognise:
                raster(reading, picture, element_id=identity, kind="image", page=index + 1)


class InertHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.pieces = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.pieces.append(data)


def read(data: bytes, filename: str, limits: ReadLimits, *, recognise: bool = False) -> dict:
    limited(len(data) > limits.input_bytes, "Input byte limit exceeded")
    if not data:
        raise ReadFailure("corrupt", "corrupt", "Empty input")
    suffix = PurePosixPath(filename).suffix.lower()
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise ReadFailure("unsupported", "unsupported", "OLE or encrypted Office container requires a separate reader")
    if data.startswith(b"%PDF-"):
        from .visual import read_pdf
        reading = Reading("native-pdf", "application/pdf", limits)
        read_pdf(data, reading, recognise=recognise)
    elif data.startswith(b"PK"):
        parts = inspect_package(data, limits)
        if "word/document.xml" in parts:
            reading = Reading("native-docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", limits)
            word(data, parts, reading)
        elif "xl/workbook.xml" in parts:
            reading = Reading("native-xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", limits)
            excel(data, parts, reading)
        else:
            raise ReadFailure("unsupported", "unsupported", "ZIP container has no supported Office document")
    elif data.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*", b"RIFF")):
        reading = Reading("native-image", "image/unknown", limits)
        image(data, reading, recognise=recognise)
    elif suffix == ".eml":
        reading = Reading("native-email", "message/rfc822", limits)
        mail(data, reading)
    elif suffix in {".txt", ".csv", ".json", ".xml", ".html", ".htm"}:
        text = data.decode("utf-8-sig", errors="strict")
        reading = Reading("native-" + suffix[1:], "text/plain", limits)
        if suffix == ".csv":
            for row_index, row in enumerate(csv.reader(StringIO(text)), 1):
                for col_index, value in enumerate(row, 1):
                    from openpyxl.utils import get_column_letter
                    reading.add("cell", CellLocator(sheet="CSV", cell=f"{get_column_letter(col_index)}{row_index}",
                        row=row_index, column=col_index), cell=CellContent(value=NativeValue(kind="text", lexical=value)),
                        text=value)
        else:
            if suffix == ".json":
                def unique(pairs):
                    result = dict(pairs)
                    if len(result) != len(pairs):
                        raise ValueError("Duplicate JSON key")
                    return result
                json.loads(text, object_pairs_hook=unique, parse_constant=lambda _: (_ for _ in ()).throw(ValueError("Nonfinite JSON")))
            if suffix == ".xml":
                xml(data)
            reading.text(text)
            if suffix in {".json", ".xml", ".html", ".htm"}:
                reading.issue("unread_content", "Exact text retained; structured paths or visual layout are not yet mapped")
    else:
        raise ReadFailure("unsupported", "unsupported", "No measured native reader for this content")
    return reading.output()


def evidence_view(response: dict) -> dict:
    """Binary children are frozen separately; structural evidence carries hashes."""
    result = {**response, "children": [{key: value for key, value in child.items() if key != "data"}
                                      for child in response["children"]]}
    if "rasters" in response:
        result["rasters"] = [{key: value for key, value in raster.items() if key != "data"}
                             for raster in response["rasters"]]
    return result
