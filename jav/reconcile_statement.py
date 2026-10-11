"""145 (backlog F-reconciliation E3b; DECISIONS 131, 132, 133, 145): a reconciliation package's result - its statement,
its approval as an unchangeable version that locks its lines, and its download.

- **The statement** (`build`) is computed in code from the package's workspace (`jav/reconcile_package.py`): every own
  invoice of the page with its result - paid, partly paid, settled by a credit note, a proposal still to decide, no
  payment found in the covered period, not covered by statements - and the lines that paid it; every statement line of
  the scope with its result - paired, partly paired, needs no invoice with its reason, left out by the code with its
  reason, open - and the invoices it paid; another party's invoice only when a line of the scope paid it, named with its
  party (133); and the warnings: a statement whose balances do not check out, and a month of a scope account without a
  statement (DECISIONS 132, 137: warnings, not blockers). Its fingerprint is the hash of its content.
- **The approval** (`approve`) records the statement the person saw (the request carries its fingerprint; a change
  since is a conflict) as the package's next version. It needs every source document from an approved run, or accepted
  by a person when it came from an evaluation on the command line (`accept_sources`, DECISIONS 132), and no decision
  left on a replaced statement (139). Open lines do not block it; the statement lists them (DECISIONS 131).
- **The lock** (enforced by `jav/reconcile_package.py`): the lines of an approved version take no new decision and lose
  none, in any package, and its package takes no decision and keeps its scope, until a person reopens the version with
  a reason (`reopen`); the version stays as an earlier one, and the next approval makes a new one (DECISIONS 132). An
  invoice is not locked: a payment from another account may still be recorded (an allocation never exceeds its amount).
- **The download** (`xlsx`, DECISIONS 145): an Excel workbook with two sheets, the invoices and the lines, each under a
  heading block naming the package, its party, period, accounts and version (or "draft"). The labels are data
  (`configs/reports.json` `reconcile_statement`; the column names of the fields from `configs/field_labels.json`), and
  a text is never written as a formula (the run export's rule, `jav/export.py`).

Nothing here calls an AI model or the network. Layer: UI/CLI → **this** (application operation) →
`jav.reconcile_package` (the workspace, the decisions and the approval tables) and `jav.work` (the package and its log).
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

from jav import cfg, export, reconcile_package as rp, store, work

FORMAT = 1  # the statement's content layout; a new layout bumps it (a stored version keeps its own)
INVOICE_RESULTS = ("paid", "partly_paid", "credited", "undecided", "no_payment_found", "not_covered")
LINE_RESULTS = ("paired", "partly_paired", "needs_no_invoice", "left_out", "open")
# the engine's invoice state (jav/reconcile.py STATUSES) -> the result, when no allocation decides it
_INVOICE_RESULT = {"confirmed": "paid", "credited": "credited", "partly_paid": "partly_paid", "proposed": "undecided",
                   "amount_only": "undecided", "amount_differs": "undecided", "rate_missing": "undecided",
                   "no_payment_found": "no_payment_found", "partly_covered": "not_covered", "not_covered": "not_covered"}
_LINE_RESULT = {"allocated": "paired", "partly_allocated": "partly_paired", "marked": "needs_no_invoice",
                "excluded": "left_out", "proposed": "open", "amount_only": "open", "open": "open"}
_INVOICE_FACTS = ("number", "supplier_name", "supplier_tax_id", "issue_date", "fulfillment_date", "due_date", "amount",
                  "currency", "payment_method", "file")
_LINE_FACTS = ("account", "statement_type", "currency", "booking_date", "value_date", "direction", "amount",
               "original_amount", "original_currency", "counterparty_name", "counterparty_account", "memo", "description",
               "bank_reference", "card_number", "purchase_date", "file", "statement_verified")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _money(value: Any) -> Decimal | None:
    try:
        return Decimal(str(value)) if value not in (None, "") else None
    except InvalidOperation:
        return None


def fingerprint(content: dict[str, Any]) -> str:
    return hashlib.sha256(json.dumps(content, sort_keys=True, ensure_ascii=False, default=str).encode("utf-8")).hexdigest()


# --- the statement ---------------------------------------------------------------------------------------------------------


def _invoice_result(row: dict[str, Any]) -> str:
    amount, allocated = _money(row.get("amount")), _money(row.get("allocated")) or Decimal(0)
    if row.get("state") == "credited":
        return "credited"
    if row.get("state") == "confirmed" or (amount is not None and allocated >= amount > 0):
        return "paid"
    if allocated > 0:
        return "partly_paid"
    return _INVOICE_RESULT.get(str(row.get("state")), "not_covered")


def build(ws: dict[str, Any], name: str) -> dict[str, Any]:
    """The statement of a package's workspace (`reconcile_package.workspace`); `name` is the package's name."""
    lines = {ln["id"]: ln for ln in ws["lines"]}
    rows = {r["id"]: r for r in ws["invoices"]}
    timing = (ws.get("overview") or {}).get("invoices") or {}
    paid_by_scope = {a["invoice_id"] for ln in ws["lines"] for a in ln["allocations"]}
    invoices = []
    for r in ws["invoices"]:
        if not r["own"] and r["id"] not in paid_by_scope:
            continue  # 133: another party's invoice only when a line of the scope paid it
        facts = timing.get(r["id"]) or {}
        paid_by = []
        for a in r["allocations"]:
            ln = lines.get(a["line_id"])
            paid_by.append({"line_id": a["line_id"], "line_amount": a["line_amount"], "invoice_amount": a["invoice_amount"],
                            "booking_date": ln.get("booking_date") if ln else None, "account": ln.get("account") if ln else None,
                            "in_scope": ln is not None})
        invoices.append({"id": r["id"], **{k: r.get(k) for k in _INVOICE_FACTS}, "result": _invoice_result(r),
                         "engine_state": r.get("state"), "allocated": r.get("allocated"), "rest": r.get("rest"),
                         "paid_by": paid_by, "paid_on": facts.get("paid_on"), "days_from_due": facts.get("days_from_due"),
                         "overdue_days": facts.get("overdue_days"), "double_payment": bool(facts.get("double_payment")),
                         "credited_by": (r.get("credited_by") or {}).get("number"),
                         "other_party": None if r["own"] else (r.get("party") or {}).get("name"), "source": r.get("source")})
    out_lines = []
    for ln in ws["lines"]:
        mark = ln.get("mark") or {}
        out_lines.append({"id": ln["id"], **{k: ln.get(k) for k in _LINE_FACTS}, "result": _LINE_RESULT.get(ln["state"], "open"),
                          "engine_state": ln["state"], "reason": mark.get("category") or ln.get("excluded_reason"),
                          "note": mark.get("note"), "allocated": ln.get("allocated"), "rest": ln.get("rest"),
                          "pays": [{"invoice_id": a["invoice_id"], "number": (rows.get(a["invoice_id"]) or {}).get("number"),
                                    "supplier_name": (rows.get(a["invoice_id"]) or {}).get("supplier_name"),
                                    "line_amount": a["line_amount"], "invoice_amount": a["invoice_amount"]}
                                   for a in ln["allocations"]]})
    warnings = [{"code": w["code"], "doc_id": w["doc_id"], "file": w["file"]} for w in ws["warnings"]]
    warnings += [{"code": "month_without_statement", "account": c["account"], "month": c["month"]}
                 for c in ws["coverage"] if not c["statements"]]
    sc = ws["scope"]
    return {"format": FORMAT, "package": {"id": ws["workpackage_id"], "name": name},
            "party": (sc.get("party") or {}).get("name"), "period": [sc["period_start"], sc["period_end"]],
            "accounts": list(sc["accounts"]), "engine_version": ws.get("engine_version"), "config_hash": ws.get("config_hash"),
            "invoices": invoices, "lines": out_lines, "warnings": warnings,
            "counts": {"invoices": dict(sorted(Counter(i["result"] for i in invoices).items())),
                       "lines": dict(sorted(Counter(ln["result"] for ln in out_lines).items())),
                       "open_lines": ws["counts"]["open_lines"]}}


def current(wp_id: str) -> dict[str, Any]:
    """The package's statement as it stands now, its fingerprint, what blocks its approval, its warnings and the
    package's approval (`reconcile_package.approval`)."""
    ws = rp.workspace(wp_id)
    content = build(ws, work.get(wp_id)["name"])
    return {"statement": content, "fingerprint": fingerprint(content), "blockers": ws["blockers"],
            "warnings": content["warnings"], "approval": ws["approval"]}


def approved_version(wp_id: str, n: int) -> dict[str, Any]:
    """An approved version of the package's statement as it was approved (KeyError: no such version)."""
    rp.scope(wp_id)
    with store.connect() as c:
        row = c.execute("SELECT * FROM reconcile_versions WHERE workpackage_id=? AND version=?", (wp_id, n)).fetchone()
    if row is None:
        raise KeyError(f"{wp_id} version {n}")
    meta = {k: row[k] for k in ("version", "fingerprint", "approved_by", "approved_at", "reopened_by", "reopened_at", "reopen_note")}
    return {"statement": json.loads(row["content"]), **meta}


# --- the approval -------------------------------------------------------------------------------------------------------


def approve(wp_id: str, *, fingerprint: str, actor: str) -> dict[str, Any]:
    """Approves the statement the person saw (its `fingerprint`) as the package's next version and locks its lines.
    `work.NotReady` with the blockers; `work.RevisionConflict` when the statement or a decision changed meanwhile."""
    state = rp._decision_state()  # read before the statement: a decision recorded after it is a conflict
    view = current(wp_id)
    if active := view["approval"]["active"]:
        raise rp.PackageError(f"version {active['version']} is approved; reopen it before approving again")
    if view["blockers"]:
        raise work.NotReady("the statement cannot be approved yet",
                            [{"code": b["code"], "message": b["file"] or b["doc_id"][:12]} for b in view["blockers"]])
    if fingerprint != view["fingerprint"]:
        raise work.RevisionConflict("the statement changed since it was shown; reload it and check it again")
    content, line_ids = view["statement"], [ln["id"] for ln in view["statement"]["lines"]]

    def write(c) -> int:
        n = (c.execute("SELECT MAX(version) FROM reconcile_versions WHERE workpackage_id=?", (wp_id,)).fetchone()[0] or 0) + 1
        c.execute("INSERT INTO reconcile_versions(workpackage_id, version, content, fingerprint, line_ids, approved_by, approved_at)"
                  " VALUES (?,?,?,?,?,?,?)", (wp_id, n, json.dumps(content, ensure_ascii=False, default=str), fingerprint,
                                              json.dumps(line_ids), actor, _now()))
        work.record_event(c, wp_id, "reconcile_approve", actor, {"version": n, "fingerprint": fingerprint, "lines": len(line_ids)})
        return n
    n = rp._write(state, write)
    return {k: v for k, v in approved_version(wp_id, n).items() if k != "statement"}


def reopen(wp_id: str, *, note: str, actor: str) -> dict[str, Any]:
    """Reopens the approved version with a reason: its lock goes, it stays as an earlier version."""
    note = (note or "").strip()
    if not note:
        raise rp.PackageError("reopening needs a reason")
    rp.scope(wp_id)
    with store.connect() as c:
        store.begin_immediate(c)
        row = c.execute("SELECT id, version FROM reconcile_versions WHERE workpackage_id=? AND reopened_at IS NULL",
                        (wp_id,)).fetchone()
        if row is None:
            raise rp.PackageError("the package has no approved statement to reopen")
        c.execute("UPDATE reconcile_versions SET reopened_by=?, reopened_at=?, reopen_note=? WHERE id=?",
                  (actor, _now(), note, row["id"]))
        work.record_event(c, wp_id, "reconcile_reopen", actor, {"version": row["version"], "note": note})
    return rp.approval(wp_id)


def accept_sources(wp_id: str, doc_ids: Iterable[str], *, actor: str) -> int:
    """Accepts the current results of source documents from an evaluation on the command line as sources (DECISIONS
    132): each must be such a blocker of the package; who and when is recorded. Returns how many were accepted."""
    ids = list(dict.fromkeys(doc_ids))
    if not ids:
        raise rp.PackageError("no document to accept")
    ws = rp.workspace(wp_id)
    blocking = {b["doc_id"] for b in ws["blockers"] if b["code"] == "source_no_run"}
    if unknown := [d for d in ids if d not in blocking]:
        raise rp.PackageError(f"not a command-line source of this package: {', '.join(d[:12] for d in unknown)}")
    snap, _sc = rp.scoped_snapshot(wp_id)
    results = {d["id"]: str(d.get("flow_run_id")) for d in [*snap["invoices"], *snap["statements"]]}
    with store.connect() as c:
        store.begin_immediate(c)
        now = _now()
        for d in ids:
            c.execute("INSERT OR IGNORE INTO reconcile_source_acceptances(doc_id, flow_run_id, workpackage_id, actor, accepted_at)"
                      " VALUES (?,?,?,?,?)", (d, results[d], wp_id, actor, now))
        work.record_event(c, wp_id, "reconcile_accept_sources", actor, {"documents": ids})
    return len(ids)


# --- the download -------------------------------------------------------------------------------------------------------


def _labels() -> dict[str, Any]:
    return cfg.load("reports")["reconcile_statement"]


def _yes_no(value: Any, lab: dict[str, Any]) -> str | None:
    return None if value is None else lab["yes"] if value else lab["no"]


def _join(parts: Iterable[Any]) -> str | None:
    text = "; ".join(str(p) for p in parts if p not in (None, ""))
    return text or None


def _invoice_table(st: dict[str, Any], lab: dict[str, Any]) -> tuple[list[str], list[list[Any]], list[str | None]]:
    fields, col = cfg.load("field_labels")["fields"], lab["columns"]
    head = [fields["invoice_number"], fields["supplier_name"], fields["supplier_tax_id"], fields["issue_date"],
            fields["fulfillment_date"], fields["due_date"], col["amount"], fields["currency"], fields["payment_method"],
            col["result"], col["allocated"], col["rest"], col["paid_on"], col["days_from_due"], col["overdue_days"],
            col["paid_by"], col["other_party"], col["source"], col["file"]]
    kinds = [None] * 6 + ["money", None, None, None, "money", "money", None, "number", "number", None, None, None, None]
    rows = [[i["number"], i["supplier_name"], i["supplier_tax_id"], i["issue_date"], i["fulfillment_date"], i["due_date"],
             i["amount"], i["currency"], i["payment_method"], lab["invoice_results"].get(i["result"], i["result"]),
             i["allocated"], i["rest"], i["paid_on"], i["days_from_due"], i["overdue_days"],
             _join(f"{p['booking_date'] or ''} {p['line_amount']}".strip() for p in i["paid_by"]), i["other_party"],
             lab["sources"].get(i["source"], i["source"]), i["file"]] for i in st["invoices"]]
    return head, rows, kinds


def _line_table(st: dict[str, Any], lab: dict[str, Any]) -> tuple[list[str], list[list[Any]], list[str | None]]:
    names, col = cfg.load("field_labels"), lab["columns"]
    cols = names["columns"]
    head = [cols["booking_date"], cols["value_date"], col["account"], cols["direction"], cols["amount"],
            names["fields"]["currency"], cols["original_amount"], cols["original_currency"], cols["counterparty_name"],
            cols["counterparty_account"], cols["memo"], cols["description"], col["result"], col["reason"], col["note"],
            col["allocated"], col["rest"], col["pays"], col["statement_file"], col["statement_verified"]]
    kinds = [None, None, None, None, "money", None, "money"] + [None] * 8 + ["money", "money", None, None, None]
    rows = [[ln["booking_date"], ln["value_date"], ln["account"], cols.get(f"direction={ln['direction']}", ln["direction"]),
             ln["amount"], ln["currency"], ln["original_amount"], ln["original_currency"], ln["counterparty_name"],
             ln["counterparty_account"], ln["memo"], ln["description"], lab["line_results"].get(ln["result"], ln["result"]),
             lab["reasons"].get(ln["reason"], ln["reason"]) if ln["reason"] else None, ln["note"], ln["allocated"], ln["rest"],
             _join(f"{p['number'] or ''} {p['supplier_name'] or ''} {p['line_amount']}".strip() for p in ln["pays"]),
             ln["file"], _yes_no(ln["statement_verified"], lab)] for ln in st["lines"]]
    return head, rows, kinds


def _heading(st: dict[str, Any], state: str, lab: dict[str, Any]) -> list[tuple[str, str]]:
    h = lab["heading"]
    return [(h["package"], st["package"]["name"]), (h["party"], st["party"] or ""), (h["period"], " – ".join(st["period"])),
            (h["accounts"], ", ".join(st["accounts"])), (h["state"], state)]


def _sheet(ws: Any, heading: list[tuple[str, str]], head: list[str], rows: list[list[Any]], kinds: list[str | None]) -> None:
    """A heading block, an empty row, then the table with its header frozen and filtered; a text is never a formula."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    for label, value in heading:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    top = ws.max_row + 1
    ws.append(head)
    for row in rows:
        ws.append([export._cell_value(v, kinds[i] if i < len(kinds) else None) for i, v in enumerate(row)])
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"  # text is never a formula (jav/export.py)
            elif isinstance(c.value, Decimal):
                c.number_format = "#,##0.00"
    for c in ws[top]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEEEF2")
    ws.freeze_panes = f"A{top + 1}"
    if rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(head))}{top + len(rows)}"
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))


_SLUG = re.compile(r"[^0-9A-Za-z]+")


def _shown(stamp: str | None) -> str:
    """A stored UTC time as the heading shows it (to the minute)."""
    return (stamp or "")[:16].replace("T", " ")


def xlsx(wp_id: str, *, version: int | None = None) -> tuple[bytes, str]:
    """The statement as an Excel workbook (DECISIONS 145: invoices and lines on two sheets) and its file name: the
    current draft, or an approved version as it was approved."""
    from openpyxl import Workbook

    lab = _labels()
    if version is None:
        st, state, suffix = current(wp_id)["statement"], lab["heading"]["draft"], lab["file"]["draft"]
    else:
        v = approved_version(wp_id, version)
        st, suffix = v["statement"], lab["file"]["version"].format(n=version)
        if v["reopened_at"] is None:
            state = lab["heading"]["approved"].format(n=version, by=v["approved_by"], at=_shown(v["approved_at"]))
        else:
            state = lab["heading"]["earlier"].format(n=version, by=v["reopened_by"], at=_shown(v["reopened_at"]),
                                                     note=v["reopen_note"] or "")
    heading = _heading(st, state, lab)
    wb = Workbook()
    used: set[str] = set()
    first = wb.active
    first.title = export._safe_sheet(lab["sheets"]["invoices"], used)
    _sheet(first, heading, *_invoice_table(st, lab))
    _sheet(wb.create_sheet(export._safe_sheet(lab["sheets"]["lines"], used)), heading, *_line_table(st, lab))
    buf = io.BytesIO()
    wb.save(buf)
    slug = _SLUG.sub("-", st["package"]["name"]).strip("-")[:40] or st["package"]["id"]
    return buf.getvalue(), f"{lab['file']['prefix']}-{slug}-{suffix}.xlsx"
