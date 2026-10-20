"""Datasets (056 U1): every list and result table of the UI is a named set of rows with a column description, which
the service delivers through one uniform query (`jav/tablequery.py`).

A dataset has a name, a label, a mandatory scope (e.g. `run_id`), a row builder and — if it is expensive — a
fingerprint. The UI draws only from the column description, so a new output needs no UI code (decision of
2026-09-28: a generic data viewer, not a custom UI per output). 172: a result table takes exactly one of a run
(`run_id`, the run's own result) and a package (`workpackage_id`, its current result: each item in use from the run
whose result counts now, `result_parts`); the run's table is built as before.

Cache: building a run's result tables is expensive (it recomputes the source locations of corrected fields; about 10 s
for 49 documents), while paging and filtering send a new request at every step. So the rows are kept until the
fingerprint of the run's data (corrections, to-dos, item status) changes; after a correction or a closed to-do, the
next request recomputes them.

Refresh behind (153): the package list is the one view a person waits on all day (it refreshes every 10 s) whose
rebuild got slow (about 5 s live on 2026-10-13: it builds every reconciliation package's workspace, which expire at
least every 120 s). Such a dataset (`refresh_behind`) answers from its last rows at once when its fingerprint has
changed, rebuilds in one background thread, and says so (`refreshing`); the next request after the rebuild gets the new
rows. A run's tables never do this: an approval names the version of the rows on screen (086).

Warm-up (158): a fresh service has no rows for the list yet, so its first request built them while the person waited
(5-8 s, reported on 2026-10-16 as "slow the first time"). The service now builds them in the background when it starts
(`warm_up`); a request for rows that are being built waits for that build instead of starting a second one; and the
cache pushes out a run's tables before a refresh-behind dataset, so the list's rows do not have to be built again after
a person has opened a few runs.

The list's own safety bound (159): the workspaces' 120 s bound (for a change their marks miss) also ended the list's
rows, so every two minutes the list built every reconciliation workspace again although nothing had changed (5-6 s on
one core, slowing every other request). The list's fingerprint and its next step now use a 15-minute bound
(`reconcile_package.LIST_KEEP_SECONDS`); a change the marks see still shows at once, and the pairing page keeps 120 s.

Column `extra` for the UI: `link` = link target (`run`, `workpackage`, `reviews`, `review`, `item`); the link is built
from the row's `_wp`, `_run`, `_stage` and `item_id` fields (`next`: the stage of the package's next step). `names`
(082, on the name column of an item list): which name the column shows (`original` or `unified`, the scope key `names`).
`wrap` (156): `anywhere` on a name-like text column (a name, a file name, a path, an e-mail address), whose long
unbroken value may break inside a word; every other column wraps at spaces only. `clamp` (183): the most lines a cell
of the column shows; the full value is its tooltip (the facts table's document, source and value, which made its rows
five to eight lines tall).
"""

from __future__ import annotations

import hashlib
import json
import logging
import threading
from collections import Counter, OrderedDict
from collections.abc import Callable
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from jav import activity_kinds, cfg, costs, logic_packs, store, work, work_views
from jav.outputs import parts
from jav.runtime import calls
from jav.tablequery import Column, Query, run_query, select, shown

Rows = tuple[list[Column], list[dict[str, Any]]]


class UnknownDataset(KeyError):
    """No dataset with this name."""


class MissingScope(ValueError):
    """The dataset needs a scope (e.g. a run id)."""


@dataclass(frozen=True)
class Dataset:
    name: str
    label: str
    scope: tuple[str, ...]                         # mandatory scope keys
    build: Callable[[dict[str, str]], Rows]
    fingerprint: Callable[[dict[str, str]], str] | None = None  # None: no cache (cheap list)
    optional_scope: tuple[str, ...] = ()
    natural_sort: tuple[tuple[str, bool], ...] = ()  # 062: (column, descending): row order when no sort is requested
    refresh_behind: bool = False  # 153: an outdated value is served at once and rebuilt in the background
    one_of: tuple[str, ...] = ()  # 172: exactly one of these scope keys (a result table: a run or a package)

    def spec(self) -> dict[str, Any]:
        return {"name": self.name, "label": self.label, "scope": list(self.scope), "optional_scope": list(self.optional_scope),
                "one_of": list(self.one_of), "natural_sort": [{"col": c, "desc": d} for c, d in self.natural_sort]}


def _labels(name: str) -> dict[str, str]:
    if name == "field":
        return cfg.load("field_labels")["fields"]
    if name == "doc_type":
        return cfg.load("field_labels")["doc_types"]
    if name == "recipe":
        return {r["id"]: r["title"] for r in work.recipes()}
    if name == "intent":  # 058 K5.1: the names of the email intents, from the registry
        from jav import intents

        return {k: v.display_name for k, v in intents.BY_KEY.items()}
    if name == "task_action":  # 058 K5.3: the Hungarian names of the task actions
        return dict(cfg.load("email_tasks")["actions"])
    if name == "provider":  # 082: the providers by the names the interface uses
        return dict(PROVIDER_NAMES)
    if name == "next_flow":
        from jav import mailbox

        return mailbox.next_flow_labels()
    return cfg.load("datasets")["labels"][name]


def _field_label(f: str) -> str:
    return cfg.load("field_labels")["fields"].get(f, f)


def _column_label(c: str) -> str:
    return cfg.load("field_labels")["columns"].get(c, c)


# 156: the keys of name-like columns; one such value cannot widen a whole table past the screen
_NAME_LIKE = frozenset({"name", "unified_name", "original_name", "file", "filename", "original", "source_path", "sender"})


def _col(key: str, label: str, kind: str = "text", *, labels: str | None = None, hidden: bool = False, **extra: Any) -> Column:
    if key in _NAME_LIKE and kind == "text":
        extra.setdefault("wrap", "anywhere")
    return Column(key, label, kind, hidden=hidden, labels=_labels(labels) if labels else None, extra=extra)  # type: ignore[arg-type]


# --- fingerprint -----------------------------------------------------------------------------------------------


def _run_fingerprint(scope: dict[str, str]) -> str:
    """Changes that affect the run's result: corrections, the to-dos of the run's subjects (including those from an
    earlier run, because of the "earlier" marker), task proposal decisions, item results, the run status. 062: the
    to-dos are narrowed to the run's subjects — until then, a to-do change in any run invalidated every run's stored
    result. 173: the decisions on the run's documents too (duplicate, reconciliation, statement choice), the same
    parts as the reviewed version (`corrections.review_version`)."""
    run_id = scope["run_id"]
    with store.connect() as c:
        row = c.execute("SELECT status, approval, input, excluded_items, workpackage_id FROM runs WHERE run_id=?", (run_id,)).fetchone()
        if row is None:
            raise KeyError(run_id)
        # 160-C (decision 163 D7): the items left out of the release (recorded at approval, until then the removed ones)
        left_out = row["excluded_items"] if row["approval"] else c.execute(
            "SELECT GROUP_CONCAT(item_id) FROM (SELECT item_id FROM workpackage_items WHERE workpackage_id=?"
            " AND removed_revision IS NOT NULL ORDER BY item_id)", (row["workpackage_id"],)).fetchone()[0]
        corr = c.execute("SELECT COUNT(*), MAX(created_at) FROM run_item_corrections WHERE run_id=?", (run_id,)).fetchone()
        rsn: list[Any] = []
        subjects = [work.review_subject(i) for i in json.loads(row["input"])["items"]]
        for kind in sorted({k for k, _ in subjects}):
            ids = sorted({i for k, i in subjects if k == kind})
            for start in range(0, len(ids), 500):  # stays under the SQLite parameter limit
                chunk = ids[start:start + 500]
                rsn += c.execute(
                    "SELECT COUNT(*), SUM(r.status='open'), MAX(r.id), MAX(r.closed_at) FROM review_reasons r"
                    f" JOIN review_queue q ON q.id = r.review_id WHERE q.subject_kind=? AND q.subject_id IN ({','.join('?' * len(chunk))})",
                    (kind, *chunk)).fetchone()
        # an email's flow id is `<run>:<item>`; a range condition, so that the key index can be used
        dec = c.execute("SELECT COUNT(*), GROUP_CONCAT(k) FROM (SELECT run_id || '#' || task_index || '=' || decision || '@' || decided_at AS k"
                        " FROM email_task_decisions WHERE run_id >= ? AND run_id < ? ORDER BY k)", (run_id + ":", run_id + ";")).fetchone()
        items = c.execute("SELECT COUNT(*), MAX(updated_at) FROM run_items WHERE run_id=?", (run_id,)).fetchone()
        from jav import native_results

        native = native_results.version_parts(run_id, c)
        from jav import duplicates, reconcile, statement_months

        documents = [i for k, i in subjects if k == "document"]
        # 126: a duplicate decision on a pair with one of the run's documents changes its mark
        pairs = [[d["pair_key"], d["decision"], d["decided_at"]] for d in duplicates.decisions_for(documents, c)]
        # 173 (plan 170, K20): so do a reconciliation decision and a statement choice, as in the reviewed version
        paid = [[d["pair_key"], d["decision"], d["decided_at"]] for d in reconcile.decisions_for(documents, c)]
        chosen = [[d["pair_key"], d["keep_doc_id"], d["decided_at"]] for d in statement_months.choices_for(documents, c)]
        # 175 (plan 170 E4): a decision kind a logic pack registers (`jav/activity_kinds.py`) beyond the parts read
        # above; a run without any keeps its earlier fingerprint
        subject_ids = [i for _k, i in subjects]
        extra = [[k.name, found] for k in activity_kinds.decisions() if k.name not in _FINGERPRINTED
                 and (found := k.rows(run_id, subject_ids, c))]
    digest = hashlib.sha256(str(dec[1]).encode("utf-8")).hexdigest()[:16]
    paired = hashlib.sha256(json.dumps(pairs).encode("utf-8")).hexdigest()[:16] if pairs else ""
    parts = [*corr, *rsn, dec[0], digest, *items, row["status"], row["approval"], native, paired,
             hashlib.sha256(str(left_out).encode("utf-8")).hexdigest()[:16]]
    if paid or chosen:  # a run without either keeps its earlier fingerprint
        parts.append(hashlib.sha256(json.dumps([paid, chosen]).encode("utf-8")).hexdigest()[:16])
    if extra:
        parts.append(hashlib.sha256(json.dumps(extra).encode("utf-8")).hexdigest()[:16])
    return "|".join(str(x) for x in parts)


# the decision kinds `_run_fingerprint` reads by itself (the built-in parts of the reviewed version)
_FINGERPRINTED = frozenset({"corrections", "tasks", "native", "duplicates", "reconcile", "statements"})


def _package_fingerprint(scope: dict[str, str]) -> str:
    """172: the package's result tables change with the items in use and their order, with which run's result counts
    for each item (and its tier), and with each counted run's own data (`_run_fingerprint`)."""
    from jav import package_result

    wp = work.get(scope["workpackage_id"])
    counted = package_result.in_use(wp)
    parts = [[i["item_id"], *((x.run_id, x.tier) if (x := counted.get(i["item_id"])) else ())] for i in wp["items"]]
    runs = sorted({x.run_id for x in counted.values()})
    return hashlib.sha256(json.dumps([parts, [_run_fingerprint({"run_id": r}) for r in runs]]).encode("utf-8")).hexdigest()


def _result_fingerprint(scope: dict[str, str]) -> str:
    """A result table's fingerprint: its run's, or for the package's current result, the package's (172)."""
    return _run_fingerprint(scope) if "run_id" in scope else _package_fingerprint(scope)


def _calls_fingerprint(scope: dict[str, str]) -> str:
    """082: also a call's settlement (a finished or a resolved uncertain call changes its state and cost)."""
    with store.connect() as c:
        row = c.execute("SELECT COUNT(*), MAX(id), MAX(finished_at), SUM(status IN ('reserved','uncertain')) FROM invocations"
                        " WHERE budget_scope=?", (scope["run_id"],)).fetchone()
    return "|".join(str(x) for x in row)


# 140 (Q-filter-slowness): what the package list reads besides the reconciliation (its packages' next step reads that):
# a package's change bumps its revision or writes an event, a run's status and approval change in place, and the
# items' results and the to-dos tell the counts.
LIST_INPUTS: dict[str, str | None] = {
    "workpackages": "SUM(revision), (SELECT group_concat(status || ':' || COALESCE(owner, ''), '|') FROM workpackages)",
    "workpackage_events": None,
    "workpackage_items": "COUNT(removed_revision)",
    "recipe_assignments": None,
    "runs": "(SELECT group_concat(run_id || ':' || status || ':' || COALESCE(approval, '') || ':' || COALESCE(finished_at, ''),"
            " '|') FROM runs)",
    "run_items": "SUM(status='done'), MAX(updated_at)",
    "review_reasons": "SUM(status='open'), MAX(closed_at)",
    "review_queue": "SUM(status='open')",
    "package_versions": None,  # 172: a released version ends the "Approve and release" step (also one recorded afterwards)
}


def _workpackages_fingerprint(scope: dict[str, str]) -> str:
    """140: the package list is built again only when what it shows may have changed (measured: 1.5-3.1 s per query
    before, at every filter change and refresh, as each reconciliation package's next step built its workspace)."""
    # first: what the logic packs' own next steps read (a reconciliation's workspace; it makes sure the tables it
    # reads are registered); 159: with the list's own safety bound. 176: from the logic packs
    own = [pack.list_mark() for pack in logic_packs.packs() if pack.list_mark is not None]
    with store.connect() as c:
        parts = [store.table_marks(c, LIST_INPUTS), *own]
    return hashlib.sha256(json.dumps(parts, default=str).encode("utf-8")).hexdigest()


def _run_items_fingerprint(scope: dict[str, str]) -> str:
    """082: the run items carry their cost, so the call log is part of their fingerprint."""
    return f"{_run_fingerprint(scope)}#{_calls_fingerprint(scope)}"


# --- row builders ----------------------------------------------------------------------------------------------


def _runs(scope: dict[str, str]) -> Rows:
    wp_id = scope.get("workpackage_id")
    if wp_id:
        work.get(wp_id)  # unknown package: 404
    cols = [
        _col("created_at", "Indítva", "datetime", link="run"),
        _col("workpackage_name", "Munkacsomag", link="workpackage"),
        _col("recipe_id", "Feldolgozás", "enum", labels="recipe"),
        _col("mode", "Mód", "enum", labels="mode"),
        _col("scope", "Run scope", "enum", labels="run_scope"),  # 160-C S3: a run over only the pending items
        _col("status", "Állapot", "enum", labels="run_status", badge=True),
        _col("approval", "Jóváhagyás", "enum", labels="approval"),
        _col("approved_by", "Jóváhagyta", hidden=True),
        _col("items_done", "Lefutott tétel", "number"),
        _col("items", "Tétel", "number"),
        _col("open_reasons", "Nyitott teendő", "number", alert=True),
        _col("cost_usd", "Költség (USD)", "number"),
        _col("actor", "Indította", hidden=True),
        _col("finished_at", "Befejeződött", "datetime", hidden=True),
        _col("run_id", "Futás-azonosító", "id", hidden=True, link="run"),
    ]
    titles = _labels("recipe")  # 058: the recipe's title (translated by the UI), not its code name
    runs = work.run_rows(wp_id)
    spent = costs.run_usd([r["run_id"] for r in runs])  # 082: the known actual cost of each run
    rows = [{**r, "_key": r["run_id"], "_run": r["run_id"], "_wp": r["workpackage_id"],
             "recipe": titles.get(r["recipe_id"], r["recipe_id"]), "cost_usd": spent[r["run_id"]]} for r in runs]
    return cols, rows


def _workpackages(scope: dict[str, str]) -> Rows:
    archived = scope.get("include_archived") == "1"  # 058: hidden packages only on request
    cols = [
        _col("name", "Név", link="workpackage"),
        *([_col("status", "Csomag", "enum", labels="wp_status", badge=True)] if archived else []),
        _col("next_label", "Következő lépés", link="next"),
        _col("last_status", "Utolsó futás", "enum", labels="run_status", badge=True),
        _col("items", "Tétel", "number"),  # 058: document or email
        _col("open_reasons", "Nyitott teendő", "number", link="reviews", alert=True),
        _col("open_reasons_all", "Nyitott teendő a korábbi futásokkal", "number", hidden=True),
        _col("older_runs", "Older runs to supersede", "number", link="process", alert=True),  # 153
        _col("source_kind", "Forrás", "enum", labels="source_kind"),
        _col("recipe_id", "Feldolgozás", "enum", labels="recipe"),
        _col("owner", "Felelős"),  # 061
        _col("last_activity", "Utolsó tevékenység", "datetime"),
        _col("created_at", "Létrehozva", "datetime"),  # 062: visible, because this is the list's default order
        _col("source_ref", "Forrás helye", hidden=True),
        _col("next_code", "Lépés kódja", "enum", hidden=True),
        _col("id", "Azonosító", "id", hidden=True),
    ]
    rows = [{**r, "_key": r["id"], "_wp": r["id"], "_run": r["last_run_id"], "_stage": r["next"]["stage"],
             "next_label": r["next"]["label"], "next_code": r["next"]["code"], "_next_params": r["next"]["params"]}
            for r in work_views.workpackage_list(include_archived=archived)]
    for r in rows:
        r.pop("next", None)
        if not logic_packs.has_runs(r["source_kind"]):  # 167 (review 153 E): a reconciliation's lines are not
            r["items"] = None                               # items; "0" read as empty
    if scope.get("owner"):  # 061: "Saját csomagjaim" (My packages) — the owner's name, case-insensitive
        who = " ".join(scope["owner"].split()).casefold()
        rows = [r for r in rows if (r.get("owner") or "").casefold() == who]
    return cols, rows


def _activity(scope: dict[str, str]) -> Rows:
    """061: a person's actions of the day ("Mai munkám", My work today); an empty day means today (local time)."""
    from jav import activity

    cols = [
        _col("at", "Időpont", "datetime"),
        _col("action", "Művelet", "enum", labels="activity_action"),
        _col("workpackage_name", "Munkacsomag", link="workpackage"),
        _col("detail", "Részlet"),
        _col("run_id", "Futás", "id", link="run"),
    ]
    titles = _labels("recipe")  # 062: the recipe's title (translated by the UI), not its code name
    rows = [{**e, "detail": titles.get(e["detail"], e["detail"]) if e["action"] == "recipe" else e["detail"],
             "_key": f"{e['at']}|{e['action']}|{n}", "_wp": e["workpackage_id"], "_run": e["run_id"]}
            for n, e in enumerate(activity.entries(scope["actor"], scope.get("day") or None))]
    return cols, rows


def _item_name(item: dict[str, Any], titles: dict[str, str]) -> str:
    return titles.get(item["item_id"]) or Path(item["source_path"]).name


NAME_CHOICES = ("original", "unified")


def _name_choice(scope: dict[str, str]) -> str:
    """082: which name the item list shows in its name column — the original (default) or the unified one."""
    choice = scope.get("names") or "original"
    if choice not in NAME_CHOICES:
        raise ValueError(f"unknown name choice: {choice} (allowed: {', '.join(NAME_CHOICES)})")
    return choice


def _name_cols() -> list[Column]:
    """082: the unified and the original name as columns that can be added, and the state of the unified name."""
    return [_col("unified_name", "Egységes név", hidden=True, link="review"),
            _col("original_name", "Eredeti név", hidden=True),
            _col("name_state", "Egységes név állapota", "enum", labels="name_state", hidden=True)]


def _name_fields(item: dict[str, Any], titles: dict[str, str], name: Any, choice: str) -> dict[str, Any]:
    """The name fields of one row: `name` follows the choice (an item without a unified name keeps its original
    name), `_name_check` says why a unified name is uncertain (for the UI's warning mark; not a column), and (183)
    `_name_reasons` the same as codes, which the UI writes in its own language."""
    original = _item_name(item, titles)
    return {"name": name.unified if choice == "unified" and name.unified else original, "original_name": original,
            "unified_name": name.unified, "name_state": name.state, "_name_check": name.check,
            "_name_reasons": list(name.reasons)}


def _workpackage_items(scope: dict[str, str]) -> Rows:
    from jav import naming

    choice = _name_choice(scope)
    wp = work.get(scope["workpackage_id"])
    titles = work_views.item_titles(wp["items"], work_views.package_root(wp))
    names = naming.package_item_names(wp)
    # 168 (plan F-item-list-and-adding): each item's state and the run whose result counts now (`work.item_states`)
    cols = [
        _col("name", "Tétel", link="review", names=choice),
        *_name_cols(),
        _col("state", "State", "enum", labels="package_item_state", badge=True),
        _col("run_mode", "Counted run", "enum", labels="mode", link="run"),
        _col("run_at", "Run started", "datetime", link="run", hidden=True),  # the mode says which run; its page has the time
        _col("final_status", "Result", "enum", labels="final_status"),
        _col("open_reasons", "Nyitott teendő", "number", link="review", alert=True),
        _col("kind", "Fajta", "enum", labels="kind", hidden=True),
        _col("added_revision", "Felvéve (verzió)", "number", hidden=True),
        _col("source_path", "Útvonal", hidden=True),
        _col("sha256", "Ujjlenyomat", "id", hidden=True),
        _col("item_id", "Tétel-azonosító", "id", hidden=True),
    ]
    open_by_subject = store.review_open_reasons_many([work.review_subject(i) for i in wp["items"]])
    owners = work_views.run_packages()  # 153: not another package's runs' to-dos
    states = work.item_states(wp["id"], wp["items"])
    rows = []
    for i in wp["items"]:
        st = states[i["item_id"]]
        reasons = len(work_views.own_reasons(open_by_subject[work.review_subject(i)], wp["id"], owners))
        rows.append({"_key": i["item_id"], "_wp": wp["id"], "_removed": False, "_run": st["run_id"], "item_id": i["item_id"],
                     **_name_fields(i, titles, names[i["item_id"]], choice), "state": st["state"], "run_mode": st["mode"],
                     "run_at": st["run_at"], "final_status": shown_final(st["final_status"], reasons),
                     "kind": i.get("kind") or "document", "source_path": i["source_path"], "sha256": i.get("sha256"),
                     "added_revision": i.get("added_revision"), "open_reasons": reasons})
    if scope.get("removed") == "1":  # 160-C S1 (decision 163 D3): the removed items too, with who, when and why
        removed = work.removed_items(wp["id"])
        titles.update(work_views.item_titles(removed, work_views.package_root(wp)))
        cols += [_col("removed_by", "Removed by"), _col("removed_at", "Removed at", "datetime"),
                 _col("removed_reason", "Reason for removal")]
        rows = [{**r, "removed_by": None, "removed_at": None, "removed_reason": None} for r in rows]
        rows += [{"_key": i["item_id"], "_wp": wp["id"], "_removed": True, "_run": None, "item_id": i["item_id"],
                  **_name_fields(i, titles, naming.NO_NAME, choice), "state": "removed", "run_mode": None, "run_at": None,
                  "final_status": None, "kind": i.get("kind") or "document",
                  "source_path": i["source_path"], "sha256": i.get("sha256"), "added_revision": i.get("added_revision"),
                  "open_reasons": None, "removed_by": i["removed_by"], "removed_at": i["removed_at"],
                  "removed_reason": i["removed_reason"]} for i in removed]
    return cols, rows


def _reconcile_inputs(scope: dict[str, str]) -> Rows:
    """169 (DECISIONS 169): a reconciliation package's inputs as its item list: the scope's statements, then the
    invoices its pairing may take (another party's only with `others`), each with the package and the run its result
    comes from, whether it blocks the approval, and where it stands in the pairing; with `left_out`, the documents left
    out of the package too, with who left them out, when and why. Read from the kept workspace."""
    from jav import reconcile_package

    ws = reconcile_package.kept_workspace(scope["workpackage_id"])
    refs = ws.get("open") or {}
    briefs = work.run_briefs(ref["run_id"] for ref in refs.values())
    blocking = {b["doc_id"] for b in ws["blockers"] if str(b["code"]).startswith("source_")}
    cols = [
        _col("name", "Document", link="review"),
        _col("kind", "Kind", "enum", labels="reconcile_input_kind"),
        _col("subject", "Number or account"),
        _col("partner", "Partner"),
        _col("date", "Date or period"),
        _col("amount", "Amount", "money"),
        _col("currency", "Currency", hidden=True),
        _col("party", "Party", hidden=scope.get("others") != "1"),
        _col("package", "Package", link="workpackage"),
        _col("source", "Counted run", "enum", labels="reconcile_input_source", link="run", badge=True),
        _col("pairing", "In the pairing", "enum", labels="reconcile_invoice_state"),
        _col("open_lines", "Open lines", "number", alert=True, hidden=True),  # a statement's only; Pairing shows its lines
        _col("blocks", "Blocks the approval", "enum", labels="yes", hidden=True),  # the Preparation's blockers card lists them
        _col("lines", "Lines in the period", "number", hidden=True),
        _col("verified", "Balances check out", "enum", labels="yes", hidden=True),
        _col("doc_id", "Document ID", "id", hidden=True),
    ]

    def common(doc_id: str, file: str | None, source: str | None) -> dict[str, Any]:
        ref = refs.get(doc_id) or {}
        brief = briefs.get(ref.get("run_id") or "") or {}
        shown_source = (("live" if brief.get("mode") == "apply" else "trial") if source == "not_approved" else source)
        return {"_key": doc_id, "_wp": ref.get("workpackage_id"), "_run": ref.get("run_id"), "_left_out": False,
                "item_id": ref.get("item_id"), "doc_id": doc_id, "name": file or doc_id[:12],
                "package": brief.get("workpackage_name"), "source": shown_source, "blocks": "igen" if doc_id in blocking else None}

    def period(row: dict[str, Any]) -> str | None:
        """A statement's period, its end without the year when it is the start's (the column stays narrow)."""
        first, last = str(row.get("period_start") or "")[:10], str(row.get("period_end") or "")[:10]
        if first[:5] and first[:5] == last[:5]:
            last = last[5:]
        return f"{first} – {last}" if first or last else None

    rows = []
    for st in sorted(ws.get("statements") or [], key=lambda s: (str(s.get("account") or ""), str(s.get("period_start") or ""))):
        rows.append({**common(st["id"], st["file"], st["source"]), "kind": "statement", "subject": st.get("account"),
                     "partner": None, "date": period(st), "amount": None, "currency": st.get("currency"),
                     "party": (ws["scope"].get("party") or {}).get("name"), "pairing": None, "open_lines": st["open_lines"],
                     "lines": st["lines"], "verified": "igen" if st["verified"] else None})
    for inv in ws["invoices"]:
        if not inv["own"] and scope.get("others") != "1":
            continue
        rows.append({**common(inv["id"], inv.get("file"), inv.get("source")), "kind": "outgoing" if inv.get("outgoing") else "incoming",
                     "subject": inv.get("number"), "partner": inv.get("partner_name"), "date": inv.get("issue_date"),
                     "amount": inv.get("amount"), "currency": inv.get("currency"), "party": (inv.get("party") or {}).get("name"),
                     "pairing": inv.get("state"), "open_lines": None, "lines": None, "verified": None})
    if scope.get("left_out") == "1":
        cols += [_col("left_out_by", "Left out by"), _col("left_out_at", "Left out at", "datetime", hidden=True),
                 _col("left_out_reason", "Reason for leaving it out")]
        rows = [{**r, "left_out_by": None, "left_out_at": None, "left_out_reason": None} for r in rows]
        for g in ws.get("left_out") or []:
            invoice = g["kind"] == "invoice"
            rows.append({**common(g["id"], g.get("file"), g.get("source")), "_left_out": True, "blocks": None,
                         "kind": ("outgoing" if g.get("outgoing") else "incoming") if invoice else "statement",
                         "subject": g.get("number") if invoice else g.get("account"), "partner": g.get("partner_name"),
                         "date": g.get("issue_date") if invoice else period(g), "amount": g.get("amount"),
                         "currency": g.get("currency"), "party": (g.get("party") or {}).get("name"), "pairing": None,
                         "open_lines": None, "lines": None, "verified": None, "left_out_by": g["actor"],
                         "left_out_at": g["left_out_at"], "left_out_reason": g["reason"]})
    return cols, rows


def shown_final(final_status: str | None, open_reasons: int) -> str | None:
    """167 (review 153 E): an item that ended with to-dos and has none open now reads "to-dos settled", not "to-dos"
    (the flow's own final status does not change when a person settles them)."""
    return "settled" if final_status == "needs_review" and not open_reasons else final_status


def _run_items(scope: dict[str, str]) -> Rows:
    from jav import naming

    choice = _name_choice(scope)
    run = work.get_run(scope["run_id"])
    titles = work_views.item_titles(run["input"]["items"], work_views.package_root(work.get(run["workpackage_id"])))
    results = {r["item_id"]: r for r in run["items"]}
    names = naming.run_item_names(run["run_id"])
    cols = [
        _col("name", "Tétel", link="review", names=choice),
        *_name_cols(),
        _col("kind", "Fajta", "enum", labels="kind", hidden=True),
        _col("status", "Futás", "enum", labels="item_status", badge=True),
        _col("final_status", "Eredmény", "enum", labels="final_status"),
        _col("open_reasons", "Teendő", "number", link="review", alert=True),
        _col("earlier_reasons", "Korábbi teendő", "number", hidden=True),
        _col("error", "Hiba", hidden=True),
        _col("updated_at", "Frissítve", "datetime", hidden=True),
        _col("item_id", "Tétel-azonosító", "id", hidden=True),
    ]
    spent = costs.run_costs(run["run_id"])
    cols += _cost_cols(spent)
    rows = []
    splits = work.items_reasons(run["run_id"], run["input"]["items"])
    for i in run["input"]["items"]:
        r = results.get(i["item_id"]) or {}
        split = splits[i["item_id"]]
        rows.append({"_key": i["item_id"], "_wp": run["workpackage_id"], "_run": run["run_id"], "item_id": i["item_id"],
                     **_name_fields(i, titles, names[i["item_id"]], choice), "kind": i.get("kind") or "document", "status": r.get("status"),
                     "final_status": shown_final(r.get("final_status"), len(split["run"])), "error": r.get("error"),
                     "updated_at": r.get("updated_at"),
                     # 150: the reasons left by trial runs are not counted (the review page hides them by default)
                     "open_reasons": len(split["run"]), "earlier_reasons": sum(not r["trial"] for r in split["earlier"]),
                     **_cost_fields(spent.items[i["item_id"]])})
    return cols, rows


PROVIDER_NAMES = {"jev": "JEV", "openai": "OpenAI", "azure_di": "Azure DI"}


def _cost_key(line: costs.Line) -> str:
    return f"cost:{line.provider}:{line.model}"


def _cost_cols(spent: costs.RunCosts) -> list[Column]:
    """082: one cost column per provider and model with a successful call in the run (a failed call has no actual model
    and no known cost: its reserved maximum is in the reserved column), the item's total, the answers taken from an
    earlier answer, and the reserved maximum of the calls without a known cost (shown only when there is any)."""
    shown = [x for x in spent.lines if x.succeeded]
    models = Counter(x.provider for x in shown)  # the header names the model only where a provider had several

    def label(x: costs.Line) -> str:
        name = PROVIDER_NAMES.get(x.provider, x.provider)
        return f"{name} · {x.model} (USD)" if models[x.provider] > 1 else f"{name} (USD)"

    per_model = [_col(_cost_key(x), label(x), "number", model=x.model) for x in shown]
    return [*per_model, _col("cost_total", "Költség összesen (USD)", "number"),
            _col("cost_reused", "Korábbi válaszból (db)", "number", hidden=True),
            _col("cost_held", "Lefoglalt, ismeretlen kimenetelű (USD)", "number", hidden=not spent.held_usd)]


def _cost_fields(item: costs.Summary) -> dict[str, Any]:
    return {**{_cost_key(x): x.usd for x in item.lines if x.succeeded}, "cost_total": item.usd, "cost_reused": item.reused,
            "cost_held": item.held_usd or None}


def _package_costs(scope: dict[str, str]) -> Rows:
    """082: the package's cost over all its runs, one row per provider and model."""
    pkg = costs.package_costs(scope["workpackage_id"])
    cols = [
        _col("provider", "AI-szolgáltató", "enum", labels="provider"),
        _col("model", "Modell"),
        _col("calls", "Fizetős hívás", "number"),
        _col("failed", "Ebből sikertelen", "number", hidden=not any(x.failed for x in pkg.lines)),
        _col("reused", "Korábbi válaszból", "number"),
        _col("usd", "Költség (USD)", "number"),
        _col("held_usd", "Lefoglalt, ismeretlen kimenetelű (USD)", "number", hidden=not pkg.held_usd),
        _col("runs", "Futás", "number"),
    ]
    used = {(x.provider, x.model): sum(1 for r in pkg.runs if any((y.provider, y.model) == (x.provider, x.model) for y in r.lines))
            for x in pkg.lines}
    rows = [{"_key": f"{x.provider}:{x.model}", "_wp": pkg.workpackage_id, "provider": x.provider, "model": x.model,
             "calls": x.calls, "failed": x.failed, "reused": x.reused, "usd": x.usd, "held_usd": x.held_usd or None,
             "runs": used[(x.provider, x.model)]} for x in pkg.lines]
    return cols, rows


# --- a result table's scope: a run, or the package's current result (172) ---------------------------------------


def result_parts(scope: dict[str, str]) -> list[tuple[dict[str, Any], frozenset[str] | None]]:
    """The runs a result table reads, each with the items it gives (None: all the run's items). A run's table reads its
    run; the package's current result reads, for each item in use, the run whose result counts now
    (`package_result.in_use`), the runs in the order they were started."""
    if "run_id" in scope:
        return [(work.get_run(scope["run_id"]), None)]
    from jav import package_result

    by_run: dict[str, set[str]] = {}
    for item, x in package_result.in_use(work.get(scope["workpackage_id"])).items():
        by_run.setdefault(x.run_id, set()).add(item)
    runs = sorted((work.get_run(r) for r in by_run), key=lambda r: (r["created_at"], r["run_id"]))
    return [(r, frozenset(by_run[r["run_id"]])) for r in runs]


def _merged(scope: dict[str, str], read: Callable[[str], list[dict[str, Any]]]) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    """The records `read(run_id)` gives for the scope (`result_parts`), and the run each item's record came from. A
    run's records keep their order; the package's follow its item order."""
    records: list[dict[str, Any]] = []
    run_of: dict[str, dict[str, Any]] = {}
    for run, items in result_parts(scope):
        for r in read(run["run_id"]):
            if items is None or r["item_id"] in items:
                records.append(r)
                run_of[r["item_id"]] = run
    if "workpackage_id" in scope:
        order = {i["item_id"]: n for n, i in enumerate(work.get(scope["workpackage_id"])["items"])}
        records.sort(key=lambda r: order.get(r["item_id"], len(order)))  # stable: an item's records keep their order
    return records, run_of


def result_records(scope: dict[str, str]) -> list[dict[str, Any]]:
    """The effective data of the scope's documents (`run_records` of a run, or of each item's counted run)."""
    return _merged(scope, run_records)[0]


def _linker(scope: dict[str, str], run_of: dict[str, dict[str, Any]]) -> Callable[[str | None], dict[str, Any]]:
    """A row's link fields by its item: its package, and the run its item's record came from (in the package's result,
    a row without an item, such as a utility-cost cell, names no run)."""
    if "run_id" in scope:
        run = next(iter(run_of.values()), None) or work.get_run(scope["run_id"])
        fixed = {"_wp": run["workpackage_id"], "_run": run["run_id"]}
        return lambda _item: fixed
    wp_id = scope["workpackage_id"]

    def link(item_id: str | None) -> dict[str, Any]:
        run = run_of.get(item_id) if item_id else None
        return {"_wp": wp_id, "_run": run["run_id"] if run else None}

    return link


def _emails(scope: dict[str, str]) -> Rows:
    """058 K5.1: the run's emails — intent (with the manual correction), suggested next step, attachments, and how much
    of the email's text the recognition saw."""
    from jav import export

    doc_types = _labels("doc_type")
    records, run_of = _merged(scope, export.email_records)
    link = _linker(scope, run_of)
    cols = [
        _col("subject", "Tárgy", link="item"),
        _col("sender", "Feladó"),
        _col("received_at", "Érkezett", "datetime"),
        _col("mailbox", "Postafiók", hidden=True),
        _col("intent", "Szándék", "enum", labels="intent"),
        _col("confidence", "Valószínűség", "number", percent=True),
        _col("corrected", "Javítva", "enum", labels="yes"),
        _col("machine_intent", "Gépi szándék", "enum", labels="intent", hidden=True),
        _col("next_flow", "Javasolt következő lépés", "enum", labels="next_flow"),
        _col("attachments", "Csatolmány", "number"),
        _col("attachment_types", "Csatolmányok típusa"),
        _col("body_status", "Levél szövege", "enum", labels="body_status"),
        _col("body_chars", "Levél hossza (karakter)", "number", hidden=True),
        _col("tasks", "Feladatjavaslat", "number", hidden=not any(r.get("tasks") for r in records)),
        _col("open_reasons", "Nyitott teendő", "number", link="item", alert=True),
        _col("message_id", "Levél-azonosító", "id", hidden=True),
        _col("item_id", "Tétel-azonosító", "id", hidden=True),
    ]
    rows = []
    for r in records:
        # 126: the type each attachment was processed as (`shown_type`), else the recognised one
        shown = [a.get("shown_type") or a.get("doc_type") for a in r["attachments"]]
        types = sorted({doc_types.get(x, x) for x in shown if x})
        rows.append({"_key": r["item_id"], **link(r["item_id"]), "item_id": r["item_id"],
                     "message_id": r["message_id"], "subject": r["subject"] or "(tárgy nélkül)", "sender": r["sender"],
                     "received_at": r["received_at"], "mailbox": r["mailbox"], "intent": r["intent"],
                     "machine_intent": r["machine_intent"], "confidence": r["confidence"],
                     "corrected": "igen" if r["corrected"] else None, "next_flow": r["next_flow"],
                     "attachments": len(r["attachments"]), "attachment_types": ", ".join(types) or None,
                     "body_status": r["body"]["status"], "body_chars": r["body"]["chars"], "open_reasons": len(r["open_reasons"]),
                     "tasks": len((r.get("tasks") or {}).get("tasks") or []) or None})
    return cols, rows


def _email_tasks(scope: dict[str, str]) -> Rows:
    """058 K5.3: the task proposals, one per row, with the evidence and the human decision (only a person accepts)."""
    from jav import export

    records, run_of = _merged(scope, export.email_records)
    link = _linker(scope, run_of)
    cols = [
        _col("subject", "Levél", link="item"),
        _col("action", "Akció", "enum", labels="task_action"),
        _col("title", "Feladat"),
        _col("due_date", "Határidő", "date"),
        _col("assignee_hint", "Felelős"),
        _col("evidence", "Bizonyíték"),
        _col("decision", "Döntés", "enum", labels="task_decision"),
        _col("decided_by", "Döntött", hidden=True),
        _col("decided_at", "Döntés ideje", "datetime", hidden=True),
        _col("done_at", "Elvégezve", "datetime"),  # 062: manual "done" on an accepted task
        _col("done_by", "Elvégezte", hidden=True),
        _col("item_id", "Tétel-azonosító", "id", hidden=True),
    ]
    rows = [{**t, "_key": f"{t['item_id']}:{t['index']}", **link(t["item_id"])} for t in export.task_rows(records)]
    return cols, rows


# the export tables (jav/export.py) by column key, in header order
_DOC_KEYS = [("file", "Irat", "text"), ("item_id", "Tétel-azonosító", "id"), ("doc_type", "Típus", "enum"), ("arm", "Út", "enum"),
             ("final_status", "Állapot", "enum"), ("open_reasons", "Nyitott teendők", "text"),
             ("source_email", "Forrás-levél", "text")]  # 058 K5.2: for an email attachment, the email's subject
# 180 (plan 170 E7): a domain's columns follow these (`jav/outputs/parts.py`; the duplicate mark and the direction)
_DP_KEYS = _DOC_KEYS[:3] + [("field", "Mező", "enum"), ("value", "Érték", "text"), ("page", "Oldal", "number"),
                            ("quote", "Forrásszöveg", "text"), ("place", "Hely", "enum"), ("corrected", "Javítva", "enum"),
                            ("field_reason", "Teendő a mezőn", "enum")]
_LI_KEYS = _DOC_KEYS[:3] + [("list", "Lista", "enum"), ("row", "Sor", "number"), ("page", "Oldal", "number")]
_HIDDEN = {"item_id"}
# 167 (review 153 E): the export's "igen" marks (corrected, a to-do on the field) are enumerated, so English shows "yes"
_LABELS = {"final_status": "final_status", "place": "provenance", "field": "field", "list": "field", "doc_type": "doc_type",
           "corrected": "yes", "field_reason": "yes"}


def _fixed(keys: list[tuple[str, str, str]]) -> list[Column]:
    """The fixed columns of the export; the document's name links to the item's review view."""
    return [_col(k, label, kind, hidden=k in _HIDDEN, labels=_LABELS.get(k), **({"link": "review"} if k == "file" else {}))
            for k, label, kind in keys]


def _zip_rows(keys: list[str], rows: list[list[Any]], link: Callable[[str | None], dict[str, Any]],
              key_of: Callable[[dict[str, Any]], str]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        d = dict(zip(keys, row, strict=False))
        to = link(d.get("item_id"))
        d.update({"_run": to["_run"], "_wp": to["_wp"], "_key": key_of(d)})
        out.append(d)
    return out


def _run_wp(to: dict[str, Any]) -> dict[str, Any]:
    """The link fields in the order the record tables have always written them."""
    return {"_run": to["_run"], "_wp": to["_wp"]}


def _records(scope: dict[str, str]) -> tuple[Callable[[str | None], dict[str, Any]], list[dict[str, Any]]]:
    """The scope's documents' effective data, and a row's link fields by its item (`_linker`)."""
    records, run_of = _merged(scope, run_records)
    return _linker(scope, run_of), records


def run_records(run_id: str) -> list[dict[str, Any]]:
    """The run's effective data (`export.run_records`) from the cache: all result tables, the utility report and the
    full export use the same — expensive — computation, and recompute only after the run's data changes."""
    from jav import export

    scope = {"run_id": run_id}
    key = (str(store.active_path()), "_records", run_id)
    fp = _run_fingerprint(scope)
    with _lock:
        hit = _cache.get(key)
        if hit and hit[0] == fp:
            _cache.move_to_end(key)
            return hit[1]  # type: ignore[return-value]
    built = export.run_records(run_id)
    _remember(key, fp, built)
    return built


def _field_kind(kind: str | None) -> str:
    return {"money": "money", "number": "number", "date": "date"}.get(kind or "", "text")


def _documents(scope: dict[str, str]) -> Rows:
    from jav import export

    link, records = _records(scope)
    head, rows = export.documents_table(records)
    kinds: dict[str, str] = {}
    for r in records:
        for f, k in r["kinds"].items():
            kinds.setdefault(f, k)
    fields = head[len(export.document_head()):]
    cols = _fixed(_DOC_KEYS)
    for part in parts.document_columns():  # 180: a domain's columns (126: the duplicate mark; 147: the direction)
        first = len(cols)
        # shown only when a row has a value in the part's first column (a document with a mark, an invoice)
        shown = any(row[first] for row in rows)
        cols += [_col(c.key, c.head, c.kind, hidden=not shown, labels=c.labels) for c in part.columns]
    cols += [_col(f"f:{f}", _field_label(f), _field_kind(kinds.get(f)), field=f) for f in fields]
    if not any(r.get("source_email") for r in records):  # only meaningful in an email package: hidden otherwise
        cols = [replace(c, hidden=True) if c.key == "source_email" else c for c in cols]
    return cols, _zip_rows([c.key for c in cols], rows, link, lambda d: d["item_id"])


def _datapoints(scope: dict[str, str]) -> Rows:
    from jav import export

    link, records = _records(scope)
    _head, rows = export.datapoints_table(records)
    cols = _fixed(_DP_KEYS)
    return cols, _zip_rows([c.key for c in cols], rows, link, lambda d: f"{d['item_id']}:{d['field']}")


def _native_facts(scope: dict[str, str]) -> Rows:
    from jav import export

    native, run_of = _merged(scope, export.native_records)
    _head, records = export.native_facts_table(native)
    # 120: the readable source column is shown; the citation JSON and the constant correctness stay technical, and the
    # machine states get their everyday labels
    hidden = {"item_id", "fact_id", "result_version", "proposal", "reading_gaps", "source_citations", "correctness"}
    states = {"machine_state", "machine_grounding", "reading_status", "interpretation_status", "correctness"}
    # 183: a document's long file name repeats in each of its facts' rows; it, the source and the value keep a row short
    lines = {"file": 2, "source": 3, "effective_value": 4}
    cols = [_col(key, label, "enum" if key in states else "text", labels="native_state" if key in states else None,
                 hidden=key in hidden, **({"clamp": lines[key]} if key in lines else {}))
            for key, label in zip(export.NATIVE_COLUMNS, export.NATIVE_HEAD, strict=True)]
    return cols, _zip_rows(list(export.NATIVE_COLUMNS), records, _linker(scope, run_of), lambda row: f"{row['item_id']}:{row['fact_id']}")


def _line_items(scope: dict[str, str]) -> Rows:
    from jav import export

    link, records = _records(scope)
    head, rows = export.line_items_table(records)
    extra = head[len(_LI_KEYS):]
    list_kinds: dict[str, str] = {}
    from jav import typepack

    for r in records:
        for lf in typepack.get(r["doc_type"]).list_fields.values():
            for c, k in lf.items():
                list_kinds.setdefault(c, k)
    cols = _fixed(_LI_KEYS) + [_col(f"c:{c}", _column_label(c), _field_kind(list_kinds.get(c)), field=c) for c in extra]
    return cols, _zip_rows([c.key for c in cols], rows, link, lambda d: f"{d['item_id']}:{d['list']}:{d['row']}")


def _file_names(scope: dict[str, str]) -> Rows:
    """078: the content-based name of each document of the run, whether its copy is ready or goes to
    the review folder, and why (`jav/naming.py`)."""
    from jav import naming

    run = work.get_run(scope["run_id"])
    cols = [
        _col("filename", "Új fájlnév", link="review"),
        _col("status", "Mappa", "enum", labels="naming_status", badge=True),
        _col("why", "Miért ellenőrzendő"),
        _col("original", "Eredeti fájlnév"),
        _col("doc_type", "Típus", "enum", labels="doc_type"),
        _col("item_id", "Tétel-azonosító", "id", hidden=True),
    ]
    rows = [{"_key": c.item_id, "_wp": run["workpackage_id"], "_run": run["run_id"], "item_id": c.item_id,
             "filename": c.filename, "status": c.status, "why": naming.reason_text(c.reasons) or None, "original": c.original,
             "doc_type": c.doc_type} for c in naming.plan(run["run_id"])]
    return cols, rows


def _calls(scope: dict[str, str]) -> Rows:
    work.get_run(scope["run_id"])
    cols = [
        _col("created_at", "Időpont", "datetime"),
        _col("step", "Lépés"),
        _col("provider", "Szolgáltató", "enum"),
        _col("model_actual", "Modell", "enum"),
        _col("status", "Állapot", "enum"),
        _col("cost_usd", "Költség (USD)", "number"),
        _col("max_cost_usd", "Lefoglalt (USD)", "number", hidden=True),
        _col("input_tokens", "Bemeneti token", "number", hidden=True),
        _col("error", "Hiba", hidden=True),
        _col("step_id", "Lépés (teljes)", "id", hidden=True),
    ]
    rows = [{**{k: c.get(k) for k in ("created_at", "provider", "model_actual", "status", "cost_usd", "max_cost_usd",
                                      "input_tokens", "error", "step_id")},
             "_key": str(c["id"]), "_run": scope["run_id"], "step": ":".join(str(c["step_id"]).split(":")[:2])}
            for c in calls.scope_journal(scope["run_id"])]
    return cols, rows


def _utility_report(scope: dict[str, str]) -> tuple[Callable[[str | None], dict[str, Any]], dict[str, Any]]:
    from jav import report_utility

    link, records = _records(scope)
    return link, report_utility.build(records)


def _utility_cost(scope: dict[str, str]) -> Rows:
    """The monthly utility-cost grid in long form: consumption point × utility × month per row (pivotable in Excel)."""
    link, rep = _utility_report(scope)
    to = link(None)
    cols = [
        _col("address", "Fogyasztási hely"),
        _col("utility", "Közmű", "enum"),
        _col("month", "Hónap", "date"),
        _col("status", "Állapot", "enum", labels="utility_status"),
        _col("amount", "Összeg (bruttó, a hónapra jutó)", "money"),
        _col("consumption", "Fogyasztás", "number"),
        _col("unit", "Egység", "enum", hidden=True),
        _col("settlement", "Elszámoló számla", "enum", labels="yes", hidden=True),
        _col("summary_only", "Csak tájékoztató", "enum", labels="yes"),
        _col("bills", "Forrásszámla", "number"),
        _col("suppliers", "Szolgáltató", hidden=True),
    ]
    rows = []
    for s_ in rep["series"]:
        for m, c in s_["cells"].items():
            rows.append({"_key": f"{s_['address']}|{s_['utility']}|{m}", "_run": to["_run"], "_wp": to["_wp"],
                         "address": s_["address"], "utility": s_["utility"], "month": m, "status": c["status"], "amount": c["amount"],
                         "consumption": c["consumption"], "unit": s_["consumption_unit"], "settlement": "igen" if c["settlement"] else None,
                         "summary_only": "igen" if s_["summary_only"] else None, "bills": len(c["sources"]),
                         "suppliers": ", ".join(s_["suppliers"])})
    return cols, rows


def _utility_sources(scope: dict[str, str]) -> Rows:
    """Source invoices of each utility-cost cell (the report's "Közmű-források" sheet): cell → document, page, sum."""
    link, rep = _utility_report(scope)
    cols = [
        _col("address", "Fogyasztási hely"),
        _col("utility", "Közmű", "enum"),
        _col("month", "Hónap", "date"),
        _col("status", "Állapot", "enum", labels="utility_status"),
        _col("amount", "Összeg (a hónapra jutó)", "money"),
        _col("file", "Irat", link="review"),
        _col("page", "Oldal", "number"),
        _col("field", "Mező", "enum", labels="field"),
        _col("corrected", "Javítva", "enum", labels="yes"),
        _col("open_reasons", "Nyitott teendő", "number", alert=True),
        _col("item_id", "Tétel-azonosító", "id", hidden=True),
    ]
    rows = []
    for s_ in rep["series"]:
        for m, c in s_["cells"].items():
            for i, x in enumerate(c["sources"]):
                rows.append({"_key": f"{s_['address']}|{s_['utility']}|{m}|{i}", **_run_wp(link(x["item_id"])),
                             "address": s_["address"], "utility": s_["utility"], "month": m,
                             "status": "settlement" if x.get("settlement") else c["status"], "amount": x["amount"], "file": x["file"],
                             "page": x["page"], "field": x["field"], "corrected": "igen" if x["corrected"] else None,
                             "open_reasons": x["open_reasons"], "item_id": x["item_id"]})
    for u in rep["unplaced"]:
        rows.append({"_key": f"unplaced|{u['item_id']}", **_run_wp(link(u["item_id"])), "status": u["reason"],
                     "file": u["file"], "item_id": u["item_id"]})
    for d in rep.get("duplicates", []):
        rows.append({"_key": f"duplicate|{d['item_id']}", **_run_wp(link(d["item_id"])), "status": "duplicate",
                     "file": d["file"], "item_id": d["item_id"]})
    return cols, rows


def _mailbox_pulls(_scope: dict[str, str]) -> Rows:
    from jav import mailbox

    cols = [
        _col("created_at", "Indítva", "datetime"),
        _col("label", "Postafiók / időszak"),
        _col("status", "Állapot", "enum", labels="pull_status", badge=True),
        _col("new", "Új levél", "number"),
        _col("changed", "Változott", "number", hidden=True),
        _col("duplicate", "Már megvolt", "number", hidden=True),
        _col("workpackage", "Munkacsomag", "id", link="workpackage"),
        _col("actor", "Indította"),
        _col("error", "Hiba"),
        _col("finished_at", "Befejeződött", "datetime", hidden=True),
    ]
    rows = []
    for p in mailbox.pulls(limit=cfg.load("datasets")["export_max_rows"]):
        res = p.get("result") or {}
        rows.append({"_key": p["id"], "_wp": res.get("workpackage"), "created_at": p["created_at"], "label": p["label"],
                     "status": p["status"], "new": res.get("new"), "changed": res.get("changed"), "duplicate": res.get("duplicate"),
                     "workpackage": res.get("workpackage"), "actor": p["actor"], "error": res.get("error"),
                     "finished_at": p.get("finished_at")})
    return cols, rows


RESULT_SCOPE = ("run_id", "workpackage_id")

REGISTRY: dict[str, Dataset] = {d.name: d for d in [
    Dataset("runs", "Futások", (), _runs, optional_scope=("workpackage_id",)),
    Dataset("workpackages", "Munkacsomagok", (), _workpackages, _workpackages_fingerprint, optional_scope=("include_archived", "owner"),
            natural_sort=(("created_at", True),), refresh_behind=True),
    Dataset("workpackage_items", "Csomag tételei", ("workpackage_id",), _workpackage_items, optional_scope=("names", "removed")),
    Dataset("reconcile_inputs", "A reconciliation's inputs", ("workpackage_id",), _reconcile_inputs,
            optional_scope=("others", "left_out")),
    Dataset("run_items", "Futás tételei", ("run_id",), _run_items, _run_items_fingerprint, optional_scope=("names",)),
    Dataset("package_costs", "A csomag költsége", ("workpackage_id",), _package_costs),
    # 172: a result table is a run's, or the package's current result (each item from its counted run)
    Dataset("emails", "Levelek", (), _emails, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("email_tasks", "Feladatjavaslatok", (), _email_tasks, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("documents", "Iratok", (), _documents, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("datapoints", "Adatpontok", (), _datapoints, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("native_facts", "Native facts", (), _native_facts, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("line_items", "Tételsorok", (), _line_items, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("utility_cost", "Közmű-költség havonta", (), _utility_cost, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("utility_sources", "Közmű-költség forrásai", (), _utility_sources, _result_fingerprint, one_of=RESULT_SCOPE),
    Dataset("file_names", "Fájlnevek", ("run_id",), _file_names, _run_fingerprint),
    Dataset("calls", "Hívásnapló", ("run_id",), _calls, _calls_fingerprint),
    Dataset("mailbox_pulls", "Postafiók-letöltések", (), _mailbox_pulls),
    Dataset("activity", "Tevékenységnapló", ("actor",), _activity, optional_scope=("day",)),
]}


# --- cache and query -------------------------------------------------------------------------------------------

_cache: OrderedDict[tuple, tuple[str, Any]] = OrderedDict()
_lock = threading.Lock()
_rebuilding: dict[tuple, threading.Thread] = {}  # 153: the background rebuild of a `refresh_behind` dataset, per key
BUILD_WAIT_SECONDS = 60.0  # 158: how long a request with no rows waits for their running build before building them itself
_log = logging.getLogger("jav.datasets")


def get(name: str) -> Dataset:
    if name not in REGISTRY:
        raise UnknownDataset(name)
    return REGISTRY[name]


def _scope_of(ds: Dataset, scope: dict[str, str] | None) -> dict[str, str]:
    scope = {k: v for k, v in (scope or {}).items() if v}
    missing = [k for k in ds.scope if k not in scope]
    if missing:
        raise MissingScope(f"dataset {ds.name} needs scope: {', '.join(missing)}")
    if ds.one_of and sum(k in scope for k in ds.one_of) != 1:
        raise MissingScope(f"dataset {ds.name} needs exactly one scope of: {', '.join(ds.one_of)}")
    return {k: scope[k] for k in (*ds.scope, *ds.one_of, *ds.optional_scope) if k in scope}


def rows(name: str, scope: dict[str, str] | None = None) -> Rows:
    """The dataset's columns and all rows (for the scope), from the cache if the fingerprint has not changed. Never
    outdated (a download, for example); `query` may serve a `refresh_behind` dataset's last rows."""
    return _rows(name, scope, stale_ok=False)[0]


def _key(name: str, sc: dict[str, str]) -> tuple:
    return (str(store.active_path()), name, tuple(sorted(sc.items())))


def _rows(name: str, scope: dict[str, str] | None, *, stale_ok: bool) -> tuple[Rows, bool]:
    """The rows, and whether they are the last ones while a newer version is being built (153, `refresh_behind`)."""
    ds = get(name)
    sc = _scope_of(ds, scope)
    key = _key(name, sc)
    if ds.fingerprint is None:
        with store.session():  # 061: one connection for all queries of the rows
            return ds.build(sc), False
    fp = ds.fingerprint(sc)
    with _lock:
        running = None if key in _cache else _rebuilding.get(key)
    if running is not None:
        running.join(BUILD_WAIT_SECONDS)  # 158: no rows yet, but they are being built (the warm-up): wait for them
    with _lock:
        hit = _cache.get(key)
        if hit and hit[0] == fp:
            _cache.move_to_end(key)
            return hit[1], False
        if hit and stale_ok and ds.refresh_behind:
            _rebuild_behind(key, ds, sc)
            return hit[1], True
    with store.session():
        built = ds.build(sc)
    _remember(key, fp, built)
    return built, False


def _rebuild_behind(key: tuple, ds: Dataset, sc: dict[str, str]) -> None:
    """Starts the background rebuild of `key` unless one is running (called under `_lock`). The thread works on the
    request's store; its fingerprint is taken before the build, so a change during the build shows at the next request."""
    if key in _rebuilding:
        return
    path = store.active_path()

    def rebuild() -> None:
        try:
            with store.use_store(path):
                fp = ds.fingerprint(sc)  # type: ignore[misc] - only datasets with a fingerprint are kept
                with store.session():
                    built = ds.build(sc)
                _remember(key, fp, built)
        except Exception:  # noqa: BLE001 - the last rows stay; the next request starts another rebuild
            _log.exception("background rebuild of dataset %s failed", ds.name)
        finally:
            with _lock:
                _rebuilding.pop(key, None)

    thread = threading.Thread(target=rebuild, name=f"dataset-{ds.name}", daemon=True)
    _rebuilding[key] = thread
    thread.start()


def wait_for_rebuilds(timeout: float = 30.0) -> None:
    """Waits for the running background rebuilds (tests, and a clean stop)."""
    with _lock:
        threads = list(_rebuilding.values())
    for thread in threads:
        thread.join(timeout)


def _remember(key: tuple, fp: str, value: Any) -> None:
    with _lock:
        _cache[key] = (fp, value)
        _cache.move_to_end(key)
        while len(_cache) > cfg.load("datasets")["cache_entries"]:
            # 158: the least recently used entry goes, but a run's table before a refresh-behind dataset
            victim = next((k for k in _cache if not _refreshes_behind(k)), None) or next(iter(_cache))
            del _cache[victim]


def _refreshes_behind(key: tuple) -> bool:
    ds = REGISTRY.get(key[1])
    return bool(ds and ds.refresh_behind)


def _warm_names() -> list[str]:
    """The datasets the service builds when it starts: those served outdated, in their default scope."""
    return [ds.name for ds in REGISTRY.values() if ds.refresh_behind and ds.fingerprint is not None]


def warm_up() -> None:
    """158: starts building the rows of `_warm_names` in the background on the active store, unless they are cached
    or being built; a request that arrives meanwhile waits for that build (`_rows`)."""
    with _lock:
        for name in _warm_names():
            ds = REGISTRY[name]
            sc = _scope_of(ds, None)
            key = _key(name, sc)
            if key not in _cache:
                _rebuild_behind(key, ds, sc)


def query(name: str, scope: dict[str, str] | None, q: Query) -> dict[str, Any]:
    """086 (audit of 2026-10-02, N01): a run's table carries the version of the reviewed result it shows
    (`review_version`), so the approval names exactly what the approver saw. The version is read before the rows: the
    rows are never older than it (the cache is keyed by a fingerprint of the same corrections and decisions), and rows
    newer than it make the approval fail with a conflict, the safe side. 172: a table of the package's current result
    carries the result's mark (`result_mark`) the same way, for Approve and release."""
    version = mark = None
    ds = get(name)
    if scope and scope.get("run_id"):
        from jav import corrections

        version = corrections.review_version(scope["run_id"])
    elif ds.one_of and scope and scope.get("workpackage_id"):
        from jav import package_release

        mark = package_release.result_mark(scope["workpackage_id"])  # 172: what Approve and release names
    (cols, all_rows), refreshing = _rows(name, scope, stale_ok=True)
    page = {"dataset": ds.spec(), **run_query(cols, all_rows, q)}
    if version is not None:
        page["review_version"] = version
    if mark is not None:
        page["result_mark"] = mark
    if refreshing:
        page["refreshing"] = True  # 153: these are the last rows; newer ones are being built
    return work_views.jsonable(page)


# --- download ---------------------------------------------------------------------------------------------------

MEDIA = {"csv": "text/csv; charset=utf-8", "json": "application/json",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def export_rows(name: str, scope: dict[str, str] | None, q: Query, rows_mode: str,
                columns: list[str] | None = None) -> tuple[list[Column], list[dict[str, Any]]]:
    """The columns and rows to download. `rows_mode`: `all` (every row, the sort order is kept), `filtered` (by the
    search and the filters), `selected` (only the selected `keys`). Paging does not apply to downloads."""
    cols, all_rows = rows(name, scope)
    if rows_mode == "all":
        q = Query(sort=q.sort)
    elif rows_mode == "selected":
        if not q.keys:
            raise ValueError("selected export needs keys")
        q = Query(sort=q.sort, keys=q.keys)
    elif rows_mode != "filtered":
        raise ValueError(f"unknown rows mode: {rows_mode}")
    by_key = {c.key: c for c in cols}
    if columns:
        unknown = [c for c in columns if c not in by_key]
        if unknown:
            raise ValueError(f"unknown column: {', '.join(unknown)}")
        out_cols = [by_key[c] for c in columns]
    else:
        out_cols = [c for c in cols if not c.hidden]
    picked = select(cols, all_rows, q)
    limit = cfg.load("datasets")["export_max_rows"]
    if len(picked) > limit:
        raise ValueError(f"export too large: {len(picked)} rows (limit {limit})")
    return out_cols, picked


def _stamp() -> str:
    from datetime import datetime

    return datetime.now().strftime("%Y%m%d-%H%M")


def export_file(name: str, scope: dict[str, str] | None, q: Query, rows_mode: str, fmt: str,
                columns: list[str] | None = None) -> tuple[bytes, str, str, int]:
    """(content, media type, file name, row count). CSV and Excel write the displayed label (instead of the enumerated
    code), JSON the raw value with the column description. Formula protection and the Hungarian CSV format belong to
    `jav/export.py`."""
    import io
    import json

    from jav import export

    if fmt not in MEDIA:
        raise ValueError(f"unknown export format: {fmt}")
    ds = get(name)
    cols, picked = export_rows(name, scope, q, rows_mode, columns)
    sc = _scope_of(ds, scope)
    fname = "-".join([name, *[v[:24] for v in sc.values()], rows_mode, _stamp()]) + f".{fmt}"
    head = [c.label for c in cols]
    table = [[shown(c, r.get(c.key)) for c in cols] for r in picked]
    if fmt == "csv":
        data = export.csv_bytes(head, table)
    elif fmt == "xlsx":
        from openpyxl import Workbook

        from jav.outputs import xlsx as xlsx_writer

        wb = Workbook()
        ws = wb.active
        ws.title = xlsx_writer.safe_sheet(ds.label, set())
        xlsx_writer.write_table(ws, head, table, [c.kind if c.kind in ("money", "number") else None for c in cols])
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
    else:
        body = {"dataset": ds.spec(), "scope": sc, "rows_mode": rows_mode, "query": q.model_dump(exclude={"offset", "limit"}),
                "exported_at": _stamp(), "columns": [c.spec() for c in cols],
                "rows": [{c.key: r.get(c.key) for c in cols} for r in picked]}
        data = json.dumps(work_views.jsonable(body), ensure_ascii=False, indent=1).encode("utf-8")
    return data, MEDIA[fmt], fname, len(picked)


def catalog() -> list[dict[str, Any]]:
    return [d.spec() for d in REGISTRY.values()]


def clear_cache() -> None:
    wait_for_rebuilds()  # 153: a rebuild finishing later would put its rows back
    with _lock:
        _cache.clear()


__all__ = ["REGISTRY", "RESULT_SCOPE", "Dataset", "MissingScope", "UnknownDataset", "catalog", "clear_cache", "export_file",
           "export_rows", "get", "query", "result_parts", "result_records", "rows", "run_records", "wait_for_rebuilds",
           "warm_up"]
