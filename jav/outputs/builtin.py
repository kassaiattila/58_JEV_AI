"""The built-in layout of an output (178, plan 170 L4): the workbook every download had until 178.

196: a Word output (`format: docx`) has a built-in Word template instead, which its domain registers
(`register_word`) and which is filled like an own template (`jav/outputs/docx.py`).

Each table of the output's data (`jav/outputs/contract.py`) on its own sheet, in order, named by the table's title and
written by the shared writer (`jav/outputs/xlsx.py`); with `heading`, the output's values stand as a block above every
table; a marked cell gets its mark's fill. The layout's settings are data (`configs/outputs.json`, `builtin`).

Layer: adapter (file format), called by the builders' modules.
"""

from __future__ import annotations

import io
from collections.abc import Callable

from jav.outputs import contract, xlsx


# 196: the built-in Word template of a Word output, made for its data (its language), by the domain that owns it
_WORD: dict[str, Callable[[contract.OutputData], bytes]] = {}


def register_word(output: str, template: Callable[[contract.OutputData], bytes]) -> None:
    """196: registers the built-in Word template of a Word output (`format: docx`); it is filled like an own one."""
    _WORD[output] = template


def render(data: contract.OutputData) -> bytes:
    """The output's file in the built-in layout: a Word output's built-in template filled (196), a workbook otherwise.
    `contract.ContractError` when the data disagrees with the contract or a Word output has no built-in template."""
    if contract.spec(data.output)["format"] == "docx":
        from jav import registrations
        from jav.outputs import docx

        registrations.load_all()  # the domains register their built-in Word templates
        if data.output not in _WORD:
            raise contract.ContractError(f"{data.output}: no built-in Word template")
        return docx.fill(data, _WORD[data.output](data))
    return _workbook(data)


def _workbook(data: contract.OutputData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    contract.check(data)
    layout = contract.spec(data.output)["builtin"]
    heading = tuple((v.label, v.value) for v in data.values) if layout["heading"] else ()
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for t in data.tables:
        ws = wb.create_sheet(xlsx.safe_sheet(t.title, used))
        top = xlsx.write_table(ws, t.head, [*t.rows, *t.footer], t.kinds, heading=heading, min_width=layout["min_width"])
        for (r, c), mark in t.marks.items():
            ws.cell(row=top + 1 + r, column=c + 1).fill = PatternFill("solid", fgColor=layout["marks"][mark])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
