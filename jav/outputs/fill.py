"""An own Excel template: its check at upload and its filling with an output's data (178, plan 170 L4, DECISIONS 178).

Plain-language summary: the person's workbook keeps its look - fonts, colours, a logo, column widths, merged cells -
and the system writes the data into its markers (`jav/outputs/markers.py`). A row holding a table's markers is
repeated once per row of the table with that row's formatting (none: the row goes), the rows below move down with
their merged cells and pictures; a value, a total or a count is written where its marker stands; the months spread
to the right. A text value is always text, never a formula, and what a value contains is never read as a marker.

The check refuses a template that would not fill safely or correctly: not an Excel workbook (.xlsx), too large, macros
or embedded programs, a link to another file or a data connection, a formula (the owner's decision of 2026-10-19), a
marker the output does not have, two tables' markers in one row, a table's markers in two rows of a sheet, a value
marker in a table's row, the months not the last cell of their row, a merged cell in a table's row, or no marker at
all. Conditional formatting and data validation are kept but do not spread to the repeated rows: a warning.

Layer: adapter (file format), called by `jav/outputs/templates.py`.
"""

from __future__ import annotations

import io
import re
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from typing import Any

from jav.outputs import contract, markers, xlsx
from jav.core.values import stored_decimal

ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DATE_FORMAT = "yyyy-mm-dd"


class TemplateError(ValueError):
    """A template that cannot be filled (it changed since its check, or the data disagrees with the contract)."""


@dataclass
class TemplateCheck:
    """The check of a template for an output: `problems` (each refuses it) and `warnings`, each `{code, where,
    detail}`; the values and the tables (name as written -> columns) it uses."""

    output: str
    problems: list[dict[str, str]] = field(default_factory=list)
    warnings: list[dict[str, str]] = field(default_factory=list)
    values: list[str] = field(default_factory=list)
    tables: dict[str, list[str]] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.problems

    def add(self, code: str, where: str = "", detail: str = "", *, warning: bool = False) -> None:
        (self.warnings if warning else self.problems).append({"code": code, "where": where, "detail": detail})

    def as_dict(self) -> dict[str, Any]:
        return {"output": self.output, "ok": self.ok, "problems": self.problems, "warnings": self.warnings,
                "values": self.values, "tables": self.tables}


# --- the check ------------------------------------------------------------------------------------------------------


def _package(data: bytes, check: TemplateCheck) -> bool:
    """The file as a package: an Excel workbook, without active content, links to other files or data connections."""
    from jav.documents.readers.limits import DEFAULT_LIMITS, ReadFailure, inspect_package

    if len(data) > contract.templates()["max_bytes"]:
        check.add("too_large", detail=str(len(data)))
        return False
    if not zipfile.is_zipfile(io.BytesIO(data)):
        check.add("not_xlsx")
        return False
    try:
        parts = inspect_package(data, DEFAULT_LIMITS)
    except ReadFailure as e:
        check.add("unsafe_content" if e.status == "excluded" else "too_large" if e.status == "resource_limited" else "unreadable",
                  detail=str(e))
        return False
    names = [n.lower() for n in parts]
    if "xl/workbook.xml" not in names:
        check.add("not_xlsx")
        return False
    if b"macroenabled" in parts.get("[Content_Types].xml", b"").lower():
        check.add("unsafe_content", detail="macro-enabled workbook")
    if any(n.startswith("xl/externallinks/") for n in names):
        check.add("external_link")
    if any(n == "xl/connections.xml" or n.startswith("xl/querytables/") for n in names):
        check.add("data_connection")
    return check.ok


def _load(data: bytes) -> Any:
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data), keep_links=False)


def _scan(ws: Any, output: str, check: TemplateCheck | None) -> tuple[dict[int, str], list[tuple[Any, str, list[markers.Marker]]]]:
    """The sheet's table rows (row -> table name) and its marker cells (cell, text, markers). With `check`, the
    problems are recorded there; without, a problem raises TemplateError."""
    rows: dict[int, str] = {}
    cells: list[tuple[Any, str, list[markers.Marker]]] = []
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str) or "{{" not in c.value:
                continue
            found, problems = markers.find(output, c.value)
            where = f"{ws.title}!{c.coordinate}"
            for p in problems:
                if check is None:
                    raise TemplateError(f"{where}: {p}")
                check.add("unknown_marker", where, p)
            cells.append((c, c.value, found))
            for m in found:
                if not m.in_row:
                    continue
                if rows.get(c.row, m.name) != m.name:
                    if check is None:
                        raise TemplateError(f"{where}: two tables in one row")
                    check.add("row_mixes_tables", where, f"{rows[c.row]}, {m.name}")
                rows.setdefault(c.row, m.name)
    return rows, cells


def _check_sheet(ws: Any, output: str, check: TemplateCheck) -> None:
    if ws.conditional_formatting:
        check.add("conditional_formatting", ws.title, warning=True)
    if ws.data_validations.dataValidation:
        check.add("data_validation", ws.title, warning=True)
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f" or (c.value is not None and type(c.value).__name__ in ("ArrayFormula", "DataTableFormula")):
                check.add("formula", f"{ws.title}!{c.coordinate}")
    rows, cells = _scan(ws, output, check)
    seen: dict[str, int] = {}
    for r, name in sorted(rows.items()):
        if name in seen:
            check.add("table_in_two_rows", ws.title, f"{name}: {seen[name]}, {r}")
        seen.setdefault(name, r)
    for rng in ws.merged_cells.ranges:
        if any(rng.min_row <= r <= rng.max_row for r in rows):
            check.add("merged_row", f"{ws.title}!{rng}")
    for c, _text, found in cells:
        where = f"{ws.title}!{c.coordinate}"
        if c.row in rows and any(not m.in_row for m in found):
            check.add("value_in_row", where, ", ".join(m.text for m in found if not m.in_row))
        spreading = [m for m in found if m.kind in ("months", "month_names")]
        if spreading:
            right = [x for x in ws[c.row] if x.column > c.column and x.value not in (None, "")]
            if len(found) > 1 or c.value.strip() != spreading[0].text or right:
                check.add("months_not_last", where)
        for m in found:
            if m.kind == "value" and m.name not in check.values:
                check.values.append(m.name)
            elif m.kind != "value":
                cols = check.tables.setdefault(m.name, [])
                if m.column and m.column not in cols:
                    cols.append(m.column)


def check(output: str, data: bytes) -> TemplateCheck:
    """The check of a workbook as a template of the output (`TemplateCheck.ok`: it can be stored)."""
    result = TemplateCheck(output)
    contract.spec(output)
    if not _package(data, result):
        return result
    try:
        wb = _load(data)
    except Exception as e:  # noqa: BLE001 - any failure of the reading library is the file's: recorded, never swallowed
        result.add("unreadable", detail=type(e).__name__)
        return result
    for ws in wb.worksheets:
        _check_sheet(ws, output, result)
    if result.ok and not result.values and not result.tables:
        result.add("no_marker")
    return result


# --- the filling ----------------------------------------------------------------------------------------------------


def _typed(value: Any, kind: str) -> Any:
    """A value as its cell holds it, by its kind: a decimal number for money and number, a whole number for a count,
    a date for an ISO date, text otherwise (a number stays a number)."""
    if value is None:
        return None
    if kind in ("money", "number"):
        return xlsx.cell_value(value, kind)
    if kind == "count":
        return value if isinstance(value, int) and not isinstance(value, bool) else xlsx.cell_value(value, "number")
    if kind == "probability":
        number = stored_decimal(value)
        return float(number) if number is not None else xlsx.cell_value(value, None)
    if kind == "date" and isinstance(value, str) and ISO_DATE.match(value):
        return date.fromisoformat(value)
    return xlsx.cell_value(value, None)


def _shown(value: Any, kind: str) -> str:
    """A value inside a cell's text: money with a space between thousands and a decimal comma (the code formats it,
    never the machine's regional settings)."""
    if value is None:
        return ""
    if kind == "money" and (number := stored_decimal(value)) is not None:
        return f"{number:,.2f}".replace(",", " ").replace(".", ",")
    return str(value)


def _put(cell: Any, value: Any, kind: str) -> None:
    v = _typed(value, kind)
    cell.value = v
    if isinstance(v, str):
        cell.data_type = "s"  # text is never a formula
    elif cell.number_format == "General" and isinstance(v, Decimal) and kind == "money":
        cell.number_format = xlsx.MONEY_FORMAT
    elif cell.number_format == "General" and isinstance(v, date):
        cell.number_format = DATE_FORMAT


@dataclass
class Lookup:
    """An output's data by name, for its markers (shared with the Word filler, `jav/outputs/docx.py`)."""

    output: str
    values: dict[str, Any]
    tables: dict[str, contract.Table]

    @classmethod
    def of(cls, data: contract.OutputData) -> Lookup:
        return cls(data.output, {v.key: v.value for v in data.values}, {t.part or t.key: t for t in data.tables})

    def table(self, name: str) -> contract.Table | None:
        return self.tables.get(name)

    def kind(self, m: markers.Marker, column: str | None = None) -> str:
        if m.kind == "value":
            return contract.spec(self.output)["values"][m.name]
        if m.kind == "count":
            return "count"
        return markers.column_kind(self.output, m.table, m.part, column or m.column)

    def single(self, m: markers.Marker) -> Any:
        """The value of a value, total or count marker."""
        if m.kind == "value":
            return self.values.get(m.name)
        t = self.table(m.name)
        if m.kind == "count":
            return len(t.rows) if t else 0
        total = Decimal(0)
        if t is not None:
            i = t.columns.index(m.column) if m.column in t.columns else None
            for row in t.rows if i is not None else ():
                number = stored_decimal(row[i]) if i < len(row) else None
                total += number if number is not None else 0
        return total

    def months(self, name: str) -> list[int]:
        t = self.table(name)
        return [i for i, c in enumerate(t.columns) if contract.MONTH.match(c)] if t else []


def _spread(ws: Any, cell: Any, values: list[tuple[Any, str]]) -> None:
    """Values in the cell and to its right, each with the cell's formatting and its column's width."""
    from openpyxl.utils import get_column_letter

    width = ws.column_dimensions[get_column_letter(cell.column)].width
    if not values:
        cell.value = None
        return
    for k, (value, kind) in enumerate(values):
        target = ws.cell(cell.row, cell.column + k)
        if k:
            target._style = copy(cell._style)
            letter = get_column_letter(cell.column + k)
            if width and not ws.column_dimensions[letter].width:
                ws.column_dimensions[letter].width = width
        _put(target, value, kind)


def _shift_below(ws: Any, row: int, by: int) -> None:
    """What `insert_rows`/`delete_rows` leave in place below `row`: merged cells, pictures and row heights."""
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > row:
            rng.shift(row_shift=by)
    for img in ws._images:
        anchor = img.anchor
        start = getattr(anchor, "_from", None)
        if start is not None and start.row >= row:  # 0-based: the picture starts below the row
            start.row += by
            end = getattr(anchor, "to", None)
            if end is not None:
                end.row += by
    heights = {i: d.height for i, d in list(ws.row_dimensions.items()) if i > row and d.height is not None}
    for i in heights:
        ws.row_dimensions[i].height = None
    for i, h in heights.items():
        ws.row_dimensions[i + by].height = h


def _text(text: str, found: list[markers.Marker], value_of: Any) -> str:
    """The cell's text with each marker replaced once; what a value contains is never read again."""
    by_text = {m.text: m for m in found}
    return markers.ANY.sub(lambda g: value_of(by_text[g.group(0)]) if g.group(0) in by_text else g.group(0), text)


def _fill_singles(ws: Any, cells: list[tuple[Any, str, list[markers.Marker]]], rows: dict[int, str], d: Lookup) -> None:
    for c, text, found in cells:
        if c.row in rows:
            continue
        if len(found) == 1 and found[0].kind == "month_names":
            t = d.table(found[0].name)
            _spread(ws, c, [(t.head[i], "text") for i in d.months(found[0].name)] if t else [])
        elif len(found) == 1 and text.strip() == found[0].text:
            _put(c, d.single(found[0]), d.kind(found[0]))
        else:
            c.value = xlsx.cell_text(_text(text, found, lambda m: _shown(d.single(m), d.kind(m))))
            c.data_type = "s"


def _fill_row(ws: Any, r: int, cells: list[tuple[Any, str, list[markers.Marker]]], d: Lookup, name: str) -> None:
    t = d.table(name)
    n = len(t.rows) if t else 0
    if n == 0:
        ws.delete_rows(r)
        _shift_below(ws, r, -1)
        return
    if n > 1:
        last = ws.max_column
        height = ws.row_dimensions[r].height
        ws.insert_rows(r + 1, n - 1)
        _shift_below(ws, r, n - 1)
        for k in range(1, n):
            for col in range(1, last + 1):
                ws.cell(r + k, col)._style = copy(ws.cell(r, col)._style)
            if height is not None:
                ws.row_dimensions[r + k].height = height
    templates = [(c.column, text, found) for c, text, found in cells]
    months = d.months(name)
    for k, row in enumerate(t.rows):
        def cell_of(m: markers.Marker, row: Any = row) -> Any:
            i = t.columns.index(m.column) if m.column in t.columns else len(row)  # a list column no row of it has
            return row[i] if i < len(row) else None
        for col, text, found in templates:
            target = ws.cell(r + k, col)
            if len(found) == 1 and found[0].kind == "months":
                _spread(ws, target, [(row[i] if i < len(row) else None, d.kind(found[0], t.columns[i])) for i in months])
            elif len(found) == 1 and text.strip() == found[0].text:
                _put(target, cell_of(found[0]), d.kind(found[0]))
            else:
                target.value = xlsx.cell_text(_text(text, found, lambda m: _shown(cell_of(m), d.kind(m))))
                target.data_type = "s"


def fill(data: contract.OutputData, template: bytes) -> bytes:
    """The output's workbook from the template. TemplateError when the template does not fill (checked when stored)."""
    contract.check(data)
    d = Lookup.of(data)
    wb = _load(template)
    for ws in wb.worksheets:
        rows, cells = _scan(ws, data.output, None)
        _fill_singles(ws, cells, rows, d)  # first: a value's position moves with the rows inserted below it
        for r in sorted(rows, reverse=True):  # from the bottom: a row's insertions never move a row above it
            _fill_row(ws, r, [x for x in cells if x[0].row == r], d, rows[r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
