"""The shared Excel table writer of every download (178, plan 170 L4).

Until 178 two hand-written copies wrote the tables: the result workbook's (`export._write_sheet`, also the dataset
download) and the reconciliation statement's (`reconcile_statement._sheet`, with a heading block above the table).
One writer now writes both. A table is written the same way everywhere:
- a text cell is always text, never a formula (153, review C8), a control character a PDF's text layer may carry
  becomes a space (145), and a money or number column holds a `Decimal` with two decimals shown;
- the header row is bold on a light grey fill, frozen and filtered;
- each column is as wide as its longest value among the first 200 rows, between a minimum and 60 characters.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from decimal import Decimal, InvalidOperation
from typing import Any

SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
HEADER_FILL = "EEEEF2"
MONEY_FORMAT = "#,##0.00"
WIDTH_MAX = 60
WIDTH_ROWS = 200  # the rows a column's width is measured on


def safe_sheet(name: str, used: set[str]) -> str:
    """A sheet name the workbook format accepts (no `[]:*?/\\`, at most 31 characters), unique in `used` ignoring case."""
    base = SHEET_BAD.sub("_", name)[:31] or "Lap"
    out, n = base, 2
    while out.casefold() in used:
        suffix = f" ({n})"
        out, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(out.casefold())
    return out


def cell_text(s: str) -> str:
    """145: a control character a PDF's text layer may carry (a broken glyph) is refused by the workbook format, and
    openpyxl raised on it, failing the whole download; it becomes a space."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub(" ", s)


def cell_value(v: Any, kind: str | None) -> Any:
    """A value as the cell holds it: a money or number column's value as a `Decimal` when it reads as one, a number as
    it is, anything else as text."""
    if v is None:
        return None
    if kind in ("money", "number") and not isinstance(v, (dict, list)):
        try:
            return Decimal(str(v))
        except InvalidOperation:
            return cell_text(str(v))
    return v if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) else cell_text(str(v))


def write_table(ws: Any, head: Sequence[str], rows: Sequence[Sequence[Any]], kinds: Sequence[str | None] | None = None, *,
                heading: Sequence[tuple[str, Any]] = (), min_width: int = 8) -> None:
    """Write a table on a sheet: an optional heading block of label and value rows with an empty row below it, then the
    header row and the rows. `kinds` gives each column's kind (`money`, `number` or None)."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    top = 1
    if heading:
        for label, value in heading:
            ws.append([label, cell_value(value, None)])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append([])
        top = ws.max_row + 1
    ws.append(list(head))
    for row in rows:
        ws.append([cell_value(v, kinds[i] if kinds and i < len(kinds) else None) for i, v in enumerate(row)])
    for row in ws.iter_rows(min_row=1):  # 153 (review C8): the header row and the heading block too
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"  # text is never a formula
            elif isinstance(c.value, Decimal):
                c.number_format = MONEY_FORMAT
    for c in ws[top]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
    ws.freeze_panes = f"A{top + 1}"
    if rows:
        ws.auto_filter.ref = ws.dimensions if top == 1 else f"A{top}:{get_column_letter(len(head))}{top + len(rows)}"
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:WIDTH_ROWS] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(WIDTH_MAX, max(min_width, width + 2))
