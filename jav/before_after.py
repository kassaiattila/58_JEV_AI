"""Before/after comparison (plan 170, E0): a free check that a refactoring step changed only what it announced.

Plain-language summary: before a step of the refactoring, a backed-up copy of the store is read through the local
service's own read views (the same lists, pages and downloads the interface shows), and every answer is written down.
After the step, a fresh copy of the same backup is read the same way with the new code. The two records are compared
view by view: a difference is allowed only where the step said it would change something.

- `static_snapshot` reads a fresh copy of the backed-up store through `jav.api.create_app` (in process, no network,
  no AI call) and writes one JSON line per view: the packages, their items, runs, reviews and next steps, each
  package's current result with its tables and downloads (172), every run's item results, journal and downloads, the
  datasets behind the tables, the reconciliation views, the parties, the workflows and the activity log of every
  person and day found in the store.
- `diff` compares two snapshots and lists the views that are missing on one side or differ, with the first differing
  places inside each view.

A snapshot holds personal data (document values, names): it is written only to the path the caller gives, which must
be outside git (the scratchpad or `store/`). The copy is made fresh for every snapshot, so a step that adds a table
or migrates one never touches the backup itself.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import sqlite3
import zipfile
from collections import Counter
from collections.abc import Iterator
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

FORMAT_VERSION = 1
BASE_URL = "http://127.0.0.1:8930"
ACTOR = "before.after"

# Keys whose values say when an answer was built rather than what it is; dropped before writing and comparing.
VOLATILE_KEYS = frozenset({"refreshing"})

# The dataset of each scope kind, so that every table the interface can show is read for every package or run.
_RUN_DATASETS = ("run_items", "emails", "email_tasks", "documents", "datapoints", "native_facts", "line_items",
                 "utility_cost", "utility_sources", "file_names", "calls")
_EXPORT_TABLES = ("documents", "datapoints", "line_items", "native_facts")
# 172: the result tables of a package's current result (each item from its counted run)
_PACKAGE_DATASETS = ("emails", "email_tasks", "documents", "datapoints", "native_facts", "line_items", "utility_cost",
                     "utility_sources")

# The tables whose (actor, time) pairs give the persons and days of the activity log (the sources of
# `jav/core/activity_kinds.py`; a test checks that every registered source's table is here).
_ACTIVITY_TIMES = (
    ("workpackage_events", "actor", "created_at"), ("recipe_assignments", "actor", "created_at"),
    ("runs", "actor", "created_at"), ("runs", "approved_by", "approved_at"),
    ("run_item_corrections", "actor", "created_at"), ("review_reasons", "actor", "closed_at"),
    ("email_task_decisions", "actor", "decided_at"), ("email_task_decisions", "done_by", "done_at"),
    ("mailbox_pulls", "actor", "created_at"),
    # 175: a person's duplicate decisions, statement choices and own workflow steps
    ("duplicate_decisions", "actor", "decided_at"), ("statement_choices", "actor", "decided_at"),
    ("workflow_events", "actor", "created_at"),
)


class SnapshotError(Exception):
    """The snapshot cannot be taken (no store in the backup, an existing output, a path inside git)."""


@dataclass(frozen=True)
class Difference:
    key: str
    kind: str  # "only_before", "only_after" or "changed"
    places: tuple[str, ...] = ()


# --- normalisation ------------------------------------------------------------------------------------------------

def _clean(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: _clean(v) for k, v in value.items() if k not in VOLATILE_KEYS}
    if isinstance(value, list):
        return [_clean(v) for v in value]
    return value


def _xlsx_cells(data: bytes) -> dict[str, list[list[str]]]:
    """The cell values of a workbook, sheet by sheet (the file's own creation time and styles are left out)."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return {ws.title: [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
                for ws in book.worksheets}
    finally:
        book.close()


def _zip_members(data: bytes) -> dict[str, str]:
    """The members of a zip file by name and content hash (the members' timestamps are left out)."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {name: hashlib.sha256(z.read(name)).hexdigest() for name in sorted(z.namelist())}


def _csv_rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


def _body(kind: str, response: Any) -> Any:
    if response.status_code >= 400:
        try:
            payload = response.json()
        except ValueError:
            return response.text
        return _clean({k: v for k, v in payload.items() if k in ("error", "message", "code")}) \
            if isinstance(payload, dict) else payload
    if kind == "json":
        return _clean(response.json())
    if kind == "csv":
        return _csv_rows(response.content)
    if kind == "xlsx":
        return _xlsx_cells(response.content)
    if kind == "zip":
        return _zip_members(response.content)
    raise ValueError(f"unknown body kind: {kind}")


# --- the static snapshot ------------------------------------------------------------------------------------------

def _store_of(baseline: Path) -> Path:
    """The store file of a backup folder (`jav.cli backup`), or the store file itself."""
    path = baseline / "jav.sqlite" if baseline.is_dir() else baseline
    if not path.is_file():
        raise SnapshotError(f"no store file in the backup: {path}")
    return path


def fresh_copy(baseline: Path, work_dir: Path) -> Path:
    """A fresh copy of the backed-up store in `work_dir`, with the files kept next to a store: the native readings
    (`native/`, from the backup's `native-artifacts.zip`), the source instances (`sources/`, the backup root's shared
    folder) and the statement files read from bank tables (`statement_tables/`, from the work store's folder: the
    backup does not hold them, they are derived). An earlier copy in `work_dir` is replaced.

    The folder is used by its full, resolved name: a short (8.3) Windows name would make the native readings' path
    check refuse every file."""
    from jav.core import store

    source = _store_of(baseline)
    work_dir = work_dir.resolve()
    if work_dir.exists():
        shutil.rmtree(work_dir)
    work_dir.mkdir(parents=True)
    target = work_dir / "jav.sqlite"
    shutil.copyfile(source, target)
    native = source.parent / "native-artifacts.zip"
    if native.is_file():
        with zipfile.ZipFile(native) as z:
            z.extractall(work_dir)
    for folder, origin in (("sources", source.parent.parent / "sources"),
                           ("statement_tables", store.STORE_PATH.parent / "statement_tables")):
        if origin.is_dir():
            shutil.copytree(origin, work_dir / folder)
    return target


def _activity_pairs(store_path: Path) -> list[tuple[str, str]]:
    """Every (person, local day) with an action in the store; the activity log is read for each."""
    pairs: set[tuple[str, str]] = set()
    with closing(sqlite3.connect(f"file:{store_path.as_posix()}?mode=ro", uri=True)) as c:  # `with` alone does not close
        tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        for table, who, when in _ACTIVITY_TIMES:
            if table not in tables:
                continue
            for actor, at in c.execute(f"SELECT {who}, {when} FROM {table} WHERE {who} IS NOT NULL AND {when} IS NOT NULL"):  # noqa: S608 - fixed names
                try:
                    day = datetime.fromisoformat(str(at)).astimezone().date().isoformat()
                except ValueError:
                    continue
                pairs.add((" ".join(str(actor).split()), day))
    return sorted(pairs)


class _Reader:
    """Reads the service's views in process and yields (key, value) pairs."""

    def __init__(self, client: Any) -> None:
        self.client = client
        self.counts: Counter[str] = Counter()

    def get(self, key: str, path: str, kind: str = "json", **params: Any) -> tuple[str, Any]:
        response = self.client.get("/api" + path, params=params or None)
        self.counts[key.split("/", 1)[0]] += 1
        return key, {"status": response.status_code, "body": _body(kind, response)}

    def dataset(self, name: str, scope: dict[str, str]) -> tuple[str, Any]:
        """Every row of a dataset (the interface pages it by 1000 at most)."""
        key = f"dataset/{name}/" + json.dumps(scope, sort_keys=True, ensure_ascii=False)
        rows: list[Any] = []
        page: dict[str, Any] = {}
        offset = 0
        while True:
            response = self.client.post(f"/api/datasets/{name}/query",
                                        json={"scope": scope, "query": {"offset": offset, "limit": 1000}})
            if response.status_code >= 400:
                self.counts["dataset"] += 1
                return key, {"status": response.status_code, "body": _body("json", response)}
            page = response.json()
            rows.extend(page.get("rows") or [])
            offset += 1000
            if offset >= int(page.get("total") or 0):
                break
        self.counts["dataset"] += 1
        head = _clean({k: v for k, v in page.items() if k != "rows"})
        return key, {"status": 200, "body": {**head, "rows": _clean(rows)}}


def _views(reader: _Reader, store_path: Path, *, items: bool, matches: bool) -> Iterator[tuple[str, Any]]:
    g = reader.get
    key, packages = g("workpackages", "/workpackages", include_archived="true")
    yield key, packages
    wps = (packages["body"] or {}).get("workpackages") or []
    yield reader.dataset("workpackages", {"include_archived": "1"})
    yield reader.dataset("runs", {})
    run_ids: list[str] = []
    for wp in sorted(wps, key=lambda w: w["id"]):
        wid = wp["id"]
        yield g(f"workpackage/{wid}", f"/workpackages/{wid}")
        yield g(f"workpackage/{wid}/reviews", f"/workpackages/{wid}/reviews")
        yield g(f"workpackage/{wid}/workflow", f"/workpackages/{wid}/workflow")
        yield g(f"workpackage/{wid}/readiness", f"/workpackages/{wid}/workflow/readiness")
        key, runs = g(f"workpackage/{wid}/runs", f"/workpackages/{wid}/runs")
        yield key, runs
        run_ids += [r["run_id"] for r in (runs["body"] or {}).get("runs") or []]
        yield reader.dataset("runs", {"workpackage_id": wid})
        yield reader.dataset("workpackage_items", {"workpackage_id": wid, "removed": "1"})
        yield reader.dataset("package_costs", {"workpackage_id": wid})
        yield g(f"workpackage/{wid}/result", f"/workpackages/{wid}/result")  # 171/172: its current result and release
        if wp.get("source_kind") == "reconcile":
            yield from _reconcile_views(reader, wid, matches=matches)
        else:
            yield g(f"workpackage/{wid}/result/export.xlsx", f"/workpackages/{wid}/result/export.xlsx", "xlsx")
            yield g(f"workpackage/{wid}/result/utility-cost", f"/workpackages/{wid}/result/utility-cost")
            for name in _PACKAGE_DATASETS:
                yield reader.dataset(name, {"workpackage_id": wid})
    for rid in sorted(set(run_ids)):
        key, run = g(f"run/{rid}", f"/runs/{rid}")
        yield key, run
        yield g(f"run/{rid}/journal", f"/runs/{rid}/journal")
        yield g(f"run/{rid}/results", f"/runs/{rid}/results")
        for table in _EXPORT_TABLES:
            yield g(f"run/{rid}/export/{table}.csv", f"/runs/{rid}/export", "csv", format="csv", table=table)
        yield g(f"run/{rid}/export/documents.json", f"/runs/{rid}/export", "json", format="json", table="documents")
        yield g(f"run/{rid}/export/documents.xlsx", f"/runs/{rid}/export", "xlsx", format="xlsx", table="documents")
        yield g(f"run/{rid}/named-copies", f"/runs/{rid}/named-copies.zip", "zip")
        yield g(f"run/{rid}/utility-cost", f"/runs/{rid}/reports/utility-cost")
        for name in _RUN_DATASETS:
            yield reader.dataset(name, {"run_id": rid})
        if items:
            for item in sorted({i["item_id"] for i in (run["body"] or {}).get("items") or []}):
                yield g(f"run/{rid}/item/{item}", f"/runs/{rid}/items/{quote(item, safe='')}")
    key, parties = g("parties", "/parties")
    yield key, parties
    for party in (parties["body"] or {}).get("parties") or []:
        yield g(f"party/{party['id']}/months", f"/parties/{party['id']}/months")
    yield g("reconcile-accounts", "/reconcile/accounts")
    yield g("workflow-templates", "/workflow-templates")
    key, flows = g("workflows", "/workflows")
    yield key, flows
    for wf in (flows["body"] or {}).get("workflows") or []:
        yield g(f"workflow/{wf['id']}", f"/workflows/{wf['id']}")
    yield reader.dataset("mailbox_pulls", {})
    for actor, day in _activity_pairs(store_path):
        yield reader.dataset("activity", {"actor": actor, "day": day})


def _reconcile_views(reader: _Reader, wid: str, *, matches: bool) -> Iterator[tuple[str, Any]]:
    g = reader.get
    key, workspace = g(f"reconcile/{wid}", f"/workpackages/{wid}/reconcile")
    yield key, workspace
    yield g(f"reconcile/{wid}/quick-review", f"/workpackages/{wid}/reconcile/quick-review")
    yield g(f"reconcile/{wid}/targets", f"/workpackages/{wid}/reconcile/targets")
    yield g(f"reconcile/{wid}/statement", f"/workpackages/{wid}/reconcile/statement")
    yield g(f"reconcile/{wid}/statement.xlsx", f"/workpackages/{wid}/reconcile/statement.xlsx", "xlsx")
    yield reader.dataset("reconcile_inputs", {"workpackage_id": wid, "others": "1", "left_out": "1"})
    if matches:
        lines = ((workspace["body"] or {}).get("lines") or []) if isinstance(workspace["body"], dict) else []
        for line in sorted({str(ln["id"]) for ln in lines if isinstance(ln, dict) and ln.get("id")}):
            yield g(f"reconcile/{wid}/matches/{line}", f"/workpackages/{wid}/reconcile/matches", line_id=line)


def _outside_git(path: Path) -> None:
    """A snapshot and its store copy hold personal data: only outside the repository or under its ignored `store/`."""
    from jav.core.config import PROJECT_ROOT

    full = path.resolve()
    if full.is_relative_to(PROJECT_ROOT.resolve()) and not full.is_relative_to((PROJECT_ROOT / "store").resolve()):
        raise SnapshotError(f"inside the repository (personal data): {path}; use the scratchpad or store/")


def _git_line() -> str:
    from jav.core import devstate

    state = devstate.git_state()
    return state.line() if state else "no git"


def static_snapshot(baseline: Path, work_dir: Path, out: Path, *, items: bool = True,
                    matches: bool = False) -> dict[str, int]:
    """Read a fresh copy of the backed-up store through the service's views and write them to `out` (JSON lines).

    Returns the number of views read per kind. The process's default store is pointed at the copy as well, so that a
    background rebuild started by a view cannot read the work store."""
    from fastapi.testclient import TestClient

    from jav.app import datasets
    from jav.core import store
    from jav.api import create_app

    for path in (out, work_dir):
        _outside_git(path)
    if out.exists():
        raise SnapshotError(f"the output already exists: {out}")
    copy = fresh_copy(baseline, work_dir)
    default = store.STORE_PATH
    store.STORE_PATH = copy
    datasets.clear_cache()  # a dataset cached from another copy at the same path would hide a difference
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        with TestClient(create_app(store_path=copy), base_url=BASE_URL, headers={"X-Actor": ACTOR}) as client, \
                out.open("w", encoding="utf-8", newline="\n") as f:
            meta = {"format": FORMAT_VERSION, "git": _git_line(), "baseline": str(_store_of(baseline)),
                    "taken_at": datetime.now().astimezone().isoformat(timespec="seconds"), "items": items,
                    "matches": matches}
            f.write(json.dumps({"meta": meta}, ensure_ascii=False) + "\n")
            reader = _Reader(client)
            for key, value in _views(reader, copy, items=items, matches=matches):
                f.write(json.dumps({"key": key, "value": value}, ensure_ascii=False, sort_keys=True, default=str) + "\n")
    finally:
        datasets.wait_for_rebuilds()
        datasets.clear_cache()
        store.STORE_PATH = default
    return dict(reader.counts)


# --- comparison -----------------------------------------------------------------------------------------------------

def load(path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    """(meta, views by key) of a snapshot file."""
    meta: dict[str, Any] = {}
    views: dict[str, Any] = {}
    with path.open(encoding="utf-8") as f:
        for line in f:
            row = json.loads(line)
            if "meta" in row:
                meta = row["meta"]
            else:
                views[row["key"]] = row["value"]
    return meta, views


def _places(before: Any, after: Any, at: str = "", limit: int = 20) -> list[str]:
    """The first places (JSON paths) where two values differ."""
    if limit <= 0:
        return []
    if isinstance(before, dict) and isinstance(after, dict):
        found: list[str] = []
        for k in sorted(set(before) | set(after), key=str):
            if k not in before:
                found.append(f"{at}/{k}: only after")
            elif k not in after:
                found.append(f"{at}/{k}: only before")
            elif before[k] != after[k]:
                found += _places(before[k], after[k], f"{at}/{k}", limit - len(found))
            if len(found) >= limit:
                break
        return found[:limit]
    if isinstance(before, list) and isinstance(after, list):
        found = [] if len(before) == len(after) else [f"{at}: {len(before)} -> {len(after)} elements"]
        for n, (b, a) in enumerate(zip(before, after, strict=False)):
            if b != a:
                found += _places(b, a, f"{at}[{n}]", limit - len(found))
            if len(found) >= limit:
                break
        return found[:limit]
    return [f"{at}: {_short(before)} -> {_short(after)}"]


def _short(value: Any, width: int = 80) -> str:
    text = json.dumps(value, ensure_ascii=False, default=str)
    return text if len(text) <= width else text[: width - 1] + "…"


def diff(before: Path, after: Path, *, limit: int = 20) -> list[Difference]:
    """The views missing on one side or differing between two snapshots, in key order."""
    _, a = load(before)
    _, b = load(after)
    out: list[Difference] = []
    for key in sorted(set(a) | set(b)):
        if key not in b:
            out.append(Difference(key, "only_before"))
        elif key not in a:
            out.append(Difference(key, "only_after"))
        elif a[key] != b[key]:
            out.append(Difference(key, "changed", tuple(_places(a[key], b[key], limit=limit))))
    return out


def kind_of(key: str) -> str:
    """The kind of a view for the summary: `run/<id>/export/documents.csv` -> `run/export/documents.csv`."""
    parts = key.split("/")
    if parts[0] == "dataset":
        return "/".join(parts[:2])
    rest = parts[2:]  # the second part is the package, run, party or workflow id
    if rest and rest[0] in ("item", "matches"):  # their last part is an item or a statement line
        rest = rest[:1]
    return "/".join([parts[0], *rest])


def summary(differences: list[Difference]) -> dict[str, Counter[str]]:
    """Differences counted per kind of view and per kind of difference."""
    by: dict[str, Counter[str]] = {}
    for d in differences:
        by.setdefault(kind_of(d.key), Counter())[d.kind] += 1
    return by
