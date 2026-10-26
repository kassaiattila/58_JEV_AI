"""231 (stage D, D-f3): the insurance quote's process pack, from the documents to the premium decision.

A case of the pack `insurance_offer` starts on a folder of the invented reference set's documents (PDFs). A person
starts the processing run; the run's results are what a correct reading gives (written as the worker would, without
asking a model), and the run is approved. The documents' facts are offered with where each was printed, the company
register and the sanctions list are asked, a person freezes the facts, the rules allow the risk, nothing is missing,
a person accepts the risk, the premium is the reference case's and, above the invented limit, another person accepts
it. D-f7: the accepted premium is printed as the offer's Word document, stored with the decision's version: its values
are the case's (the policyholder, the product, each module's premium, the approver), and the download and the
Decision and release view name that file. Invented data; no paid call.
"""

from __future__ import annotations

import hashlib
import io
import time

import pytest

from jav.app import (approved_document, case_decisions, case_views, document_processing, field_check, package_release,
                     process_runtime, product_sets, rated_offer, work, work_runs)
from jav.core import connectors, store, trace
from jav.core.process import cases, engine
from jav.devtools import invented_services
from jav.evals import property_documents as pd

ANNA, BELA = "Minta Anna", "Minta Bela"
SET = next(s for s in pd.load().sets if s.id == "reference")
CASE = next(c for c in product_sets.load().cases if c.id == "reference")


@pytest.fixture()
def folder(tmp_path):
    with store.use_store(tmp_path / "offer.sqlite"), connectors.use_transport(invented_services.transport()):
        product_sets.install(product_sets.load(), maker=ANNA, approver=BELA)
        out = tmp_path / "documents"
        out.mkdir()
        for document in SET.documents:
            from jav.documents.synthetic_pdf import write_unicode_pdf_pages

            write_unicode_pdf_pages(out / f"{document.id}.pdf", pd.pages(document))
        yield out


_CLOCK = [time.time()]


def _turn(case_id: str, n: int = 2) -> dict:
    for _ in range(n):
        _CLOCK[0] = max(_CLOCK[0], time.time()) + 10  # never behind the real time the steps' timers are set from
        engine.turn(case_id, now=_CLOCK[0])
    return engine.state(case_id)


def _read_as_correct(run_id: str, wp_id: str) -> None:
    """The run's items finished with what a correct reading of each gives (the worker would ask a model)."""
    by_name = {d.id: d for d in SET.documents}
    for item in work.get(wp_id)["items"]:
        document = by_name[str(item["source_path"]).replace("\\", "/").rsplit("/", 1)[-1].removesuffix(".pdf")]
        flow = work_runs.flow_run_id(run_id, item["item_id"])
        store.insert_datapoints(run_id=flow, doc_id=item["item_id"], doc_type=document.pack, arm="G", datapoints=document.values,
                                field_conf={}, validation=[], route="auto", review_reasons=[], final_status="done")
        work_runs.record_item_result(run_id, item["item_id"], status="done", final_status="done", flow_run_id=flow)
    with store.connect() as c:  # the run's jobs done, as the worker leaves them
        c.execute("UPDATE jobs SET status='done' WHERE run_id=?", (run_id,))
    work_runs.refresh_run_status(run_id)


def _decide(case_id: str, step: str, actor: str) -> None:
    case_views.act(case_id, step, "decide", actor, values={"choice": "accept",
                                                          "expected_basis": case_decisions.view(case_id, step)["basis"]})


def test_a_quote_runs_from_its_documents_to_the_accepted_premium(folder) -> None:
    case_id = process_runtime.start_case("insurance_offer", {"source": {"kind": "folder", "folder": str(folder)}},
                                         name="Minta Kereskedő Kft.", actor=ANNA)
    _turn(case_id)
    case_views.act(case_id, "documents", "start_run", ANNA)
    wp = document_processing.package_of(case_id, "documents")
    [run] = work_runs.run_rows(wp)
    _read_as_correct(run["run_id"], wp)
    _turn(case_id)
    work_runs.approve_run(run["run_id"], actor=ANNA)
    steps = _turn(case_id, 3)["steps"]
    assert [steps[s]["state"] for s in ("documents", "facts", "company", "sanctions")] == ["done"] * 4
    facts = steps["facts"]["result"]
    assert len(facts["documents"]) == 3 and not any(d["unread"] for d in facts["documents"])
    assert steps["company"]["result"]["answer"]["company"]["activity_code"] == "4719"  # the register agrees with the form

    page = case_views.view(case_id)  # D-f4: the Data check view's elements and the page's next step
    shown = {(e["type"], e["step"]): e["data"] for tab in page["tabs"] if tab["id"] == "data" for e in tab["elements"]}
    review = shown[("field_review", "check")]
    assert review["waits"] and review["conflicts"] == 0 and review["catalogue"]["id"] == "property"
    assert page["next"]["code"] == "facts_check" and page["next"]["tab"] == "data"
    assert [d["claims"] > 0 for d in shown[("document_facts", "facts")]["documents"]] == [True] * 3
    assert shown[("lookup", "company")]["question"] == {"tax_number": "12345676-2-41"}  # by fact, not the register's names
    assert shown[("lookup", "sanctions")]["catalogue"]["facts"]["sanctions_listed"]["type"] == "bool"
    view = field_check.view(case_id, "check")
    holder = next(u for u in view["units"] if u["level"] == "holder")
    assert holder["facts"]["company_status"] == "active" and holder["facts"]["sanctions_listed"] is False
    assert holder["sources"]["company_name"]["source"]["file"].endswith(".pdf")
    assert not [c for c in view["claims"] if c["open"]]  # the documents and the registers agree
    case_views.act(case_id, "check", "freeze", ANNA, values={"cover_start": CASE.input["cover_start"],
                                                             "expected_state": view["state"]})
    steps = _turn(case_id, 3)["steps"]
    assert steps["evaluate"]["result"]["verdict"] == CASE.expected.verdict and steps["missing"]["state"] == "done"

    _decide(case_id, "risk", ANNA)
    steps = _turn(case_id, 3)["steps"]
    assert steps["rate"]["result"]["total"] == {"amount": CASE.expected.total, "currency": "HUF"}
    calc = {(e["type"], e["step"]): e["data"] for tab in case_views.view(case_id)["tabs"] if tab["id"] == "calculation"
            for e in tab["elements"]}  # D-f5: the Calculation and versions view
    evaluation, premium, missing = calc[("evaluation", "evaluate")], calc[("calculation", "rate")], calc[("missing_info", "missing")]
    assert evaluation["latest"]["verdict"] == "allow" and len(evaluation["versions"]) == 1
    holder = next(u for u in evaluation["latest"]["units"] if u["level"] == "holder")
    assert holder["named"]["company_name"] == "Minta Kereskedő Kft." and holder["combined"]
    assert premium["latest"]["total"] == {"amount": CASE.expected.total, "currency": "HUF"} and premium["tariff"]["title"]
    assert set(premium["modules"]) == {"property", "burglary", "surge", "interruption"} and premium["latest"]["lines"]
    assert (missing["state"], missing["task"], missing["next"]) == ("done", None, None)
    assert case_decisions.view(case_id, "premium")["second_person"]["needed"] is True
    page = case_views.view(case_id)  # D-f6: the Decision and release view
    shown = {e["step"]: e["data"] for tab in page["tabs"] if tab["id"] == "decision" for e in tab["elements"]}
    assert shown["risk"]["decided"]["choice"] == "accept" and not shown["risk"]["waits"]
    assert shown["premium"]["waits"] and shown["premium"]["second_person"]["earlier_by"] == ANNA
    assert page["next"]["code"] == "decision_make" and page["next"]["tab"] == "decision"
    with pytest.raises(case_decisions.AnotherDeciderNeeded):
        _decide(case_id, "premium", ANNA)
    _decide(case_id, "premium", BELA)
    state = _turn(case_id, 3)
    assert state["steps"]["premium"]["state"] == "done"
    with store.connect() as c:
        assert cases.find_case_in(c, case_id)["status"] == cases.CLOSED

    printed = state["steps"]["document"]  # D-f7: the offer's Word document, stored with the decision's version
    assert printed["state"] == "done" and printed["result"]["output"] == "rated_offer"
    version = printed["result"]["version"]
    assert version == package_release.own_approval(case_id, step="premium")["active"]["version"]
    content, name = approved_document.step_file(case_id, "document")
    assert name == f"biztositasi-ajanlat-Minta-Kereskedo-Kft-v{version}.docx"
    assert printed["result"]["file"] == {"name": name, "sha256": hashlib.sha256(content).hexdigest()}
    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
    assert "Minta Kereskedő Kft." in text and "Kitalált vállalkozói vagyonbiztosítás" in text and "Minta Bela" in text
    cells = [c.text for t in Document(io.BytesIO(content)).tables for row in t.rows for c in row.cells]
    assert "Betöréses lopás és rablás" in cells and len(package_release.own_files(case_id, version)) == 1
    values = {v.key: v.value for v in rated_offer.data(package_release.own_version(case_id, version)["content"], case_id=case_id,
                                                      package="Minta Kereskedő Kft.", version=version).values}
    assert (values["cover_start"], values["total"], values["verdict"]) == (CASE.input["cover_start"], CASE.expected.total, "elfogadható")
    assert "Részösszeg" in text and "Díjfizetési gyakoriság" in cells  # the modules' subtotal, then the factor on the total
    shown = {e["step"]: e["data"] for tab in case_views.view(case_id)["tabs"] if tab["id"] == "decision" for e in tab["elements"]}
    assert shown["document"]["file"]["name"] == name and shown["document"]["from_step"] == "premium"
    from fastapi.testclient import TestClient

    from jav.service.api import create_app

    with TestClient(create_app(store_path=store.current_path()), base_url="http://127.0.0.1:8930") as client:
        got = client.get(f"/api/cases/{case_id}/steps/document/file")
        assert got.status_code == 200 and got.content == content and name in got.headers["content-disposition"]
        assert client.get(f"/api/cases/{case_id}/steps/premium/file").status_code == 404  # it prints no document
    assert all(trace.replay(r) == r["outcome"] for r in trace.rows(case_id=case_id))
