"""172 (plan 170 L1, DECISIONS 167, 172): the result tables of the package's current result.

The Result section shows the package's current result by default: each item from the run whose result counts now
(`package_result.counted_results`), an item removed from the package left out. The run's own tables stay as they were.
A page of the package's table carries the mark of the result it shows, so Approve and release names what the person
saw. The package's next step is Approve and release, also when a run was approved without a released version. Synthetic
PDFs, a fake JEV client, no paid calls."""

import io

from openpyxl import load_workbook

from jav import datasets, package_release, work
from jav.runtime import worker
from tests import test_flexible_package_160c as fp
from tests.pdfgen import INVOICE_LINES, write_text_pdf
from tests.test_flexible_package_160c import HUMAN, _assign, _package, _start, _start_pending
from tests.test_package_release_171 import _live, _result

env = fp.env  # the 160-C fixture: a store of its own, two synthetic invoices, a fake JEV client


def _query(c, name: str, scope: dict) -> dict:
    r = c.post(f"/api/datasets/{name}/query", json={"scope": scope, "query": {"limit": 500}})
    assert r.status_code == 200, r.text
    return r.json()


def _third(env, wp_id: str) -> str:
    third = env["folder"] / "invoice_3.pdf"
    write_text_pdf(third, [line.replace("MINTA-2026-001", "MINTA-2026-003") for line in INVOICE_LINES])
    work.add_documents(wp_id, [third], expected_revision=work.get(wp_id)["revision"])
    return next(i["item_id"] for i in work.get(wp_id)["items"] if i["source_path"].endswith("invoice_3.pdf"))


def test_the_package_tables_show_each_item_from_its_counted_run(env):
    c = env["client"]
    wp = _package(c, env["folder"])
    _assign(c, wp["id"])
    _start(c, wp["id"], mode="shadow")  # a trial first, so that a trial over the pending item is offered later
    worker.run_worker(once=True)
    live = _live(c, wp["id"])
    new = _third(env, wp["id"])
    r = _start_pending(c, wp["id"], mode="shadow")
    assert r.status_code == 201, r.text
    trial = r.json()["run_id"]
    worker.run_worker(once=True)
    scope = {"workpackage_id": wp["id"]}
    page = _query(c, "documents", scope)
    runs = {row["item_id"]: row["_run"] for row in page["rows"]}
    assert runs == {**{i["item_id"]: live for i in wp["items"]}, new: trial}
    assert {row["_wp"] for row in page["rows"]} == {wp["id"]}
    assert page["result_mark"] == _result(c, wp["id"])["mark"] and "review_version" not in page
    assert {row["item_id"] for row in _query(c, "datapoints", scope)["rows"]} == set(runs)
    # the rows follow the package's item order
    assert [row["item_id"] for row in page["rows"]] == [i["item_id"] for i in work.get(wp["id"])["items"]]
    # the run's own table is the run's, as before
    assert {row["item_id"] for row in _query(c, "documents", {"run_id": live})["rows"]} == {i["item_id"] for i in wp["items"]}
    tables = _result(c, wp["id"])["tables"]
    assert "documents" in tables and "datapoints" in tables
    assert "file_names" not in tables  # two runs: the copies are made per run (the run's view has them)


def test_a_single_run_package_also_has_its_file_names_and_a_removed_item_is_left_out(env):
    c = env["client"]
    wp = _package(c, env["folder"])
    _assign(c, wp["id"])
    live = _live(c, wp["id"])
    assert "file_names" in _result(c, wp["id"])["tables"]
    first, second = (i["item_id"] for i in wp["items"])
    r = c.post(f"/api/workpackages/{wp['id']}/items/{first}/remove", headers=HUMAN,
               json={"expected_revision": work.get(wp["id"])["revision"], "reason": "not this package's"})
    assert r.status_code == 200, r.text
    assert [row["item_id"] for row in _query(c, "documents", {"workpackage_id": wp["id"]})["rows"]] == [second]
    # the run's own table has left a removed item out since 160-C (DECISIONS 163 D7)
    assert [row["item_id"] for row in _query(c, "documents", {"run_id": live})["rows"]] == [second]


def test_a_result_table_needs_exactly_one_of_a_run_and_a_package(env):
    c = env["client"]
    wp = _package(c, env["folder"])
    _assign(c, wp["id"])
    live = _live(c, wp["id"])
    for scope in ({}, {"run_id": live, "workpackage_id": wp["id"]}):
        r = c.post("/api/datasets/documents/query", json={"scope": scope})
        assert r.status_code in (400, 422), (scope, r.text)


def test_the_package_result_downloads_as_a_workbook_and_a_utility_report(env):
    c = env["client"]
    wp = _package(c, env["folder"])
    _assign(c, wp["id"])
    _live(c, wp["id"])
    r = c.get(f"/api/workpackages/{wp['id']}/result/export.xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    book = load_workbook(io.BytesIO(r.content))
    first = book[book.sheetnames[0]]
    assert first.max_row == 3  # the header and the two documents
    rep = c.get(f"/api/workpackages/{wp['id']}/result/utility-cost")
    assert rep.status_code == 200 and "series" in rep.json()


def test_the_next_step_is_approve_and_release_until_a_version_holds_the_approval(env):
    c = env["client"]
    wp = _package(c, env["folder"])
    _assign(c, wp["id"])
    run_id = _live(c, wp["id"])
    nxt = c.get(f"/api/workpackages/{wp['id']}").json()["next"]
    assert (nxt["code"], nxt["stage"], nxt["label"]) == ("release", "result", "Jóváhagyás és kiadás")
    assert c.post(f"/api/runs/{run_id}/approve", json={}, headers=HUMAN).status_code == 200  # the way before v1.20.0
    assert c.get(f"/api/workpackages/{wp['id']}").json()["next"]["code"] == "release"  # approved, not released
    row = next(r for r in _query(c, "workpackages", {})["rows"] if r["id"] == wp["id"])
    assert row["next_code"] == "release"
    package_release.backfill()
    assert c.get(f"/api/workpackages/{wp['id']}").json()["next"]["code"] == "done"
    _query(c, "workpackages", {})  # the list answers from its last rows and rebuilds behind (153) ...
    datasets.wait_for_rebuilds()
    row = next(r for r in _query(c, "workpackages", {})["rows"] if r["id"] == wp["id"])
    assert row["next_code"] == "done"  # ... once it sees that a version was recorded
