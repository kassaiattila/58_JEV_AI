"""136 (backlog F-revolut-csv, E4): a bank's tabular statement export read by code, one statement per account.

The owner's decisions (DECISIONS 132, 136): a general tabular reader driven by profiles (`configs/statement_tables.json`),
no model call; only the accounts a person chooses go into the store. The workbooks here are synthetic: made-up
accounts and merchants, the sample account numbers the data guard allows, shaped like the 2026 Revolut export (the
CSV records in the first column, UTF-8 read as cp1250, the summaries first, then one table per account).
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from jav import statement_table as st

IBAN_HUF = "HU88 9963 5873 0162 1314 9181 0034"  # synthetic, allowed by the data guard
IBAN_EUR = "HU65 5763 4061 3447 7315 4358 0066"
IBAN_KID = "HU42 1177 3016 1111 1018 0000 0000"
NAME = "statement_2026-01-01_2026-03-31_hu.xlsx"
HEADER = ["Dátum", "Leírás", "Kategória", "Pénz be-/kifizetése", "Egyenleg", "Általad befizetett adó", "Egyéb adók", "Díjak"]
HEADER_FX = ["Dátum", "Leírás", "Kategória", "Pénz be-/kifizetése", "Pénz be-/kifizetése", "Egyenleg", "Egyenleg",
             "Általad befizetett adó", "Általad befizetett adó", "Egyéb adók", "Egyéb adók", "Díjak", "Díjak"]


def _summary(title: str, iban: str, opening: str, closing: str) -> list[list[str]]:
    return [[title], [], ["Folyószámla adatai"], ["Számlaszám (IBAN)", iban, "BIC", "EXAMPLEXX"],
            ["Befizetett összeg"], ["Átlagos egyenleg", " I. negyedév 2026", opening, "Nyitóegyenleg", opening],
            ["Átlagos egyenleg a II. negyedévben 2026", closing, "Záróegyenleg", closing], ["---------"], []]


HUF_LINES = [["2026. jan. 2.", "Example Market", "Kereskedő", "-1 200,00 HUF", "8 800,00", "0,00 HUF", "0,00 HUF", "0,00 HUF"],
             ["2026. jan. 5.", "Top-up by *1234", "Feltöltés", "5 000,00 HUF", "13 800,00", "0,00 HUF", "0,00 HUF", "0,00 HUF"],
             ["2026. febr. 10.", "Example Stream", "Egyéb", "-2 990,00 HUF", "10 810,00", "0,00 HUF", "0,00 HUF", "10,00 HUF"],
             ["2026. márc. 1.", "Jane Example", "Egyéb", "-810,00 HUF", "10 000,00", "0,00 HUF", "0,00 HUF", "0,00 HUF"]]
EUR_LINES = [["2026. febr. 8.", "Example Hotel", "Kereskedő", "-50,00€", "-19 500,00 HUF", "150,00€", "58 500,00 HUF",
              "0,00€", "0,00 HUF", "0,00€", "0,00 HUF", "0,00€", "0,00 HUF"],
             ["2026. márc. 3.", "Exchanged to EUR", "Átváltás", "100,00€", "39 000,00 HUF", "250,00€", "97 500,00 HUF",
              "0,00€", "0,00 HUF", "0,00€", "0,00 HUF", "0,00€", "0,00 HUF"]]
KID_LINES = [["2026. jan. 9.", "Example Toys", "Kereskedő", "-500,00 HUF", "1 500,00", "0,00 HUF", "0,00 HUF", "0,00 HUF"]]


def _total(lines: list[list[str]], index: int = 3) -> list[str]:
    total = sum(Decimal(r[index].replace("€", "").replace("HUF", "").replace(" ", "").replace(",", ".")) for r in lines)
    return ["Végösszeg", "", "", f"{total:.2f}".replace(".", ",") + " HUF"]


def _table(title: str, header: list[str], lines: list[list[str]]) -> list[list[str]]:
    return [[title], [], ["Tranzakciókivonat"], header, *(list(r) for r in lines), _total(lines), [], ["---------"], []]


def _records() -> list[list[str]]:
    rows = [["Folyószámlák Összefoglalók"], []]
    rows += _summary("Személyes számla (HUF)", IBAN_HUF, "10 000,00 HUF", "10 000,00 HUF")
    rows += _summary("Személyes számla (EUR)", IBAN_EUR, "200,00€", "250,00€")
    rows += _summary("Kids &amp; Teens (HUF)", IBAN_KID, "2 000,00 HUF", "1 500,00 HUF")
    rows += _summary("Kids &amp; Teens (HUF)", IBAN_HUF.replace("88", "77", 1), "0,00 HUF", "0,00 HUF")
    rows += _summary("Megtakarítások  (HUF)", IBAN_EUR.replace("65", "66", 1), "1 000,00 HUF", "1 010,00 HUF")
    rows += [["Folyószámlák Tranzakciókivonatok"], []]
    rows += _table("Személyes számla (HUF)", HEADER, HUF_LINES)
    rows += _table("Személyes számla (EUR)", HEADER_FX, EUR_LINES)
    rows += _table("Kids &amp; Teens (HUF)", HEADER, KID_LINES)
    rows += _table("Kids &amp; Teens (HUF)", HEADER, [])
    rows += [["Megtakarítások  (HUF)"], ["Tranzakciókivonat (kamatbevételi nyugta)"],
             ["Dátum", "Leírás", "Gross rate", "Net rate", "Gross interest"], ["2026. 01. 01.", "Interest", "2%", "1,5%", "0,05 HUF"]]
    return rows


def _as_csv(row: list[str]) -> str:
    out = io.StringIO()
    csv.writer(out, lineterminator="").writerow(row)
    return out.getvalue()


def _mojibake(text: str) -> str:
    """UTF-8 bytes decoded as cp1250; a byte cp1250 leaves undefined comes through as its own code point."""
    return "".join(bytes([b]).decode("cp1250", errors="ignore") or chr(b) for b in text.encode("utf-8"))


def _workbook(tmp_path: Path, rows: list[list[str]] | None = None, *, name: str = NAME, garble: bool = True,
              split_at: int | None = None) -> Path:
    """The records as an export opened in Excel: each CSV record in column A (`split_at`: that record continued in B)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, row in enumerate(rows if rows is not None else _records(), start=1):
        text = _as_csv(row) if len(row) > 1 else (row[0] if row else "")
        text = _mojibake(text) if garble else text
        if split_at is not None and i == split_at:
            cut = text.index(",", len(text) // 2)
            ws.cell(row=i, column=1, value=text[:cut])
            ws.cell(row=i, column=2, value=text[cut + 1:])
        elif text:
            ws.cell(row=i, column=1, value=text)
    path = tmp_path / name
    wb.save(path)
    return path


def _by_title(reading: st.Reading) -> dict[tuple[str, str, int], st.Account]:
    return {(a.title, a.currency, a.occurrence): a for a in reading.accounts}


def test_the_accounts_with_a_transaction_table_are_found_with_their_summaries(tmp_path):
    reading = st.read(_workbook(tmp_path))
    assert reading.profile == "revolut_consolidated_hu" and reading.institution == "Revolut"
    assert set(reading.repairs) == {"csv_in_first_columns", "mojibake", "html_entities"}
    accounts = _by_title(reading)
    assert sorted(accounts) == [("Kids & Teens", "HUF", 1), ("Kids & Teens", "HUF", 2), ("Személyes számla", "EUR", 1),
                               ("Személyes számla", "HUF", 1)]
    huf, eur = accounts[("Személyes számla", "HUF", 1)], accounts[("Személyes számla", "EUR", 1)]
    assert (huf.currency, eur.currency) == ("HUF", "EUR")
    assert huf.key == IBAN_HUF.replace(" ", "") + ":HUF" and eur.key == IBAN_EUR.replace(" ", "") + ":EUR"
    assert accounts[("Kids & Teens", "HUF", 2)].key == IBAN_HUF.replace("88", "77", 1).replace(" ", "") + ":HUF"  # its own summary
    assert [len(a.statement["transactions"]) for a in reading.accounts] == [4, 2, 1, 0]


def test_a_statement_has_the_bank_statement_fields_and_its_checks_pass(tmp_path):
    huf = _by_title(st.read(_workbook(tmp_path)))[("Személyes számla", "HUF", 1)]
    s = huf.statement
    assert {k: s[k] for k in ("statement_type", "account_no", "account_iban", "currency", "period_start", "period_end")} == {
        "statement_type": "bank_account", "account_no": None, "account_iban": IBAN_HUF.replace(" ", ""), "currency": "HUF",
        "period_start": "2026-01-01", "period_end": "2026-03-31"}
    assert (s["opening_balance"], s["closing_balance"], s["total_debit"], s["total_credit"]) == ("10000.00", "10000.00", "5000.00", "5000.00")
    first, top_up = s["transactions"][:2]
    assert first == {"booking_date": "2026-01-02", "value_date": None, "direction": "debit", "amount": "1200.00",
                     "running_balance": "8800.00", "description": "Kereskedő", "counterparty_name": "Example Market",
                     "counterparty_account": None, "memo": None}
    assert (top_up["direction"], top_up["amount"]) == ("credit", "5000.00")
    assert huf.checks == {"running_balance": {"ok": 4, "lines": 4}, "total_row": True, "closing_balance": True}
    assert huf.problems == []


def test_a_repeated_column_takes_the_accounts_own_currency(tmp_path):
    eur = _by_title(st.read(_workbook(tmp_path)))[("Személyes számla", "EUR", 1)]
    lines = eur.statement["transactions"]
    assert [(t["direction"], t["amount"], t["running_balance"]) for t in lines] == [("debit", "50.00", "150.00"), ("credit", "100.00", "250.00")]
    assert (eur.statement["opening_balance"], eur.statement["closing_balance"]) == ("200.00", "250.00")
    assert eur.checks["running_balance"] == {"ok": 2, "lines": 2} and eur.checks["closing_balance"] is True


def test_a_clean_file_is_not_repaired_and_a_record_split_over_two_columns_is_joined(tmp_path):
    rows = _records()
    split = next(i for i, r in enumerate(rows, start=1) if r and r[0] == "2026. febr. 10.")
    reading = st.read(_workbook(tmp_path, rows, garble=False, split_at=split))
    assert "mojibake" not in reading.repairs
    huf = _by_title(reading)[("Személyes számla", "HUF", 1)]
    assert huf.statement["transactions"][2]["counterparty_name"] == "Example Stream" and huf.checks["running_balance"]["ok"] == 4


def test_a_csv_file_without_summaries_takes_its_balances_and_period_from_the_lines(tmp_path):
    path = tmp_path / "export.csv"
    rows = [["Személyes számla (HUF)"], ["Tranzakciókivonat"], HEADER, *HUF_LINES, _total(HUF_LINES)]
    path.write_text("\n".join(_as_csv(r) if len(r) > 1 else r[0] for r in rows) + "\n", encoding="utf-8")
    huf = st.read(path).accounts[0]
    assert huf.key == "Személyes számla (HUF)#1"  # no account number printed: the title and its occurrence
    s = huf.statement
    assert (s["period_start"], s["period_end"]) == ("2026-01-02", "2026-03-01")
    assert (s["opening_balance"], s["closing_balance"]) == ("10000.00", "10000.00")  # 8800 - (-1200); the last balance
    assert huf.checks["closing_balance"] is True


def test_a_broken_balance_and_an_unreadable_row_are_reported(tmp_path):
    rows = _records()
    for r in rows:
        if r and r[0] == "2026. febr. 10.":
            r[4] = "10 900,00"  # the running balance does not follow
        if r and r[0] == "2026. márc. 1.":
            r[3] = "n/a"
    huf = _by_title(st.read(_workbook(tmp_path, rows)))[("Személyes számla", "HUF", 1)]
    assert huf.problems == ["line 4: unreadable amount"]
    assert len(huf.statement["transactions"]) == 3
    assert huf.checks["running_balance"] == {"ok": 2, "lines": 3} and huf.checks["total_row"] is False


def test_a_file_no_profile_knows_is_refused(tmp_path):
    with pytest.raises(st.StatementTableError, match="no statement table"):
        st.read(_workbook(tmp_path, [["Date", "Amount"], ["2026-01-01", "1"]]))


def test_a_file_over_the_row_limit_is_refused(tmp_path, monkeypatch):
    monkeypatch.setitem(st._conf()["limits"], "max_rows", 10)
    with pytest.raises(st.StatementTableError, match="row limit"):
        st.read(_workbook(tmp_path))


def test_the_survey_lists_the_accounts_without_their_values(tmp_path):
    survey = st.survey(st.read(_workbook(tmp_path)))
    assert survey["profile"] == "revolut_consolidated_hu" and len(survey["accounts"]) == 4
    huf = survey["accounts"][0]
    assert huf == {"key": IBAN_HUF.replace(" ", "") + ":HUF", "account": "…" + IBAN_HUF.replace(" ", "")[-4:], "title": "Személyes számla",
                   "occurrence": 1, "currency": "HUF", "lines": 4, "first": "2026-01-02", "last": "2026-03-01",
                   "period_start": "2026-01-01", "period_end": "2026-03-31", "checks_ok": True, "problems": 0,
                   "balance_checked": True, "companion": None}  # 137: an export with its own balances
    assert "Example Market" not in repr(survey)


def test_one_account_number_with_two_currencies_gives_two_accounts(tmp_path):
    rows = [r if r[:1] != ["Számlaszám (IBAN)"] or r[1] != IBAN_EUR else [r[0], IBAN_HUF, *r[2:]] for r in _records()]
    accounts = _by_title(st.read(_workbook(tmp_path, rows)))
    huf, eur = accounts[("Személyes számla", "HUF", 1)], accounts[("Személyes számla", "EUR", 1)]
    assert huf.statement["account_iban"] == eur.statement["account_iban"] and huf.key != eur.key
