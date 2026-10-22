"""178 (plan 170 L4, DECISIONS 178): an approved reconciliation statement's download is stored, not rebuilt.

Until 178 the download of an approved version was rebuilt from the stored statement at every download, so a later
change of a label or of the layout (and, with the second round of L4, of a chosen template) would have changed a file
already approved. The file is now stored when the version is approved, and again when it is reopened (its heading then
names it an earlier version, as before); the download returns the stored bytes. A version approved before 178 is
rebuilt as before until `reconcile-versions-backfill` records its file. Synthetic data only; no AI call, no network.
"""

from __future__ import annotations

import copy
import hashlib
import io

import pytest
from openpyxl import load_workbook

from jav import package_release, cli, store
from jav import reconcile_statement as rs
from tests.test_characterisation_178 import _parts
from tests.test_reconcile_statement_145 import _world


@pytest.fixture
def db(tmp_path):
    with store.use_store(tmp_path / "f178.sqlite"):
        yield


def _approved() -> str:
    wp_id, _inv, _ids = _world()
    rs.approve(wp_id, fingerprint=rs.current(wp_id)["fingerprint"], actor="Reviewer")
    return wp_id


def test_the_approved_file_is_stored_and_a_later_label_change_leaves_it(db, monkeypatch):
    wp_id = _approved()
    stored, name = rs.xlsx(wp_id, version=1)
    files = package_release.own_files(wp_id, 1)
    assert [(f["state"], f["name"], f["template"], f["template_version"], f["backfilled_at"]) for f in files] == [
        ("approved", name, "builtin", "1.1.0", None)]
    assert files[0]["sha256"] == hashlib.sha256(stored).hexdigest()
    labels = copy.deepcopy(rs._labels())
    labels["heading"]["package"] = "Package (renamed later)"
    monkeypatch.setattr(rs, "_labels", lambda: labels)
    assert rs.xlsx(wp_id, version=1) == (stored, name)  # the approved file stays as it was approved
    draft = load_workbook(io.BytesIO(rs.xlsx(wp_id)[0]))
    assert draft.worksheets[0].cell(1, 1).value == "Package (renamed later)"  # the draft follows the labels


def test_reopening_stores_the_earlier_file_as_the_download_built_it(db):
    wp_id = _approved()
    approved = rs.xlsx(wp_id, version=1)[0]
    rs.reopen(wp_id, note="a payment was missed", actor="Reviewer")
    earlier, _name = rs.xlsx(wp_id, version=1)
    assert [f["state"] for f in package_release.own_files(wp_id, 1)] == ["approved", "earlier"]
    assert _parts(earlier) != _parts(approved)  # the heading names it an earlier version
    with store.connect() as c:
        c.execute("DELETE FROM package_version_files WHERE state IN ('approved', 'earlier')")
    assert _parts(rs.xlsx(wp_id, version=1)[0]) == _parts(earlier)  # the same file the download built before 178


def test_a_version_approved_before_178_is_rebuilt_until_its_file_is_recorded(db):
    wp_id = _approved()
    with store.connect() as c:
        c.execute("DELETE FROM package_version_files WHERE state IN ('approved', 'earlier')")  # as a version approved before 178
    rebuilt = rs.xlsx(wp_id, version=1)[0]
    assert package_release.own_files(wp_id, 1) == []  # a download writes nothing
    assert rs.backfill_files(dry_run=True) == [{"workpackage_id": wp_id, "version": 1, "state": "approved"}]
    assert package_release.own_files(wp_id, 1) == []
    assert rs.backfill_files() == [{"workpackage_id": wp_id, "version": 1, "state": "approved"}]
    files = package_release.own_files(wp_id, 1)
    assert [f["state"] for f in files] == ["approved"] and files[0]["backfilled_at"]
    assert _parts(rs.xlsx(wp_id, version=1)[0]) == _parts(rebuilt)
    assert rs.backfill_files() == []  # idempotent


def test_the_command_records_the_missing_files(db, capsys):
    wp_id = _approved()
    with store.connect() as c:
        c.execute("DELETE FROM package_version_files WHERE state IN ('approved', 'earlier')")
    assert cli.main(["reconcile-versions-backfill", "--dry-run"]) == 0
    assert f"{wp_id} version 1 (approved) (dry run)" in capsys.readouterr().out
    assert cli.main(["reconcile-versions-backfill"]) == 0
    assert "versions: 1" in capsys.readouterr().out and len(package_release.own_files(wp_id, 1)) == 1
