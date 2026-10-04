"""Native contract and immutable storage regression checks with synthetic data."""
import json
from pathlib import Path
import sqlite3

from pydantic import ValidationError
import pytest

from jav.native_contracts import NativeItemResult, ReceiptRef, SourcePage
from jav import native_results, store
from jav.readers import pipeline
from native_fixtures_109 import contract_examples, publish_runtime, runtime_source


@pytest.fixture
def isolated(tmp_path):
    with store.use_store(tmp_path / "native.sqlite"):
        yield tmp_path


def test_all_shared_native_contract_examples_validate():
    examples = contract_examples()
    for payload in examples["items"].values():
        assert NativeItemResult.model_validate_json(json.dumps(payload)).model_dump(mode="json") == payload
    for payload in examples["source_pages"].values():
        assert SourcePage.model_validate_json(json.dumps(payload)).model_dump(mode="json") == payload


def test_rejected_and_uncertain_are_not_empty_successes():
    cases = contract_examples()["items"]
    assert cases["empty_success"]["interpretation"]["status"] == "empty"
    for name in ("rejected", "uncertain", "failed"):
        assert cases[name]["interpretation"] is None
        assert cases[name]["interpretation_outcome"]["status"] == name
    assert cases["rejected"]["interpretation_outcome"]["receipt_refs"][0]["cost_known"]
    assert cases["uncertain"]["interpretation_outcome"]["receipt_refs"][0]["cost_usd"] is None


def test_unknown_cost_cannot_be_serialised_as_zero():
    receipt = contract_examples()["items"]["uncertain"]["interpretation_outcome"]["receipt_refs"][0]
    receipt["cost_usd"] = "0"
    with pytest.raises(ValidationError, match="unknown cost remains null"):
        ReceiptRef.model_validate_json(json.dumps(receipt))


def test_unpublished_result_cannot_claim_publication_identity():
    payload = contract_examples()["items"]["not_started"]
    payload["result_version"] = "a" * 64
    with pytest.raises(ValidationError, match="Unpublished"):
        NativeItemResult.model_validate_json(json.dumps(payload))


def test_all_delivery_collections_roundtrip_and_backup_inventory(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch, raster=True)
    delivery = native_results.load_reading(ref.reading_id)
    assert all(getattr(delivery, name) for name in ("objects", "evidence", "texts", "rasters"))
    assert delivery.bundle.manifest.occurrences[0].original_name == "original.txt"
    with store.connect() as c:
        artifacts = native_results.referenced_artifacts(c)
    assert {Path(a.relative_path).parts[-2] for a in artifacts} >= {"objects", "evidence", "texts", "rasters"}
    for artifact in artifacts:
        data = (isolated / artifact.relative_path).read_bytes()
        assert len(data) == artifact.byte_size and pipeline.digest(data) == artifact.sha256


@pytest.mark.parametrize("collection", ["objects", "evidence", "texts", "rasters"])
def test_every_saved_collection_is_verified_before_publication(isolated, monkeypatch, collection):
    _, ref = runtime_source(isolated, monkeypatch, raster=True)
    path = next((isolated / ref.artifact_relpath / collection).iterdir())
    path.write_bytes(b"tampered")
    with pytest.raises(native_results.NativeIntegrityError, match="damaged"):
        publish_runtime(ref)
    assert native_results.get_publication("native-run", "native-item") is None


def test_once_only_publication_and_stable_fact_identity(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    first = publish_runtime(ref)
    assert publish_runtime(ref) == first
    assert native_results.machine_facts(first) == native_results.machine_facts(first)
    with pytest.raises(native_results.PublicationConflict):
        publish_runtime(ref, value="Order")
    assert native_results.get_publication("native-run", "native-item") == first


def test_caller_rollback_keeps_publication_and_outcome_atomic(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    with store.connect() as c:
        c.execute("BEGIN IMMEDIATE")
        publication = publish_runtime(ref, c=c)
        native_results.verify_publication(publication, c=c)
        native_results.machine_facts(publication, c=c)
        native_results.validate_sources(publication, {}, c=c)
        assert c.in_transaction
        c.rollback()
    assert native_results.get_publication("native-run", "native-item") is None
    with store.connect() as c:
        assert c.execute("SELECT COUNT(*) FROM native_outcomes").fetchone()[0] == 0


def test_database_publication_error_does_not_leave_visible_result(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    with store.connect() as c:
        c.execute("CREATE TRIGGER publication_failure BEFORE INSERT ON native_publications BEGIN SELECT RAISE(ABORT,'synthetic failure'); END")
    with pytest.raises(sqlite3.IntegrityError, match="synthetic failure"):
        publish_runtime(ref)
    assert native_results.get_publication("native-run", "native-item") is None


def test_interrupted_file_write_has_no_saved_reading(isolated, monkeypatch):
    original = pipeline.Delivery.save

    def interrupted(self, destination):
        Path(destination).mkdir(parents=True)
        (Path(destination) / "partial").write_bytes(b"incomplete")
        raise OSError("Synthetic interrupted write")

    monkeypatch.setattr(pipeline.Delivery, "save", interrupted)
    with pytest.raises(OSError, match="interrupted"):
        runtime_source(isolated, monkeypatch)
    with store.connect() as c:
        assert c.execute("SELECT COUNT(*) FROM native_readings").fetchone()[0] == 0
    monkeypatch.setattr(pipeline.Delivery, "save", original)
    ref = native_results.prepare_reading(isolated / "frozen.txt", original_name="original.txt",
                                        expected_sha256=pipeline.digest((isolated / "frozen.txt").read_bytes()))
    native_results.load_reading(ref.reading_id)


def test_changed_reader_configuration_creates_new_reading_without_overwrite(isolated, monkeypatch):
    path, first = runtime_source(isolated, monkeypatch)
    assert native_results.prepare_reading(path, original_name="original.txt", expected_sha256=first.source_sha256) == first
    monkeypatch.setattr(pipeline, "implementation_version", lambda: "native-synthetic-version-two")
    second = native_results.prepare_reading(path, original_name="original.txt", expected_sha256=first.source_sha256)
    assert second.reading_id != first.reading_id and second.bundle_sha256 != first.bundle_sha256
    assert native_results.load_reading(first.reading_id).bundle.digest() == first.bundle_sha256


def test_source_mismatch_and_relative_path_escape_are_rejected(isolated, monkeypatch):
    path, ref = runtime_source(isolated, monkeypatch)
    with pytest.raises(native_results.NativeIntegrityError, match="frozen"):
        native_results.prepare_reading(path, original_name="original.txt", expected_sha256="0" * 64)
    with store.connect() as c:
        damaged = ref.model_copy(update={"artifact_relpath": "native/../../outside"})
        payload = damaged.model_dump_json().encode()
        c.execute("UPDATE native_readings SET payload=?,payload_sha256=?", (payload.decode(), pipeline.digest(payload)))
    with pytest.raises(native_results.NativeIntegrityError, match="path"):
        native_results.load_reading(ref.reading_id)


def test_citation_identity_literal_quote_and_codepoint_position(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    pub = publish_runtime(ref)
    cite = native_results.machine_facts(pub)[0].native_citations[0]
    assert cite.quote_span.start == 2  # one emoji codepoint, then a space
    assert cite.quote_match_count == 1
    page = native_results.source_elements(pub, offset=0, limit=1)
    assert page.total == 1 and not page.has_more
    assert page.texts[cite.quote_span.text_sha256][cite.quote_span.start:cite.quote_span.end] == cite.quote
    for raw in ({"occurrence_id": "other", "element_id": cite.element_id, "quote": cite.quote},
                {"occurrence_id": cite.occurrence_id, "element_id": cite.element_id, "quote": "invented"}):
        with pytest.raises(native_results.NativeIntegrityError):
            native_results.resolve_citations(pub, [raw])


def test_repeated_quote_never_claims_unique_span(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch, text="0012 and 0012")
    pub = publish_runtime(ref, quote="0012")
    cite = native_results.machine_facts(pub)[0].native_citations[0]
    assert cite.quote_match_count == 2 and cite.quote_span is None


def test_version_fingerprint_refuses_changed_evidence(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    publish_runtime(ref)
    next((isolated / ref.artifact_relpath / "texts").iterdir()).write_bytes(b"changed")
    with store.connect() as c, pytest.raises(native_results.NativeIntegrityError):
        native_results.version_parts("native-run", c)


@pytest.mark.parametrize("extension", ["docx", "xlsx", "csv", "txt"])
def test_all_native_formats_keep_full_structured_source_elements(isolated, monkeypatch, extension):
    from io import BytesIO
    from jav.native_contracts import Succeeded
    from jav.readers.interpretation import ProposedExtraction, ground

    content = b"Order code: 0012\n"
    if extension == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("Order code: 0012")
        output = BytesIO()
        document.save(output)
        content = output.getvalue()
    elif extension == "xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active["A1"] = "0012"
        workbook.active["C1"] = "=1+2"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        content = output.getvalue()
    elif extension == "csv":
        content = b"Code,Blank,Value\n0012,,3\n"
    _, ref = runtime_source(isolated, monkeypatch, name="frozen." + extension, content=content)
    delivery = native_results.load_reading(ref.reading_id)
    interpretation = ground(delivery, ProposedExtraction(facts=()), provider="synthetic", model="synthetic-test", execution="synthetic_test")
    pub = native_results.publish(run_id="native-run", item_id="native-item", source_sha256=ref.source_sha256,
        recipe_hash="a" * 16, reading_id=ref.reading_id, outcome=Succeeded(), interpretation=interpretation)
    expected = [e for r in delivery.bundle.results for e in r.elements]
    collected, offset = [], 0
    while True:
        page = native_results.source_elements(pub, offset=offset, limit=1)
        assert page.total == len(expected) and page.result_version == pub.result_version
        collected.extend(page.elements)
        if not page.has_more:
            break
        offset = page.next_offset
    assert [e.model_dump(exclude={"occurrence_id", "attempt_id", "reader_key"}) for e in collected] == [e.model_dump() for e in expected]
    assert any("0012" in (e.text or "") for e in collected)
    if extension == "csv":
        assert any(e.cell is not None and e.cell.value.lexical == "" for e in collected)
        assert any(e.locator.kind == "cell" and e.locator.cell == "B2" for e in collected)
    if extension == "xlsx":
        assert any(e.kind == "sheet" for e in collected)
        assert any(e.cell is not None and e.cell.value.kind == "empty" for e in collected)
        assert any(e.cell is not None and e.cell.formula == "=1+2" for e in collected)


def test_canonical_bundle_envelope_cannot_change_after_publication(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    pub = publish_runtime(ref)
    path = isolated / ref.artifact_relpath / "bundle.json"
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(native_results.NativeIntegrityError, match="frozen envelope"):
        native_results.verify_publication(pub)


def test_foreign_run_or_approved_run_cannot_gain_publication(isolated, monkeypatch):
    _, ref = runtime_source(isolated, monkeypatch)
    with store.connect() as c:
        c.execute("UPDATE runs SET recipe_hash=? WHERE run_id='native-run'", ("f" * 16,))
    with pytest.raises(native_results.NativeIntegrityError, match="frozen work run"):
        publish_runtime(ref)
    with store.connect() as c:
        c.execute("UPDATE runs SET recipe_hash=?,approval='approved' WHERE run_id='native-run'", ("a" * 16,))
    with pytest.raises(native_results.PublicationConflict, match="approved"):
        publish_runtime(ref)
    assert native_results.get_publication("native-run", "native-item") is None
