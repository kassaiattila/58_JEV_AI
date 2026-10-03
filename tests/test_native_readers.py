"""Real, entirely synthetic files exercise the independent reader pipeline."""
from email.message import EmailMessage
from io import BytesIO
import zipfile
import os

import pytest

from jav.readers.pipeline import read_files

pytest.importorskip("docx", reason="Native reader dependencies await their shared integration window")
pytest.importorskip("defusedxml")
pytestmark = pytest.mark.skipif(os.name != "nt", reason="Windows Job Object reader boundary")



def workbook_bytes():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.append(["Code", "Empty", "Amount"])
    sheet.append(["0007", None, "=1+2"])
    sheet.merge_cells("A4:C4")
    sheet["A4"] = "Merged heading"
    sheet.row_dimensions[2].hidden = True
    book.create_sheet("Hidden").sheet_state = "hidden"
    book["Hidden"]["A1"] = "Private synthetic note"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def document_bytes():
    from docx import Document
    from PIL import Image
    picture = BytesIO()
    Image.new("RGB", (16, 12), "white").save(picture, format="PNG")
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "0007"
    document.add_paragraph("After table")
    document.add_picture(BytesIO(picture.getvalue()))
    document.sections[0].header.paragraphs[0].text = "Synthetic header"
    document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def read_one(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    return read_files([path])


def test_real_excel_preserves_empty_formula_and_hidden_state(tmp_path):
    delivery = read_one(tmp_path, "orders.xlsx", workbook_bytes())
    result = delivery.bundle.results[0]
    cells = {e.locator.cell: e for e in result.elements
             if e.kind == "cell" and e.locator.sheet == "Orders"}
    assert cells["A2"].cell.value.lexical == "0007"
    assert cells["B2"].cell.value.kind == "empty"
    assert cells["C2"].cell.formula == "=1+2"
    assert cells["C2"].cell.cached_state == "missing"
    assert cells["A4"].cell.merged_range == "A4:C4"
    assert cells["A2"].hidden
    assert any(e.kind == "sheet" and e.hidden for e in result.elements)
    assert result.attempt.execution == "reader"
    assert result.status == "partial"
    delivery.verify()


def test_real_word_order_headers_and_image_inventory(tmp_path):
    delivery = read_one(tmp_path, "document.docx", document_bytes())
    result = delivery.bundle.results[0]
    roots = [e for e in result.elements if e.parent_id is None]
    assert [e.kind for e in roots[:3]] == ["paragraph", "table", "paragraph"]
    assert roots[0].text == "Before table"
    assert roots[2].text == "After table"
    assert {"Synthetic header", "Synthetic footer"} <= {e.text for e in result.elements}
    assert any(e.kind == "image" for e in result.elements)
    assert any(o.role == "embedded" for o in delivery.bundle.manifest.occurrences)
    assert any(i.code == "needs_ocr" for i in result.issues)
    delivery.verify()


def test_same_file_from_folder_and_email_shares_bytes_not_occurrence(tmp_path):
    data = workbook_bytes()
    (tmp_path / "orders.xlsx").write_bytes(data)
    message = EmailMessage()
    message["From"] = "sender@example.invalid"
    message["To"] = "recipient@example.invalid"
    message["Subject"] = "Synthetic order"
    message.set_content("Please read the attached order.")
    message.add_attachment(data, maintype="application", subtype="octet-stream", filename="orders.xlsx")
    message.add_attachment(b"unknown payload", maintype="application", subtype="octet-stream", filename="unknown.bin")
    (tmp_path / "message.eml").write_bytes(message.as_bytes())
    delivery = read_files([tmp_path])
    occurrences = delivery.bundle.manifest.occurrences
    books = [o for o in occurrences if o.original_name == "orders.xlsx"]
    assert len(books) == 2 and books[0].occurrence_id != books[1].occurrence_id
    assert books[0].object_sha256 == books[1].object_sha256
    email = next(o for o in occurrences if o.role == "email")
    inventory = next(i for i in delivery.bundle.manifest.inventories if i.occurrence_id == email.occurrence_id)
    assert inventory.expected_children == 2
    assert any(r.status == "unsupported" for r in delivery.bundle.results)
    assert delivery.reading_invocations[books[0].object_sha256] == 1
    delivery.verify()


@pytest.mark.parametrize("name,data", [
    ("empty.txt", b""), ("broken.docx", b"PK\x03\x04broken"),
    ("unsupported.bin", b"\x00\xff"),
])
def test_unreadable_files_remain_visible(tmp_path, name, data):
    delivery = read_one(tmp_path, name, data)
    assert len(delivery.bundle.manifest.occurrences) == 1
    assert delivery.bundle.results[0].status not in {"complete", "partial"}
    assert delivery.bundle.results[0].issues


def test_archive_path_escape_is_excluded(tmp_path):
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as package:
        package.writestr("../escaped.txt", "untrusted")
        package.writestr("word/document.xml", "<root/>")
    delivery = read_one(tmp_path, "unsafe.docx", output.getvalue())
    assert delivery.bundle.results[0].status == "excluded"
    assert not (tmp_path.parent / "escaped.txt").exists()


def test_limited_cells_fail_visibly(tmp_path):
    from jav.readers.limits import DEFAULT_LIMITS
    path = tmp_path / "orders.xlsx"
    path.write_bytes(workbook_bytes())
    delivery = read_files([path], limits=DEFAULT_LIMITS.model_copy(update={"visited_cells": 2}))
    assert delivery.bundle.results[0].status == "resource_limited"


def test_frozen_delivery_cannot_be_overwritten_or_reparsed(tmp_path):
    delivery = read_one(tmp_path, "data.txt", b"Code: 0007\nQuantity: 3")
    output = tmp_path / "evidence"
    delivery.save(output)
    from jav.readers.pipeline import Delivery
    loaded = Delivery.load(output)
    assert loaded.bundle == delivery.bundle
    object_file = next((output / "objects").iterdir())
    object_file.write_bytes(b"tampered")
    with pytest.raises(ValueError):
        Delivery.load(output)


def test_symlink_outside_folder_is_not_read(tmp_path):
    folder = tmp_path / "incoming"
    folder.mkdir()
    target = tmp_path / "outside.txt"
    target.write_text("do not read", encoding="utf-8")
    try:
        (folder / "link.txt").symlink_to(target)
    except OSError:
        pytest.skip("Host does not allow creating a synthetic symbolic link")
    delivery = read_files([folder])
    assert delivery.bundle.manifest.occurrences[0].acquisition == "excluded"
    assert not delivery.bundle.manifest.objects


def test_loaded_bundle_cannot_change_values_without_changing_evidence(tmp_path):
    import json
    from jav.readers.pipeline import Delivery
    delivery = read_one(tmp_path, "input.txt", b"Literal source")
    output = tmp_path / "saved"
    delivery.save(output)
    path = output / "bundle.json"
    bundle = json.loads(path.read_bytes())
    bundle["results"][0]["elements"][0]["text"] = "Invented source"
    path.write_text(json.dumps(bundle), encoding="utf-8")
    with pytest.raises(ValueError, match="frozen evidence"):
        Delivery.load(output)


def test_saved_metadata_cannot_request_unbounded_object_loading(tmp_path):
    import json
    from jav.readers.pipeline import Delivery
    delivery = read_one(tmp_path, "input.txt", b"Literal source")
    output = tmp_path / "saved"
    delivery.save(output)
    path = output / "bundle.json"
    bundle = json.loads(path.read_bytes())
    bundle["manifest"]["objects"][0]["byte_size"] = 10**12
    path.write_text(json.dumps(bundle), encoding="utf-8")
    with pytest.raises(ValueError, match="loading bounds"):
        Delivery.load(output)


def test_large_embedded_image_does_not_consume_structural_evidence_budget(tmp_path):
    from docx import Document
    from PIL import Image
    image = BytesIO()
    Image.effect_noise((1800, 1800), 100).save(image, format="PNG")
    document = Document()
    document.add_paragraph("Text beside large embedded picture")
    document.add_picture(BytesIO(image.getvalue()))
    source = BytesIO()
    document.save(source)
    delivery = read_one(tmp_path, "large-picture.docx", source.getvalue())
    result = delivery.bundle.results[0]
    assert result.status == "partial"
    assert any(e.text == "Text beside large embedded picture" for e in result.elements)
    assert result.raw_evidence.byte_size < 100_000
    assert len(delivery.bundle.manifest.occurrences) == 2
    delivery.verify()
