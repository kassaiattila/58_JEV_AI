"""178 (plan 170 L4, DECISIONS 178): the output layer - the shared Excel writer and the output contract.

The contract (`configs/outputs.json`) lists what each output is filled with; the builders emit only those names and the
built-in layout refuses anything else, so a template can rely on the list. Found while pinning the downloads (K20: a small visible fault found during the refactoring is fixed in its own commit):
the reconciliation statement styled the empty row below its heading block as the table's header, froze the rows above
the real header and filtered from the empty row, so the header scrolled away and the last row was outside the filter.
Synthetic values only.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from jav import export, typepack
from jav.core import store
from jav.domains.finance import reconcile_statement as rs
from jav.outputs import builtin, contract, xlsx
from tests import test_characterisation_178 as c178


def test_a_table_below_a_heading_block_has_its_own_header_frozen_and_filtered():
    wb = Workbook()
    ws = wb.active
    xlsx.write_table(ws, ["Number", "Amount"], [["A-1", "10.00"], ["A-2", "20.00"], ["A-3", "30.00"]], [None, "money"],
                     heading=[("Package", "April"), ("Period", "2026-04-01 – 2026-04-30")], min_width=10)
    assert [ws.cell(r, 1).value for r in range(1, 8)] == ["Package", "Period", None, "Number", "A-1", "A-2", "A-3"]
    assert ws.cell(4, 1).font.b and ws.cell(4, 1).fill.fgColor.rgb == "00EEEEF2"  # the header row
    assert not ws.cell(3, 1).fill.fill_type  # the empty row between stays plain
    assert ws.freeze_panes == "A5"  # the header stays in view
    assert ws.auto_filter.ref == "A4:B7"  # from the header to the last row


def test_a_table_without_a_heading_starts_on_the_first_row():
    wb = Workbook()
    ws = wb.active
    xlsx.write_table(ws, ["Number"], [["A-1"]])
    assert ws.cell(1, 1).font.b and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:A2"


# --- the output contract --------------------------------------------------------------------------------------------


@pytest.fixture
def db(tmp_path):
    with store.use_store(tmp_path / "o178.sqlite"):
        yield


def test_every_name_of_the_contract_is_an_identifier_and_every_kind_is_known():
    outputs = contract.outputs()
    assert set(outputs) == {"package_result", "utility_cost", "reconcile_statement", "quote"}  # 196: the quote
    for name, out in outputs.items():
        assert out["format"] in ("xlsx", "docx") and out["what"]  # 196: a Word output
        layout = {"version", "heading", "min_width", "marks"} if out["format"] == "xlsx" else {"version"}
        assert layout <= set(out["builtin"])
        for key, kind in out["values"].items():
            assert contract.KEY.match(key) and kind in contract.KINDS, (name, key)
        for table, spec in out["tables"].items():
            assert contract.KEY.match(table), (name, table)
            for key, kind in spec["columns"].items():
                assert contract.KEY.match(key) and kind in contract.KINDS, (name, table, key)


def test_a_type_packs_fields_never_collide_with_the_fixed_columns():
    tables = contract.spec("package_result")["tables"]
    fixed_documents, fixed_lines = set(tables["documents"]["columns"]), set(tables["line_items"]["columns"])
    for key in typepack.keys():
        pack = typepack.get(key)
        assert not fixed_documents & set(pack.fields), key
        assert not fixed_lines & {c for lf in pack.list_fields.values() for c in lf}, key


def test_the_builders_emit_only_what_the_contract_lists(db):
    data = export.result_data(c178._records(), {"emails": c178._mails()}, [])
    contract.check(data)
    by_key = {(t.key, t.part): t for t in data.tables}
    assert list(by_key) == [("emails", None), ("tasks", None), ("documents", "invoice_hu"), ("documents", "villamos_energia_szamla"),
                            ("datapoints", None), ("line_items", None), ("utility_cost", None), ("utility_sources", None)]
    assert by_key[("documents", "invoice_hu")].columns[10:] == typepack.get("invoice_hu").header_fields
    assert by_key[("utility_cost", None)].columns[5:] == ("month_2026_01", "month_2026_02", "month_2026_03", "month_2026_04")
    assert set(by_key[("utility_cost", None)].marks.values()) == {"missing", "overlap"}
    wp_id, _inv, _ids = c178._world()
    statement = rs.statement_data(rs.current(wp_id)["statement"], "draft")
    contract.check(statement)
    assert [v.key for v in statement.values] == ["package", "party", "period", "accounts", "state"]
    assert [t.key for t in statement.tables] == ["invoices", "lines"]


def _data(**over):
    table = contract.Table(key="invoices", title="Invoices", columns=("number", "amount"), head=("Number", "Amount"),
                           kinds=(None, "money"), rows=[["A-1", "10.00"]])
    fields = {"output": "reconcile_statement", "values": (contract.Value("package", "Package", "April"),), "tables": (table,)}
    return contract.OutputData(**{**fields, **over})


def test_a_name_the_contract_does_not_list_is_refused():
    builtin.render(_data())  # the listed names write
    unknown = contract.Table(key="invoices", title="Invoices", columns=("number", "secret"), head=("Number", "Secret"),
                             kinds=(None, None), rows=[])
    with pytest.raises(contract.ContractError, match="column invoices.secret"):
        builtin.render(_data(tables=(unknown,)))
    with pytest.raises(contract.ContractError, match="value owner"):
        contract.check(_data(values=(contract.Value("owner", "Owner", "x"),)))
    with pytest.raises(contract.ContractError, match="table payments"):
        contract.check(_data(tables=(contract.Table(key="payments", title="P", columns=(), head=(), kinds=(), rows=[]),)))
    with pytest.raises(contract.ContractError, match="unknown output"):
        contract.check(_data(output="invoice_letter"))


def test_a_column_whose_kind_disagrees_with_the_contract_is_refused():
    as_text = contract.Table(key="invoices", title="Invoices", columns=("number", "amount"), head=("Number", "Amount"),
                             kinds=(None, None), rows=[])
    with pytest.raises(contract.ContractError, match="column invoices.amount: kind None, the contract's money"):
        contract.check(_data(tables=(as_text,)))


def test_a_table_per_document_type_takes_only_its_own_type_packs_fields():
    out = "package_result"
    own = contract.Table(key="documents", title="invoice_hu", columns=("file", "gross_total"), head=("File", "Gross"),
                         kinds=(None, None), rows=[], part="invoice_hu")
    contract.check(contract.OutputData(out, (), (own,)))
    other = contract.Table(key="documents", title="invoice_hu", columns=("file", "consumption_kwh"), head=("Irat", "kWh"),
                           kinds=(None, None), rows=[], part="invoice_hu")
    with pytest.raises(contract.ContractError, match="column documents.consumption_kwh"):
        contract.check(contract.OutputData(out, (), (other,)))
    without_type = contract.Table(key="documents", title="x", columns=("file",), head=("Irat",), kinds=(None,), rows=[])
    with pytest.raises(contract.ContractError, match="part None"):
        contract.check(contract.OutputData(out, (), (without_type,)))
