"""196 (plan F-Word-output W2, DECISIONS 196): the approved quote as a Word document.

The owner's decisions of 2026-10-22: a built-in template in Hungarian and English, the language estimated by code from
the request and changed by a person, no PDF. The approval stores the document with its version; a person's own Word
template replaces the built-in one, chosen per package. Synthetic requests and documents; a stand-in JEV and match; no
paid call.
"""

from __future__ import annotations

import copy
import io
import json
from pathlib import Path

import pytest
from docx import Document

from jav import package_release
from jav.core import store
from jav.domains.quotes import quote_document, quote_draft, invented_golden
from jav.outputs import builtin, docx, templates
from tests import test_quote_draft_192 as q192

env = q192.env
quote = q192.quote


def _content(language: str = "hu", **header) -> dict:
    head = {"customer_name": "Példa Kft.", "customer_tax_number": None, "contact_name": "Minta Anna",
            "contact_email": "anna@pelda.hu", "contact_phone": None, "deadline": "2026-11-05", "reference": None,
            "note": None, **header}
    return {"header": head, "language": language, "currency": "HUF", "vat_percent": "27",
            "lines": [{"sku": "HW-MON-24", "name": "24 colos monitor", "quantity": "3", "unit": "db", "list_price": "54900",
                       "unit_price": "54900", "discount_percent": "0", "net": "164700"}],
            "totals": {"net": "164700", "vat": "44469", "gross": "209169"},
            "price_list": {"name": "IT 2026", "version": 2}, "request": {"title": "Ajánlatkérés"}}


def _text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells + [doc.sections[0].footer.paragraphs[0].text])


# --- the language ---------------------------------------------------------------------------------------------------


def test_the_language_of_every_invented_request_is_estimated_by_code():
    golden = invented_golden.load()
    for case in golden.cases:
        items = case.expected.request.items if case.expected.request else ()
        texts = [case.message.subject, *(i.written for i in items)]
        if not items:  # a decoy has no items: its body stands in for them
            texts.append(case.message.body)
        assert quote_document.language_of(texts) == case.language, case.id


def test_nothing_to_go_by_is_hungarian():
    assert quote_document.language_of([]) == "hu" and quote_document.language_of(["HW-MON-24", "3"]) == "hu"


# --- the data and the built-in template ----------------------------------------------------------------------------


def test_the_data_is_the_approved_quotes_and_an_unpriced_line_is_refused():
    d = quote_document.data(_content(), package="Office", version=2, approved_by="anna", approved_at="2026-10-22T09:00:00Z")
    values = {v.key: v.value for v in d.values}
    assert values["version"] == 2 and values["state"] is None and values["quote_date"] == "2026-10-22"
    assert values["gross_total"] == "209169" and values["price_list"] == "IT 2026 (2)"
    assert d.tables[0].rows == [[1, "HW-MON-24", "24 colos monitor", "3", "db", "54900", "54900", "0", "164700"]]
    draft = {v.key: v.value for v in quote_document.data(_content(), package="Office").values}
    assert draft["state"] == "Piszkozat" and draft["version"] is None
    unpriced = _content()
    unpriced["lines"][0].pop("net")
    with pytest.raises(quote_document.QuoteDocumentError):
        quote_document.data(unpriced, package="Office")


def test_the_built_in_template_writes_hungarian_or_english_and_leaves_out_an_empty_line():
    hu = _text(builtin.render(quote_document.data(_content(), package="Office", version=1, approved_at="2026-10-22")))
    assert "Árajánlat" in hu and "Ajánlatkérő: Példa Kft." in hu and "Kért határidő: 2026. 11. 05." in hu
    assert "Bruttó összesen: 209 169 HUF" in hu and "ÁFA (27 %): 44 469 HUF" in hu
    assert "Adószám" not in hu and "Hivatkozás" not in hu and "Piszkozat" not in hu
    assert hu.endswith("Office · 1. változat · . oldal")  # the page number is a field, outside the paragraph's text
    en = _text(builtin.render(quote_document.data(_content("en", reference="RFQ-7"), package="Office")))
    assert "Quotation – Draft" in en and "Reference: RFQ-7" in en and "Gross total: 209,169 HUF" in en
    assert "Requested by: 5 November 2026" in en and "54,900" in en


def test_the_built_in_template_is_a_word_template_of_the_output():
    data = quote_document.builtin_template(quote_document.data(_content(note="Első\nMásodik"), package="Office"))
    check = docx.check(quote_document.OUTPUT, data)
    assert check.ok, check.problems
    assert "lines" in check.tables and "gross_total" in check.values


def test_the_file_name_names_the_quote_its_package_and_its_version():
    assert quote_document.file_name(_content(), package="Példa Kft. ajánlat", version=3) == "arajanlat-Pelda-Kft-ajanlat-v3.docx"
    assert quote_document.file_name(_content("en"), package="Office", version=None) == "quotation-Office-draft.docx"


# --- the approval stores it, the download gives it -------------------------------------------------------------------


def test_the_approval_stores_the_document_and_the_download_never_changes(quote):
    wp_id = quote["wp"]
    v = q192._done(wp_id)
    assert v["language"] == "hu" and v["languages"] == ["hu", "en"]
    draft, name = quote_draft.draft_docx(wp_id)
    assert name.endswith("-piszkozat.docx") and "Árajánlat – Piszkozat" in _text(draft)
    quote_draft.approve(wp_id, fingerprint=v["fingerprint"], actor="anna")
    files = package_release.own_files(wp_id, 1)
    assert [(f["state"], f["template"]) for f in files] == [("approved", "builtin")]
    stored, name = quote_draft.version_docx(wp_id, 1)
    assert name == "arajanlat-Office-v1.docx" and "1 228 979 HUF" in _text(stored)
    kept = quote_draft.approved_version(wp_id, 1)["quote"]
    assert kept["language"] == "hu"
    after = quote_draft.reopen(wp_id, note="Another monitor", actor="anna")
    assert quote_draft.version_docx(wp_id, 1)[0] == stored  # a reopened version keeps the document it was approved with
    assert after["approval"]["active"] is None


def test_a_person_changes_the_language_and_the_approval_records_it(quote):
    wp_id = quote["wp"]
    v = q192._done(wp_id)
    with pytest.raises(quote_draft.QuoteDraftError, match="language"):
        quote_draft.save(wp_id, q192._sent(v, language="de"), actor="anna", expected_state=v["state"])
    v2 = quote_draft.save(wp_id, q192._sent(v, language="en"), actor="anna", expected_state=v["state"])
    assert v2["language"] == "en" and v2["fingerprint"] != v["fingerprint"]
    quote_draft.approve(wp_id, fingerprint=v2["fingerprint"], actor="anna")
    assert "Quotation" in _text(quote_draft.version_docx(wp_id, 1)[0])


def test_a_version_approved_before_196_gets_its_document_from_the_built_in_template(quote):
    wp_id = quote["wp"]
    v = q192._done(wp_id)
    quote_draft.approve(wp_id, fingerprint=v["fingerprint"], actor="anna")
    with store.connect() as c:
        c.execute("DELETE FROM package_version_files")
        row = c.execute("SELECT id, content FROM package_versions WHERE workpackage_id=?", (wp_id,)).fetchone()
        content = json.loads(row["content"])
        content.pop("language")
        c.execute("UPDATE package_versions SET content=? WHERE id=?", (json.dumps(content), row["id"]))
    made, name = quote_draft.version_docx(wp_id, 1)
    assert name == "arajanlat-Office-v1.docx" and "Árajánlat" in _text(made)


def test_an_own_word_template_chosen_for_the_package_makes_the_stored_document(quote, tmp_path):
    wp_id = quote["wp"]
    doc = Document()
    doc.add_paragraph("Our offer to {{customer_name}}: {{gross_total}} {{currency}}")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "{{lines.name}}"
    path = tmp_path / "own.docx"
    doc.save(path)
    made = templates.add("quote", "Own offer", path.name, path.read_bytes(), actor="anna")
    templates.activate(made["template"], 1, actor="anna")
    templates.choose(wp_id, "quote", made["template"], actor="anna")
    preview, preview_name = templates.preview(made["template"], 1)
    assert preview_name == "preview-own.docx" and "Our offer to" in _text(preview)
    v = q192._done(wp_id)
    quote_draft.approve(wp_id, fingerprint=v["fingerprint"], actor="anna")
    assert package_release.own_files(wp_id, 1)[0]["template"] == made["template"]
    text = _text(quote_draft.version_docx(wp_id, 1)[0])
    assert "Our offer to " in text and "1 228 979 HUF" in text and "Monitor 24 colos" in text


def test_the_service_gives_the_draft_and_the_versions_document(quote, tmp_path):
    from fastapi.testclient import TestClient

    from jav.api import create_app

    wp_id = quote["wp"]
    human = {"X-Actor": "teszt.elek"}
    with TestClient(create_app(store_path=Path(tmp_path) / "q.sqlite"), base_url="http://127.0.0.1:8930") as c:
        assert c.get(f"/api/workpackages/{wp_id}/quote/draft/quote.docx").status_code == 422  # no draft yet
        v = q192._done(wp_id)
        got = c.get(f"/api/workpackages/{wp_id}/quote/draft/quote.docx")
        assert got.status_code == 200 and got.headers["content-type"] == package_release.DOCX
        assert "piszkozat.docx" in got.headers["content-disposition"]
        c.post(f"/api/workpackages/{wp_id}/quote/approve", headers=human, json={"fingerprint": v["fingerprint"]})
        version = c.get(f"/api/workpackages/{wp_id}/quote/versions/1/quote.docx")
        assert version.status_code == 200 and "Árajánlat" in _text(version.content)
        assert c.get(f"/api/workpackages/{wp_id}/quote/versions/9/quote.docx").status_code == 404
        choices = c.get(f"/api/workpackages/{wp_id}/output-templates").json()["outputs"]
        assert list(choices) == ["quote"]


def test_a_draft_with_a_line_not_priced_has_no_document(quote):
    v = quote_draft.start(quote["wp"], source="request", actor="anna", expected_state=0)
    assert v["totals"] is None
    with pytest.raises(quote_draft.QuoteDraftError, match="not priced"):
        quote_draft.draft_docx(quote["wp"])
    assert copy.deepcopy(v)["language"] == "hu"
