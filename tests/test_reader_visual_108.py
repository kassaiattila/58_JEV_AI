"""Frozen PDF/image evidence, using synthetic sources and local recognition only."""
from io import BytesIO
import json
import os
import re

from PIL import Image, ImageDraw, ImageFont
import pytest

from jav.readers import pipeline
from jav.readers.limits import DEFAULT_LIMITS, ReadFailure
from tests.pdfgen import write_text_pdf

pytestmark = pytest.mark.skipif(os.name != "nt", reason="Measured Windows reader boundary")


def picture_bytes(*, frames=1, format="PNG"):
    pictures = [Image.new("RGB", (600, 200), "white") for _ in range(frames)]
    for index, picture in enumerate(pictures):
        ImageDraw.Draw(picture).text((40, 60), f"SYNTHETIC {index + 1}", fill="black",
                                    font=ImageFont.load_default(size=40))
    out = BytesIO()
    if frames > 1:
        pictures[0].save(out, format=format, save_all=True, append_images=pictures[1:])
    else:
        pictures[0].save(out, format=format)
    return out.getvalue()


def tsv(text="SYNTHETIC", left=40, top=60, width=220, height=42):
    return ("level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
            f"5\t1\t1\t1\t1\t1\t{left}\t{top}\t{width}\t{height}\t95\t{text}\n")


@pytest.fixture
def fake_ocr(monkeypatch):
    from jav.readers import visual_ocr
    calls = []

    class FakeOCR:
        fingerprint = "a" * 64

        def recognise(self, png, limits, *, timeout):
            assert timeout > 0
            assert png.startswith(b"\x89PNG")
            calls.append(png)
            return tsv(f"SYNTHETIC-{len(calls)}")

    monkeypatch.setattr(visual_ocr.LocalOCR, "discover", lambda: FakeOCR())
    return calls


def write_image(tmp_path, name="sample.png", **kwargs):
    path = tmp_path / name
    path.write_bytes(picture_bytes(**kwargs))
    return path


def raw_result(delivery, index=0):
    result = delivery.bundle.results[index]
    return json.loads(delivery.evidence[result.raw_evidence.sha256])


def test_native_pdf_has_frozen_word_layer_and_exact_line_references(tmp_path):
    path = write_text_pdf(tmp_path / "text.pdf", ["Order SYN-001", "Total 125.00 EUR"])
    delivery = pipeline.read_files([path])
    result = delivery.bundle.results[0]
    assert result.status == "complete"
    assert result.attempt.parser_name == "native-pdf"
    raw = raw_result(delivery)
    assert raw["source_layers"]
    for element in result.elements:
        assert element.locator.kind == "pdf"
        layer = raw["source_layers"][element.locator.source_layer_id]
        assert layer["doc_id"] == result.attempt.source_sha256
        selected = [w for w in layer["words"] if w["id"] in element.locator.word_ids]
        assert " ".join(w["text"] for w in selected) == element.text
        assert all(0 <= w["x0"] < w["x1"] <= 1 for w in selected)


def test_image_recognition_is_bound_to_original_pixels_and_raw_tsv(tmp_path, fake_ocr):
    path = write_image(tmp_path)
    delivery = pipeline.read_files([path], ocr=True)
    result = delivery.bundle.results[0]
    assert len(fake_ocr) == 1
    assert result.status == "partial"
    assert any(issue.code == "protection_unavailable" for issue in result.issues)
    assert result.attempt.protections.network == result.attempt.protections.paths == "unavailable"
    words = [e for e in result.elements if e.text]
    assert len(words) == 1
    assert words[0].text == "SYNTHETIC-1"
    assert words[0].locator.region == (40, 60, 260, 102)
    assert words[0].locator.image_sha256 == pipeline.digest(path.read_bytes())
    assert result.attempt.models_sha256 == "a" * 64
    raw = raw_result(delivery)
    assert raw["ocr"][0]["tsv"] == tsv("SYNTHETIC-1")
    assert raw["rasters"][0]["sha256"] in delivery.rasters


def test_multiframe_image_preserves_frame_identity(tmp_path, fake_ocr):
    path = write_image(tmp_path, "frames.tiff", frames=2, format="TIFF")
    result = pipeline.read_files([path], ocr=True).bundle.results[0]
    assert len(fake_ocr) == 2
    assert [(e.locator.frame, e.text) for e in result.elements if e.text] == [
        (1, "SYNTHETIC-1"), (2, "SYNTHETIC-2")]


def test_saved_evidence_reopens_without_recognition_and_detects_raster_tampering(tmp_path, fake_ocr, monkeypatch):
    path = write_image(tmp_path)
    delivery = pipeline.read_files([path], ocr=True)
    destination = tmp_path / "delivery"
    delivery.save(destination)
    monkeypatch.setattr(pipeline, "run", lambda *a, **k: pytest.fail("Reopening invoked the reader"))
    loaded = pipeline.Delivery.load(destination)
    assert loaded.bundle == delivery.bundle
    assert len(fake_ocr) == 1
    raster = next((destination / "rasters").iterdir())
    raster.write_bytes(b"changed synthetic raster")
    with pytest.raises(ValueError):
        pipeline.Delivery.load(destination)


def test_select_keeps_only_referenced_frozen_rasters(tmp_path, fake_ocr):
    path = write_image(tmp_path)
    delivery = pipeline.read_files([path], ocr=True)
    selected = delivery.select({delivery.bundle.manifest.occurrences[0].occurrence_id})
    assert selected.rasters == delivery.rasters
    selected.verify()


def test_identical_inputs_share_one_reading_and_one_recognition(tmp_path, fake_ocr):
    path = write_image(tmp_path)
    delivery = pipeline.read_files([path, path], ocr=True)
    assert len(delivery.bundle.manifest.occurrences) == 2
    assert len(fake_ocr) == 1
    assert list(delivery.reading_invocations.values()) == [1]


def test_ocr_timeout_retains_original_and_explicit_unread_frame(tmp_path, monkeypatch):
    from jav.readers import visual_ocr

    class FailingOCR:
        fingerprint = "b" * 64

        def recognise(self, png, limits, *, timeout):
            raise ReadFailure("resource_limited", "resource_limit", "Synthetic recognition timeout")

    monkeypatch.setattr(visual_ocr.LocalOCR, "discover", lambda: FailingOCR())
    path = write_image(tmp_path)
    delivery = pipeline.read_files([path], ocr=True)
    result = delivery.bundle.results[0]
    assert result.status == "partial"
    assert any(i.code == "resource_limit" for i in result.issues)
    assert delivery.objects[pipeline.digest(path.read_bytes())] == path.read_bytes()
    assert not any(e.text for e in result.elements)


def test_missing_local_ocr_does_not_use_remote_provider(tmp_path, monkeypatch):
    from jav.readers import visual_ocr

    def missing():
        raise ReadFailure("unsupported", "needs_ocr", "Local OCR is not installed")

    monkeypatch.setattr(visual_ocr.LocalOCR, "discover", missing)
    result = pipeline.read_files([write_image(tmp_path)], ocr=True).bundle.results[0]
    assert result.status == "partial"
    assert any(i.code == "needs_ocr" for i in result.issues)


def test_ocr_out_of_bounds_coordinates_are_rejected(tmp_path, monkeypatch):
    from jav.readers import visual_ocr

    class InvalidOCR:
        fingerprint = "c" * 64

        def recognise(self, png, limits, *, timeout):
            return tsv(left=590, width=200)

    monkeypatch.setattr(visual_ocr.LocalOCR, "discover", lambda: InvalidOCR())
    result = pipeline.read_files([write_image(tmp_path)], ocr=True).bundle.results[0]
    assert result.status == "partial"
    assert not any(e.text for e in result.elements)
    assert any(i.code == "corrupt" for i in result.issues)


def test_scanned_pdf_maps_ocr_pixels_to_page_coordinates(tmp_path, fake_ocr):
    path = tmp_path / "scan.pdf"
    with Image.open(BytesIO(picture_bytes())) as image:
        image.save(path, "PDF", resolution=150)
    delivery = pipeline.read_files([path], ocr=True)
    result = delivery.bundle.results[0]
    assert len(fake_ocr) == 1
    text = next(e for e in result.elements if e.text)
    assert text.locator.kind == "pdf"
    layer = raw_result(delivery)["source_layers"][text.locator.source_layer_id]
    word = layer["words"][text.locator.word_ids[0]]
    raster = raw_result(delivery)["rasters"][0]
    assert word["x0"] == pytest.approx(40 / raster["width"], abs=0.00001)
    assert word["y0"] == pytest.approx(60 / raster["height"], abs=0.00001)


def test_pdf_pixels_are_bounded_before_rendering(tmp_path):
    path = tmp_path / "scan.pdf"
    with Image.open(BytesIO(picture_bytes())) as image:
        image.save(path, "PDF", resolution=150)
    limits = DEFAULT_LIMITS.model_copy(update={"image_pixels": 20})
    result = pipeline.read_files([path], limits=limits, ocr=True).bundle.results[0]
    assert result.status == "resource_limited"


def test_real_local_recognition_returns_text_and_a_saved_region(tmp_path):
    from jav.readers.visual_ocr import LocalOCR
    LocalOCR.discover()  # Missing installation must remain visible in this host acceptance test.
    result = pipeline.read_files([write_image(tmp_path)], ocr=True).bundle.results[0]
    text = " ".join(e.text for e in result.elements if e.text)
    assert "SYNTHETIC" in text
    regions = [e.locator.region for e in result.elements if e.text]
    assert regions and all(0 <= box[0] < box[2] <= 600 and 0 <= box[1] < box[3] <= 200 for box in regions)


def test_mixed_pdf_keeps_native_words_when_visual_recognition_fails(tmp_path, monkeypatch):
    import pypdfium2 as pdfium
    from jav.readers import visual_ocr

    native = write_text_pdf(tmp_path / "native.pdf", ["Order SYN-002", "Total 125.00 EUR"])
    scanned = tmp_path / "scanned.pdf"
    with Image.open(BytesIO(picture_bytes())) as image:
        image.save(scanned, "PDF", resolution=150)
    combined = tmp_path / "combined.pdf"
    with pdfium.PdfDocument.new() as document:
        for path in (native, scanned):
            with pdfium.PdfDocument(path) as source:
                document.import_pages(source)
        document.save(combined)

    def unavailable():
        raise ReadFailure("unsupported", "needs_ocr", "Synthetic missing local engine")

    monkeypatch.setattr(visual_ocr.LocalOCR, "discover", unavailable)
    result = pipeline.read_files([combined], ocr=True).bundle.results[0]
    assert result.status == "partial"
    assert any("SYN-002" in (e.text or "") and e.locator.page == 1 for e in result.elements)
    assert any(e.locator.page == 2 and e.availability == "unreadable" for e in result.elements)


@pytest.mark.parametrize("rotation", [90, 180, 270])
def test_rotated_scanned_pdf_uses_actual_rotated_raster_dimensions(tmp_path, fake_ocr, rotation):
    import pypdfium2 as pdfium
    source = tmp_path / "scan.pdf"
    with Image.open(BytesIO(picture_bytes())) as image:
        image.save(source, "PDF", resolution=150)
    rotated = tmp_path / "rotated.pdf"
    with pdfium.PdfDocument(source) as document:
        page = document[0]
        page.set_rotation(rotation)
        page.close()
        document.save(rotated)
    delivery = pipeline.read_files([rotated], ocr=True)
    raw = raw_result(delivery)
    assert raw["status"] == "partial", raw["issues"]
    assert [issue["code"] for issue in raw["issues"]] == ["protection_unavailable"]
    raster = raw["rasters"][0]
    assert (raster["height"] > raster["width"]) == (rotation in (90, 270))
    text = next(e for e in delivery.bundle.results[0].elements if e.text)
    layer = raw["source_layers"][text.locator.source_layer_id]
    assert layer["words"][0]["x0"] == pytest.approx(40 / raster["width"], abs=0.00001)


def test_blank_recognition_stays_partial(tmp_path, monkeypatch):
    from jav.readers import visual_ocr

    class EmptyOCR:
        fingerprint = "d" * 64

        def recognise(self, png, limits, *, timeout):
            return tsv().splitlines()[0] + "\n"

    monkeypatch.setattr(visual_ocr.LocalOCR, "discover", lambda: EmptyOCR())
    result = pipeline.read_files([write_image(tmp_path)], ocr=True).bundle.results[0]
    assert result.status == "partial"
    assert any("blankness" in i.message for i in result.issues)


def test_pdf_saved_word_reference_cannot_name_an_absent_word(tmp_path):
    from jav.readers.visual import verify_layers
    delivery = pipeline.read_files([write_text_pdf(tmp_path / "source.pdf", ["Order SYN-003"])])
    raw = raw_result(delivery)
    raw["elements"][0]["locator"]["word_ids"] = [99999]
    with pytest.raises(ValueError, match="invalid words"):
        verify_layers(raw, delivery.bundle.results[0].attempt.source_sha256)


def test_corrupt_pdf_keeps_a_named_failure_and_original_bytes(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4\nnot a valid synthetic PDF")
    delivery = pipeline.read_files([path])
    assert delivery.bundle.results[0].status == "corrupt"
    assert delivery.objects[pipeline.digest(path.read_bytes())] == path.read_bytes()


def pdf_with_catalog_entry(tmp_path, entry):
    path = write_text_pdf(tmp_path / "catalog.pdf", ["Synthetic PDF content"])
    bodies = re.findall(rb"[0-9]+ 0 obj\n(.*?)\nendobj", path.read_bytes(), re.DOTALL)
    bodies[0] = bodies[0].replace(b"/Pages 2 0 R", b"/Pages 2 0 R " + entry)
    output, offsets = bytearray(b"%PDF-1.4\n"), []
    for index, body in enumerate(bodies, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(bodies) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(bodies) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(output)
    return path


@pytest.mark.parametrize("entry", [
    b"/OpenAction << /S /JavaScript /JS (synthetic) >>",
    b"/Open#41ction << /S /Java#53cript /JS (synthetic) >>",
    b"/AcroForm << /Fields [] >>",
    b"/Names << /EmbeddedFiles << /Names [] >> >>",
])
def test_active_or_unrepresented_pdf_structures_are_excluded(tmp_path, entry):
    result = pipeline.read_files([pdf_with_catalog_entry(tmp_path, entry)]).bundle.results[0]
    assert result.status == "excluded"
    assert not result.elements


def test_embedded_word_image_uses_child_pixels_without_invented_office_page(tmp_path, fake_ocr):
    from docx import Document
    document = Document()
    document.add_paragraph("Synthetic image host")
    document.add_picture(BytesIO(picture_bytes()))
    path = tmp_path / "image.docx"
    document.save(path)
    delivery = pipeline.read_files([path], ocr=True)
    parent, child = delivery.bundle.results
    image = next(e for e in parent.elements if e.kind == "image")
    assert image.locator.host.kind == "word"
    text = next(e for e in child.elements if e.text)
    assert text.locator.image_sha256 == image.locator.image_sha256
    assert text.locator.region == (40, 60, 260, 102)
    assert delivery.bundle.manifest.occurrences[1].parent_id == parent.attempt.occurrence_id
    assert len(fake_ocr) == 1
