"""196 (plan F-Word-output W1, DECISIONS 196): own Word templates - the check at upload and the filling.

The owner's decisions of 2026-10-22: a Word template uses the markers of an Excel template and keeps its look; a
table's row (or a paragraph outside a table) repeats with its formatting; no calculating field (totals are markers).
The templates here are made-up Word documents built in the test, the output a made-up one beside the contract; the
data is synthetic. No AI call, no network.
"""

from __future__ import annotations

import io
import zipfile
from decimal import Decimal

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt

from jav.app.outputs import builtin, contract, docx, templates

OUTPUT = "test_word"
SPEC = {
    "what": "A made-up Word output of the tests.",
    "format": "docx",
    "builtin": {"version": "1.0.0"},
    "values": {"customer": "text", "language": "text", "issued": "date", "gross": "money", "note": "text"},
    "tables": {
        "lines": {"columns": {"name": "text", "quantity": "number", "unit_price": "money", "net": "money"}},
        "grid": {"columns": {"label": "text"}, "months": "money"},
    },
}


@pytest.fixture(autouse=True)
def word_output(monkeypatch):
    every = {**contract.outputs(), OUTPUT: SPEC}
    monkeypatch.setattr(contract, "outputs", lambda: every)


def _data(rows=None, **values) -> contract.OutputData:
    rows = [["Monitor", Decimal("2"), Decimal("54900"), Decimal("109800")],
            ["Licence", Decimal("1.5"), Decimal("1234.5"), Decimal("1851.75")]] if rows is None else rows
    base = {"customer": "Példa Kft.", "language": "hu", "issued": "2026-10-22", "gross": Decimal("141800.50"), "note": None}
    base.update(values)
    return contract.OutputData(OUTPUT, tuple(contract.Value(k, k, v) for k, v in base.items()), (
        contract.table("lines", "lines", ("name", "quantity", "unit_price", "net"), ("Name", "Qty", "Price", "Net"), rows,
                       ("text", "number", "money", "money")),))


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _template() -> bytes:
    """A heading with a value split across runs, a table with a header and a row template, a total and a count, a list
    paragraph, a header and a footer with a page number."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Ajánlat ").bold = True
    p.add_run("{{cus")
    p.add_run("tomer}}").italic = True
    p.add_run(", kelt: ")
    p.add_run("{{issued}}")
    table = doc.add_table(rows=2, cols=4)
    for cell, text in zip(table.rows[0].cells, ("Megnevezés", "Mennyiség", "Egységár", "Nettó"), strict=True):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("{{lines.name}}", "{{lines.quantity}}", "{{lines.unit_price}}",
                                                 "{{lines.net}}"), strict=True):
        cell.text = text
    table.rows[1].cells[0].paragraphs[0].runs[0].font.size = Pt(9)
    total = doc.add_paragraph("Nettó összesen: {{lines.net:sum}} Ft, {{lines:count}} tétel, bruttó {{gross}} Ft")
    total.runs[0].bold = True
    doc.add_paragraph("• {{lines.name}}")
    doc.add_paragraph("Megjegyzés: {{note}}")
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Ügyfél: {{customer}}"
    footer = section.footer.paragraphs[0]
    footer.text = "Oldal "
    footer._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    return _bytes(doc)


def _open(data: bytes):
    return Document(io.BytesIO(data))


def _xml(data: bytes, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(part).decode("utf-8")


# --- the filling ----------------------------------------------------------------------------------------------------


def test_the_values_go_where_their_markers_stand_and_the_text_around_keeps_its_formatting():
    out = _open(docx.fill(_data(), _template()))
    runs = out.paragraphs[0].runs
    assert out.paragraphs[0].text == "Ajánlat Példa Kft., kelt: 2026. 10. 22."
    assert runs[0].text == "Ajánlat " and runs[0].bold
    assert runs[1].text == "Példa Kft." and not runs[1].italic  # the value takes its marker's first character's run
    assert runs[2].text == "" and runs[3].text == ", kelt: "
    assert out.sections[0].header.paragraphs[0].text == "Ügyfél: Példa Kft."


def test_a_tables_row_repeats_with_its_formatting_and_numbers_are_written_by_code():
    out = _open(docx.fill(_data(), _template()))
    rows = out.tables[0].rows
    assert [[c.text for c in r.cells] for r in rows] == [
        ["Megnevezés", "Mennyiség", "Egységár", "Nettó"],
        ["Monitor", "2", "54\u00a0900", "109\u00a0800"],
        ["Licence", "1,5", "1\u00a0234,50", "1\u00a0851,75"]]
    assert all(r.cells[0].paragraphs[0].runs[0].font.size == Pt(9) for r in rows[1:])
    assert out.paragraphs[1].text == "Nettó összesen: 111\u00a0651,75 Ft, 2 tétel, bruttó 141\u00a0800,50 Ft"
    assert out.paragraphs[1].runs[0].bold
    assert [p.text for p in out.paragraphs[2:4]] == ["• Monitor", "• Licence"]


def test_english_writes_numbers_and_dates_the_english_way():
    out = _open(docx.fill(_data(language="en"), _template()))
    assert out.paragraphs[0].text.endswith("kelt: 22 October 2026")
    assert [c.text for c in out.tables[0].rows[2].cells] == ["Licence", "1.5", "1,234.50", "1,851.75"]


def test_without_rows_the_row_and_the_list_go_and_the_header_stays():
    out = _open(docx.fill(_data(rows=[]), _template()))
    assert [[c.text for c in r.cells] for r in out.tables[0].rows] == [["Megnevezés", "Mennyiség", "Egységár", "Nettó"]]
    assert [p.text for p in out.paragraphs[1:3]] == ["Nettó összesen: 0 Ft, 0 tétel, bruttó 141\u00a0800,50 Ft",
                                                      "Megjegyzés: "]


def test_a_table_of_only_its_repeated_row_goes_without_rows():
    doc = Document()
    doc.add_paragraph("{{customer}}")
    doc.add_table(rows=1, cols=1).rows[0].cells[0].text = "{{lines.name}}"
    out = _open(docx.fill(_data(rows=[]), _bytes(doc)))
    assert out.tables == []


def test_a_value_is_never_read_as_a_marker_and_a_line_break_is_a_break():
    out = docx.fill(_data(customer="{{gross}} Kft.", note="első sor\nmásodik sor"), _template())
    doc = _open(out)
    assert doc.paragraphs[0].text.startswith("Ajánlat {{gross}} Kft.")
    assert doc.paragraphs[4].text == "Megjegyzés: első sor\nmásodik sor"
    assert "<w:br/>" in _xml(out) and "ns0:" not in _xml(out)


def test_a_filled_template_keeps_its_page_number_and_leaves_no_marker():
    out = docx.fill(_data(), _template())
    assert "{{" not in _xml(out) and "{{" not in _xml(out, "word/header1.xml")
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        footer = next(n for n in z.namelist() if n.startswith("word/footer"))
        assert 'w:instr=" PAGE "' in z.read(footer).decode("utf-8")


def test_a_character_xml_cannot_hold_is_left_out():
    out = _open(docx.fill(_data(customer="A\x0bB"), _template()))
    assert out.paragraphs[0].text.startswith("Ajánlat AB")


# --- the check ------------------------------------------------------------------------------------------------------


def test_a_template_with_markers_passes_and_names_what_it_uses():
    check = docx.check(OUTPUT, _template())
    assert check.ok, check.problems
    assert check.values == ["customer", "issued", "gross", "note"]
    assert check.tables == {"lines": ["name", "quantity", "unit_price", "net"]}


def _doc(*texts: str) -> object:
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    return doc


def _codes(data: bytes) -> list[str]:
    return [p["code"] for p in docx.check(OUTPUT, data).problems]


def test_not_a_word_document_or_without_markers_is_refused():
    assert _codes(b"not a document") == ["not_docx"]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook/>")
    assert _codes(buf.getvalue()) == ["not_docx"]
    assert _codes(_bytes(_doc("Nothing to fill"))) == ["no_marker"]


@pytest.mark.parametrize(("build", "code"), [
    (lambda: _doc("{{nobody}}"), "unknown_marker"),
    (lambda: _doc("{{grid.months}}"), "months_in_word"),
])
def test_a_marker_the_output_cannot_fill_is_refused(build, code):
    assert code in _codes(_bytes(build()))


def _table_doc(*cells: str, merged: bool = False) -> object:
    doc = Document()
    table = doc.add_table(rows=1, cols=len(cells))
    for cell, text in zip(table.rows[0].cells, cells, strict=True):
        cell.text = text
    if merged:
        table.rows[0].cells[0]._tc.get_or_add_tcPr().append(parse_xml(f'<w:vMerge {nsdecls("w")} w:val="restart"/>'))
    return doc


def test_a_row_holds_one_tables_columns_without_a_value_or_a_vertical_merge():
    assert "row_mixes_tables" in _codes(_bytes(_table_doc("{{lines.name}}", "{{grid.label}}")))
    assert "value_in_row" in _codes(_bytes(_table_doc("{{lines.name}}", "{{customer}}")))
    assert "merged_row" in _codes(_bytes(_table_doc("{{lines.name}}", merged=True)))
    assert docx.check(OUTPUT, _bytes(_table_doc("{{lines.name}}", "{{lines.net}}"))).ok


def test_a_repeated_row_inside_another_is_refused():
    doc = _table_doc("{{lines.name}}")
    inner = doc.tables[0].rows[0].cells[0].add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = "{{grid.label}}"
    assert "nested_rows" in _codes(_bytes(doc))


def _with_field(instr: str, complex_field: bool = False) -> bytes:
    doc = _doc("{{customer}}")
    p = doc.paragraphs[0]._p
    if complex_field:
        for xml in (f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>',
                    f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve">{instr[:3]}</w:instrText></w:r>',
                    f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve">{instr[3:]}</w:instrText></w:r>',
                    f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>',
                    f'<w:r {nsdecls("w")}><w:t>x</w:t></w:r>',
                    f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'):
            p.append(parse_xml(xml))
    else:
        attr = instr.replace('"', "&quot;")
        p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{attr}"><w:r><w:t>x</w:t></w:r></w:fldSimple>'))
    return _bytes(doc)


@pytest.mark.parametrize("complex_field", [False, True])
def test_a_calculating_or_importing_field_is_refused_a_page_number_stays(complex_field):
    assert "field" in _codes(_with_field(" = 2 * 3 ", complex_field))
    assert "field" in _codes(_with_field(' DDEAUTO excel "x" ', complex_field))
    assert "field" in _codes(_with_field(' INCLUDETEXT "c:/x.docx" ', complex_field))
    assert docx.check(OUTPUT, _with_field(" NUMPAGES ", complex_field)).ok


def test_field_names():
    assert docx.field_name(" PAGE  \\* MERGEFORMAT ") == "PAGE"
    assert docx.field_name("=SUM(ABOVE)") == "="
    assert docx.field_name(" hyperlink \"https://example.org\" ") == "HYPERLINK"


def _with_relationship(kind: str, part: str = "word/_rels/document.xml.rels") -> bytes:
    """The template with an external relationship of `kind` added to `part`."""
    data = _bytes(_doc("{{customer}}"))
    buf = io.BytesIO()
    rel = (f'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}"'
           ' Target="file:///C:/Templates/Company.dotx" TargetMode="External"/>')
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for info in zin.infolist():
            content = zin.read(info)
            if info.filename == part:
                content = content.replace(b"</Relationships>", rel.encode() + b"</Relationships>")
            if info.filename == "word/settings.xml" and kind == "attachedTemplate":
                content = content.replace(b"</w:settings>", b'<w:attachedTemplate r:id="rIdX"/></w:settings>')
            zout.writestr(info, content)
        if part not in names:
            zout.writestr(part, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                          '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                          + rel + "</Relationships>")
    return buf.getvalue()


def test_a_link_to_another_file_is_refused_a_web_link_stays():
    assert _codes(_with_relationship("oleObject")) == ["external_link"]
    assert _codes(_with_relationship("subDocument")) == ["external_link"]
    assert docx.check(OUTPUT, _with_relationship("hyperlink")).ok


def test_the_attached_template_is_a_warning_and_goes_when_filled():
    data = _with_relationship("attachedTemplate", "word/_rels/settings.xml.rels")
    assert "w:attachedTemplate" in _xml(data, "word/settings.xml")
    check = docx.check(OUTPUT, data)
    assert check.ok and [w["code"] for w in check.warnings] == ["attached_template"]
    out = docx.fill(_data(), data)
    assert "attachedTemplate" not in _xml(out, "word/settings.xml")
    assert "Company.dotx" not in _xml(out, "word/_rels/settings.xml.rels")
    _open(out)


def test_macros_embedded_content_and_imported_files_are_refused():
    data = _bytes(_doc("{{customer}}"))
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info))
        zout.writestr("word/vbaProject.bin", b"\x00" * 10)
    assert _codes(buf.getvalue()) == ["unsafe_content"]
    doc = _doc("{{customer}}")
    doc.element.body.insert(0, parse_xml(f'<w:altChunk {nsdecls("w", "r")} r:id="rId99"/>'))
    assert "unsafe_content" in _codes(_bytes(doc))


def test_a_template_larger_than_the_limit_is_refused(monkeypatch):
    monkeypatch.setattr(contract, "templates", lambda: {"max_bytes": 100, "sample_rows": 3, "word_fields": []})
    assert _codes(_template()) == ["too_large"]


# --- the format's dispatch ------------------------------------------------------------------------------------------


def test_the_store_checks_and_fills_by_the_outputs_format():
    assert templates.engine(OUTPUT) is docx
    assert templates.media(OUTPUT) == docx.MEDIA and templates.extension(OUTPUT) == ".docx"
    assert templates.extension("reconcile_statement") == ".xlsx"


def test_a_word_outputs_built_in_template_is_its_domains(monkeypatch):
    monkeypatch.setattr(builtin, "_WORD", {})
    with pytest.raises(contract.ContractError):
        builtin.render(_data())
    builtin.register_word(OUTPUT, lambda data: _template())
    assert _open(builtin.render(_data())).paragraphs[0].text.startswith("Ajánlat Példa Kft.")
