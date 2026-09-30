"""Reports (054 K4): export (CSV / XLSX / JSON) and utility cost time series. Synthetic data, no AI calls.

Decisions of 2026-09-28: the source is the valid data of the selected run; breakdown by consumption point + utility;
monthly grid with day-proportional allocation; the shared water summary is informative only (against double
counting)."""

import csv
import io
import json
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from jav import export, report_utility
from jav.runtime import worker
from tests import test_api
from tests.test_api import _ready_wp, _start

env = test_api.env


def _bill(item, doc_type, start, end, amount, *, address="1111 Budapest, Minta utca 11", consumption=None, supplier="MVM"):
    fields = {"supplier_name": supplier, "consumption_address": address, "billing_period_start": start,
              "billing_period_end": end, "gross_total": amount}
    if consumption is not None:
        fields["consumption_kwh" if doc_type == "villamos_energia_szamla" else "consumption_m3"] = consumption
    return {"item_id": item, "file": f"{item}.pdf", "doc_type": doc_type, "fields": fields, "pages": {"gross_total": 1},
            "corrected": [], "open_reasons": []}


# --- monthly allocation ------------------------------------------------------------------------------------


def test_split_by_month_is_day_proportional_and_exact():
    parts = report_utility.split_by_month(date(2026, 2, 5), date(2026, 3, 4), Decimal("28000"))
    assert list(parts) == ["2026-02", "2026-03"]
    assert parts["2026-02"] == Decimal("24000.00") and parts["2026-03"] == Decimal("4000.00")  # 24 + 4 days
    odd = report_utility.split_by_month(date(2026, 1, 1), date(2026, 3, 31), Decimal("100"))
    assert sum(odd.values()) == Decimal("100")  # exact total even after rounding


def test_series_by_address_and_utility_with_missing_and_overlap_months():
    records = [
        _bill("a", "villamos_energia_szamla", "2026-01-01", "2026-01-31", "10000", consumption="300"),
        _bill("b", "villamos_energia_szamla", "2026-03-01", "2026-03-31", "12000", consumption="350"),  # February missing
        _bill("c", "villamos_energia_szamla", "2026-03-15", "2026-04-30", "9000"),                     # overlaps in March
        _bill("d", "foldgaz_szamla", "2026-01-01", "2026-01-31", "20000", address="1111 BUDAPEST Minta utca 11."),
        _bill("e", "villamos_energia_szamla", "2026-01-01", "2026-01-31", "5000", address="1011 Budapest, Fő utca 1"),
    ]
    rep = report_utility.build(records)
    assert rep["months"] == ["2026-01", "2026-02", "2026-03", "2026-04"]
    keys = {(s["address"], s["utility"]) for s in rep["series"]}
    assert len(rep["series"]) == 3 and ("1111 Budapest, Minta utca 11", "Földgáz") in keys  # address spellings merge
    villany = next(s for s in rep["series"] if s["utility"] == "Villamos energia" and "Minta" in s["address"])
    cells = villany["cells"]
    assert cells["2026-01"]["status"] == "ok" and cells["2026-01"]["amount"] == "10000.00"
    assert cells["2026-02"]["status"] == "missing" and cells["2026-02"]["amount"] is None
    assert cells["2026-03"]["status"] == "overlap" and {s["item_id"] for s in cells["2026-03"]["sources"]} == {"b", "c"}
    assert villany["total"] == "31000.00" and villany["consumption_unit"] == "kWh"
    assert cells["2026-01"]["consumption"] == "300.00"


def test_district_numeral_does_not_split_an_address():
    records = [_bill("a", "villamos_energia_szamla", "2026-01-01", "2026-01-31", "100", address="1111 BUDAPEST MINTA UTCA 11."),
               _bill("b", "villamos_energia_szamla", "2026-02-01", "2026-02-28", "100", address="1111 Budapest XI., Minta utca 11"),
               _bill("c", "villamos_energia_szamla", "2026-02-01", "2026-02-28", "100", address="1111 Budapest XI. ker. Minta utca 11")]
    rep = report_utility.build(records)
    assert len(rep["series"]) == 1 and rep["series"][0]["cells"]["2026-02"]["status"] == "overlap"


def test_summary_statement_is_shown_but_not_summed():
    records = [_bill("s", "viz_szamla", "2026-01-01", "2026-01-31", "30000"), _bill("v", "vizmuvek_szamla", "2026-01-01", "2026-01-31", "10000")]
    rep = report_utility.build(records)
    summary = next(s for s in rep["series"] if s["summary_only"])
    assert summary["total"] == "30000.00"
    assert rep["grand_total"] == "10000.00"  # only the individual bill counts


def test_same_invoice_in_two_documents_counts_once():
    a = _bill("m1", "mohu_szamla", "2026-01-01", "2026-01-31", "3266")
    b = _bill("m2", "mohu_szamla", "2026-01-01", "2026-01-31", "3266")
    a["fields"]["invoice_number"] = b["fields"]["invoice_number"] = "HU-123"
    rep = report_utility.build([a, b])
    [s] = rep["series"]
    assert s["total"] == "3266.00" and s["cells"]["2026-01"]["status"] == "ok"
    assert rep["duplicates"] == [{"item_id": "m2", "file": "m2.pdf", "doc_type": "mohu_szamla", "same_as": "m1", "same_as_file": "m1.pdf"}]


def test_settlement_bill_is_booked_at_its_end_month_and_does_not_break_coverage():
    records = [_bill(f"p{m}", "foldgaz_szamla", f"2026-0{m}-01", f"2026-0{m}-{28 if m == 2 else 30 if m in (4, 6) else 31}", "1000")
               for m in range(1, 7)]
    records.append(_bill("yr", "foldgaz_szamla", "2026-01-01", "2026-06-30", "-500"))  # annual settlement: difference
    rep = report_utility.build(records)
    [s] = rep["series"]
    assert all(c["status"] == "ok" for c in s["cells"].values())
    assert s["cells"]["2026-06"]["amount"] == "500.00" and s["cells"]["2026-06"]["settlement"] is True
    assert s["total"] == "5500.00"


def test_bills_without_period_or_amount_are_listed_not_guessed():
    records = [_bill("x", "mohu_szamla", None, "2026-01-31", "1000"), _bill("y", "mohu_szamla", "2026-01-01", "2026-01-31", None),
               {"item_id": "z", "file": "z.pdf", "doc_type": "invoice_hu", "fields": {}, "pages": {}, "corrected": [], "open_reasons": []}]
    rep = report_utility.build(records)
    assert rep["series"] == []
    assert {u["item_id"]: u["reason"] for u in rep["unplaced"]} == {"x": "no_period", "y": "no_amount"}  # non-utility doc left out


# --- export --------------------------------------------------------------------------------------------------


def test_csv_guards_formulas_but_keeps_numbers():
    data = export.csv_bytes(["a", "b"], [["=HYPERLINK(\"x\")", "-1200.50"], ["@SUM(A1)", "normál"]])
    assert data.startswith(b"\xef\xbb\xbf")  # BOM: Hungarian Excel opens it as UTF-8
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))
    assert rows[1] == ["'=HYPERLINK(\"x\")", "-1200.50"] and rows[2][0] == "'@SUM(A1)"


def test_run_exports_over_the_service(env):
    c = env["client"]
    wp = _ready_wp(c, env["folder"])
    run_id = _start(c, wp["id"]).json()["run_id"]
    worker.run_worker(once=True)
    base = f"/api/runs/{run_id}/export"

    r = c.get(base, params={"format": "csv", "table": "datapoints"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"] and r.headers["x-export-rows"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    gross = [x for x in rows if x["Mező"] == "gross_total"]
    assert len(gross) == 2 and all(x["Oldal"] for x in gross)  # every number can be traced back to its source

    r = c.get(base, params={"format": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Adatpontok", "Tételsorok", "Közmű-költség"} <= set(wb.sheetnames)
    for ws in wb.worksheets:  # text is never a formula
        assert not any(isinstance(c_.value, str) and c_.data_type == "f" for row in ws.iter_rows() for c_ in row)

    r = c.get(base, params={"format": "json"})
    body = json.loads(r.content)
    assert body["run_id"] == run_id and len(body["documents"]) == 2

    r = c.get(f"/api/runs/{run_id}/reports/utility-cost")
    assert r.status_code == 200 and r.json()["series"] == [] and len(r.json()["unplaced"]) == 0  # not utility documents

    assert c.get(base, params={"format": "exe"}).status_code == 422


def test_result_tables_offer_only_views_with_data(env):
    # 058: no utility cost view for Hungarian invoices; there is one for utility bill records
    c = env["client"]
    wp = _ready_wp(c, env["folder"])
    run_id = _start(c, wp["id"]).json()["run_id"]
    worker.run_worker(once=True)
    tables = c.get(f"/api/runs/{run_id}").json()["tables"]
    assert tables[:2] == ["documents", "datapoints"] and "utility" not in tables
    assert report_utility.has_utility([_bill("a", "villamos_energia_szamla", "2026-01-01", "2026-01-31", "1000")])
    assert not report_utility.has_utility([{"item_id": "b", "doc_type": "invoice_hu"}])
