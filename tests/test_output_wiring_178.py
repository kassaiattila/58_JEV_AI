"""178 (plan 170 L4, DECISIONS 178): own templates over the service - upload, activation, the choice per package, and
the downloads, the release and the reconciliation's approval made with the chosen template.

The owner's decisions of 2026-10-19: a template is chosen per package; the utility cost has its own download; what a
release or an approval stores records the template and its version. Synthetic PDFs, made-up templates built in the
test, a fake JEV client; no paid calls.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from jav import report_utility, store
from jav import reconcile_statement as rs
from jav.outputs import templates
from tests import test_flexible_package_160c as fp
from tests.test_characterisation_178 import _world
from tests.test_flexible_package_160c import HUMAN, _assign, _package
from tests.test_output_templates_178 import _statement_template
from tests.test_package_release_171 import _live, _release

env = fp.env  # the 160-C fixture: a store of its own, two synthetic invoices, a fake JEV client


def _template(folder, name: str, **cells) -> str:
    """A made-up template in a folder of its own (not the package's documents)."""
    wb = Workbook()
    for coord, value in cells.items():
        wb.active[coord] = value
    path = folder.parent / "templates" / name
    path.parent.mkdir(exist_ok=True)
    wb.save(path)
    return str(path)


def _upload(c, path: str, output: str = "package_result", name: str = "Ours", **extra):
    return c.post("/api/output-templates", headers=HUMAN, json={"output": output, "name": name, "path": path, **extra})


def test_a_template_is_uploaded_activated_chosen_and_used_by_the_download_and_the_release(env):
    c = env["client"]
    refused = _upload(c, _template(env["folder"], "bad.xlsx", A1="=1+1"))
    assert refused.status_code == 422 and refused.json()["error"] == "template_refused"
    assert [p["code"] for p in refused.json()["check"]["problems"]] == ["formula"]
    path = _template(env["folder"], "ours.xlsx", A1="Result of {{package}}", A3="{{invoice_hu.invoice_number}}",
                     B3="{{invoice_hu.gross_total}}", A5="{{invoice_hu:count}} invoices")
    up = _upload(c, path)
    assert up.status_code == 201, up.text
    tpl = up.json()["template"]
    listing = c.get("/api/output-templates").json()
    assert [t["template"] for t in listing["templates"]] == [tpl]
    assert set(listing["outputs"]) == {"package_result", "utility_cost", "reconcile_statement", "quote"}  # 196: the quote
    preview = c.get(f"/api/output-templates/{tpl}/versions/1/preview.xlsx")
    assert preview.status_code == 200 and load_workbook(io.BytesIO(preview.content)).active["A7"].value == "3 invoices"  # below three made-up rows
    assert c.get(f"/api/output-templates/{tpl}/versions/1/file").content == open(path, "rb").read()

    wp = _package(c, env["folder"])
    _assign(c, wp["id"])
    assert c.put(f"/api/workpackages/{wp['id']}/output-templates/package_result", headers=HUMAN,
                 json={"template": tpl}).status_code == 409  # a draft is not chosen
    assert c.post(f"/api/output-templates/{tpl}/versions/1/activate", headers=HUMAN, json={}).status_code == 200
    chosen = c.put(f"/api/workpackages/{wp['id']}/output-templates/package_result", headers=HUMAN, json={"template": tpl})
    assert chosen.status_code == 200 and chosen.json()["outputs"]["package_result"]["chosen"]["version"] == 1
    assert c.get(f"/api/workpackages/{wp['id']}/output-templates").json()["outputs"]["utility_cost"]["chosen"] is None

    _live(c, wp["id"])
    sheet = load_workbook(io.BytesIO(c.get(f"/api/workpackages/{wp['id']}/result/export.xlsx").content)).active
    assert sheet["A1"].value == f"Result of {wp['name']}" and sheet["A6"].value == "2 invoices"
    utility = c.get(f"/api/workpackages/{wp['id']}/result/utility-cost.xlsx")
    titles = [t.title for t in report_utility.output_data([]).tables]  # the built-in template's two sheets
    assert utility.status_code == 200 and load_workbook(io.BytesIO(utility.content)).sheetnames == titles

    released = _release(c, wp["id"])
    assert released.status_code == 200, released.text
    files = {f["name"]: f for f in released.json()["version"]["files"]}
    assert (files["result.xlsx"]["template"], files["result.xlsx"]["template_version"]) == (tpl, "1")
    assert files["result.json"]["template"] is None and "utility_cost.xlsx" not in files  # no utility bills
    stored = c.get(f"/api/workpackages/{wp['id']}/versions/1/files/result.xlsx").content
    assert load_workbook(io.BytesIO(stored)).active["A3"].value == sheet["A3"].value

    assert c.post(f"/api/output-templates/{tpl}/retire", headers=HUMAN, json={}).status_code == 200
    after = load_workbook(io.BytesIO(c.get(f"/api/workpackages/{wp['id']}/result/export.xlsx").content))
    assert after.sheetnames[0] == "invoice_hu"  # the built-in template again
    assert c.get(f"/api/workpackages/{wp['id']}/versions/1/files/result.xlsx").content == stored  # the release keeps its file
    with store.connect() as conn:
        events = [r[0] for r in conn.execute("SELECT action FROM workpackage_events WHERE workpackage_id=?", (wp["id"],))]
    assert "output_template_choose" in events


def test_a_reconciliation_is_approved_with_its_chosen_template_and_reopened_with_the_same(tmp_path):
    with store.use_store(tmp_path / "w178.sqlite"):
        wp_id, _inv, _ids = _world()
        v = templates.add("reconcile_statement", "Ours", "ours.xlsx", _statement_template(), actor="Reviewer")
        templates.activate(v["template"], 1, actor="Reviewer")
        templates.choose(wp_id, "reconcile_statement", v["template"], actor="Reviewer")
        assert load_workbook(io.BytesIO(rs.xlsx(wp_id)[0]))["Statement"]["A1"].value.startswith("Reconciliation of April")
        rs.approve(wp_id, fingerprint=rs.current(wp_id)["fingerprint"], actor="Reviewer")
        assert [(f["state"], f["template"], f["template_version"]) for f in rs.version_files(wp_id, 1)] == [
            ("approved", v["template"], "1")]
        v2 = templates.add("reconcile_statement", "Ours", "ours-2.xlsx", _statement_template(), actor="Reviewer",
                           template=v["template"])
        templates.activate(v2["template"], 2, actor="Reviewer")
        rs.reopen(wp_id, note="a payment was missed", actor="Reviewer")
        assert [(f["state"], f["template_version"]) for f in rs.version_files(wp_id, 1)] == [("approved", "1"), ("earlier", "1")]
        assert load_workbook(io.BytesIO(rs.xlsx(wp_id, version=1)[0]))["Statement"]["A1"].value.startswith("Reconciliation of")


def test_the_command_line_checks_stores_activates_and_chooses(tmp_path, capsys):
    from jav import cli

    with store.use_store(tmp_path / "c178.sqlite"):
        wp_id, _inv, _ids = _world()
        bad = _template(tmp_path / "x", "bad.xlsx", A1="=1+1", A2="{{package}}")
        assert cli.main(["output-template-add", bad, "--output", "reconcile_statement", "--name", "Bad", "--actor", "R"]) == 2
        assert "problem: formula Sheet!A1" in capsys.readouterr().out
        good = tmp_path / "templates" / "good.xlsx"
        good.write_bytes(_statement_template())
        assert cli.main(["output-template-add", str(good), "--output", "reconcile_statement", "--name", "Good", "--actor", "R"]) == 0
        tpl = templates.listing()[0]["template"]
        assert cli.main(["output-template-activate", tpl, "1", "--actor", "R"]) == 0
        assert cli.main(["output-template-choose", wp_id, "reconcile_statement", tpl, "--actor", "R"]) == 0
        assert f"{wp_id} reconcile_statement: {tpl} version 1" in capsys.readouterr().out
        out = tmp_path / "preview.xlsx"
        assert cli.main(["output-template-preview", tpl, "1", "--out", str(out)]) == 0 and out.exists()
        assert cli.main(["output-template-choose", wp_id, "reconcile_statement", "--builtin", "--actor", "R"]) == 0
        assert "built-in template" in capsys.readouterr().out.splitlines()[-1]


def test_a_store_from_before_178_gets_the_template_columns_of_a_versions_files(tmp_path):
    import sqlite3

    from jav import package_release

    path = tmp_path / "old.sqlite"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE package_version_files (version_id INTEGER NOT NULL, name TEXT NOT NULL, media TEXT NOT NULL,"
                 " rows INTEGER NOT NULL, sha256 TEXT NOT NULL, content BLOB NOT NULL, PRIMARY KEY (version_id, name))")
    conn.commit()
    package_release._migrate(conn)
    package_release._migrate(conn)  # idempotent
    assert {"template", "template_version"} <= {r[1] for r in conn.execute("PRAGMA table_info(package_version_files)")}
    conn.close()
