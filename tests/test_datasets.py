"""Adatkészletek a szolgáltatáson át (056 U1): egységes lekérdezés, gyorsítótár, letöltés. Mesterséges adat, AI-hívás nélkül."""

import csv
import io
import json

from openpyxl import load_workbook

from jav import datasets
from jav.runtime import worker
from tests import test_api
from tests.test_api import HUMAN, _ready_wp, _start

env = test_api.env


def _q(c, name, scope=None, **query):
    r = c.post(f"/api/datasets/{name}/query", json={"scope": scope or {}, "query": query})
    assert r.status_code == 200, r.text
    return r.json()


def _run(c, folder):
    wp = _ready_wp(c, folder)
    run_id = _start(c, wp["id"]).json()["run_id"]
    worker.run_worker(once=True)
    return wp, run_id


def test_catalog_lists_every_dataset_with_its_scope(env):
    body = env["client"].get("/api/datasets").json()
    by_name = {d["name"]: d for d in body["datasets"]}
    assert {"runs", "workpackages", "workpackage_items", "run_items", "documents", "datapoints", "line_items", "calls",
            "mailbox_pulls", "utility_cost", "utility_sources"} <= set(by_name)
    assert by_name["datapoints"]["scope"] == ["run_id"] and by_name["runs"]["optional_scope"] == ["workpackage_id"]
    assert by_name["workpackages"]["natural_sort"] == [{"col": "created_at", "desc": True}]  # 062: a fejlécen jelölve


def test_lists_are_queried_on_the_service_without_silent_cut(env):
    c = env["client"]
    wp, run_id = _run(c, env["folder"])
    runs = _q(c, "runs")
    assert runs["total"] == 1 and runs["rows"][0]["run_id"] == run_id and runs["rows"][0]["workpackage_name"] == "Mesterséges számlák"
    status = next(col for col in runs["columns"] if col["key"] == "status")
    assert status["labels"]["done"] == "Kész"  # a felirat a szolgáltatásból jön: a keresés arra is talál
    assert _q(c, "runs", q="Éles")["matched"] == 1
    assert _q(c, "runs", {"workpackage_id": wp["id"]})["total"] == 1

    assert _q(c, "workpackages")["rows"][0]["name"] == "Mesterséges számlák"
    items = _q(c, "workpackage_items", {"workpackage_id": wp["id"]}, sort=[{"col": "name", "desc": True}])
    assert [r["name"] for r in items["rows"]] == ["szamla_2.pdf", "szamla_1.pdf"]
    run_items = _q(c, "run_items", {"run_id": run_id})
    assert run_items["total"] == 2 and all(r["status"] == "done" for r in run_items["rows"])


def test_result_tables_filter_sort_and_page(env):
    c = env["client"]
    _wp, run_id = _run(c, env["folder"])
    scope = {"run_id": run_id}
    docs = _q(c, "documents", scope)
    assert docs["total"] == 2 and any(col["key"] == "f:gross_total" and col["kind"] == "money" for col in docs["columns"])
    dps = _q(c, "datapoints", scope, filters=[{"col": "field", "op": "eq", "value": "gross_total"}])
    assert dps["matched"] == 2 and all(r["page"] for r in dps["rows"])
    page = _q(c, "datapoints", scope, sort=[{"col": "field"}], offset=1, limit=3)
    assert len(page["rows"]) == 3 and page["matched"] == page["total"] > 3
    assert "field" in page["facets"] and "gross_total" in page["facets"]["field"]
    assert _q(c, "calls", scope)["total"] >= 0


def test_named_errors(env):
    c = env["client"]
    _wp, run_id = _run(c, env["folder"])
    assert c.post("/api/datasets/nincs/query", json={}).status_code == 404
    assert c.post("/api/datasets/datapoints/query", json={}).status_code == 422  # hatókör nélkül
    r = c.post("/api/datasets/datapoints/query", json={"scope": {"run_id": run_id}, "query": {"sort": [{"col": "nincs"}]}})
    assert r.status_code == 422 and "unknown column" in r.json()["message"]
    assert c.post("/api/datasets/datapoints/query", json={"scope": {"run_id": "run-nincs"}}).status_code == 404


def test_cache_is_refreshed_after_a_correction(env):
    c = env["client"]
    wp, run_id = _run(c, env["folder"])
    datasets.clear_cache()
    scope = {"run_id": run_id}
    flt = [{"col": "field", "op": "eq", "value": "invoice_number"}, {"col": "value", "op": "eq", "value": "JAVITOTT-1"}]
    assert _q(c, "datapoints", scope, filters=flt)["matched"] == 0
    item_id = wp["items"][0]["item_id"]
    ok = c.post(f"/api/runs/{run_id}/items/{item_id}/correction", headers=HUMAN,
                json={"fields": {"invoice_number": "JAVITOTT-1"}, "expected_revision": 0})
    assert ok.status_code == 200, ok.text
    hit = _q(c, "datapoints", scope, filters=flt)
    assert hit["matched"] == 1 and hit["rows"][0]["corrected"] == "igen"


def test_run_fingerprint_follows_only_this_runs_subjects_and_task_decisions(env):
    """062 (Q-ujjlenyomat): eddig bármely futás teendő-változása minden futás tárolt eredményét érvénytelenítette."""
    from jav import store

    c = env["client"]
    wp, run_id = _run(c, env["folder"])
    fp = datasets._run_fingerprint({"run_id": run_id})
    other = store.review_enqueue(subject_kind="document", subject_id="f" * 64, run_id="run-masik:ffff", reasons=["pick:x"], producer="m2")
    assert datasets._run_fingerprint({"run_id": run_id}) == fp  # más alany teendője nem érinti
    store.review_resolve(store.review_open_reasons("document", "f" * 64)[0]["id"], actor="Minta Anna")
    assert datasets._run_fingerprint({"run_id": run_id}) == fp and other
    own = wp["items"][0]["item_id"]
    store.review_enqueue(subject_kind="document", subject_id=own, run_id="run-korabbi:" + own[:16], reasons=["pick:y"], producer="m2")
    fp2 = datasets._run_fingerprint({"run_id": run_id})
    assert fp2 != fp  # a futás alanyának (akár korábbi futásból jött) teendője a „korábbi” jelölés miatt számít
    store.email_task_decide(f"{run_id}:{own[:16]}", 0, decision="accepted", actor="Minta Anna", note=None)
    fp3 = datasets._run_fingerprint({"run_id": run_id})
    assert fp3 != fp2
    store.email_task_decide(f"{run_id}:{own[:16]}", 0, decision="rejected", actor="Minta Anna", note=None)
    assert datasets._run_fingerprint({"run_id": run_id}) != fp3  # ugyanabban a másodpercben átdöntve is


def test_export_scope_columns_and_formats(env):
    c = env["client"]
    _wp, run_id = _run(c, env["folder"])
    url = "/api/datasets/datapoints/export"
    scope = {"run_id": run_id}
    flt = {"filters": [{"col": "field", "op": "eq", "value": "gross_total"}]}

    r = c.post(url, json={"scope": scope, "query": flt, "rows": "filtered", "format": "csv", "columns": ["file", "field", "value"]})
    assert r.status_code == 200, r.text
    assert r.headers["x-export-rows"] == "2" and "datapoints-" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["Irat", "Mező", "Érték"] and len(rows) == 3

    everything = c.post(url, json={"scope": scope, "query": flt, "rows": "all", "format": "json"})
    body = json.loads(everything.content)
    assert int(everything.headers["x-export-rows"]) == len(body["rows"]) > 2  # a szűrő a „minden sor” letöltésre nem hat

    keys = [row["_key"] for row in _q(c, "datapoints", scope, limit=3)["rows"]]
    sel = c.post(url, json={"scope": scope, "query": {"keys": keys}, "rows": "selected", "format": "xlsx"})
    ws = load_workbook(io.BytesIO(sel.content)).active
    assert sel.headers["x-export-rows"] == "3" and ws.max_row == 4 and ws.title == "Adatpontok"

    assert c.post(url, json={"scope": scope, "query": {}, "rows": "selected", "format": "csv"}).status_code == 422
    assert c.post(url, json={"scope": scope, "rows": "all", "format": "csv", "columns": ["nincs"]}).status_code == 422


def test_export_guards_formulas(env):
    c = env["client"]
    wp, run_id = _run(c, env["folder"])
    item_id = wp["items"][0]["item_id"]
    c.post(f"/api/runs/{run_id}/items/{item_id}/correction", headers=HUMAN,
           json={"fields": {"invoice_number": "=HYPERLINK(\"x\")"}, "expected_revision": 0})
    r = c.post("/api/datasets/datapoints/export", json={"scope": {"run_id": run_id}, "rows": "all", "format": "csv",
                                                       "query": {"q": "HYPERLINK"}})
    text = r.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text


def test_utility_cost_long_form_and_sources(env, monkeypatch):
    from jav import export
    from tests.test_reports import _bill

    c = env["client"]
    _wp, run_id = _run(c, env["folder"])
    records = [_bill("a", "villamos_energia_szamla", "2026-01-01", "2026-01-31", "10000"),
               _bill("b", "villamos_energia_szamla", "2026-03-01", "2026-03-31", "12000")]
    monkeypatch.setattr(export, "run_records", lambda _run_id: records)
    datasets.clear_cache()
    scope = {"run_id": run_id}
    cost = _q(c, "utility_cost", scope, sort=[{"col": "month"}])
    assert [(r["month"], r["status"], r["amount"]) for r in cost["rows"]] == [
        ("2026-01", "ok", "10000.00"), ("2026-02", "missing", None), ("2026-03", "ok", "12000.00")]
    assert _q(c, "utility_cost", scope, q="hiányzik")["matched"] == 1  # a keresés az állapot feliratára is talál
    src = _q(c, "utility_sources", scope, filters=[{"col": "month", "op": "eq", "value": "2026-03"}])
    assert [r["file"] for r in src["rows"]] == ["b.pdf"]


def test_export_with_accented_scope_value_keeps_the_name(env):
    """066 Á22: az ékezetes személynév (pl. a tevékenységnapló szerzője) a letöltés fájlnevébe kerül; eddig a fejléc
    Latin-1 kódolása hibát adott (422). Most ASCII-tartalék + UTF-8 fájlnév (RFC 5987)."""
    r = env["client"].post("/api/datasets/activity/export", json={"scope": {"actor": "Kőműves Ödön"}, "format": "csv", "rows": "all"})
    assert r.status_code == 200, r.text
    cd = r.headers["content-disposition"]
    assert cd.startswith("attachment; filename=\"activity-Kom") and "filename*=UTF-8''activity-K%C5%91m%C5%B1ves" in cd
    assert all(ord(ch) < 128 for ch in cd)
