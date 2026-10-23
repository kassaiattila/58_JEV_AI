"""160 (review S-2): the native facts' downloads and the JSON download mask a full card number too.

The masking of 153 wrapped the documents, data points and line items tables; the native facts table (its CSV, its
Excel sheet and its dataset) and the JSON download wrote a full card number from a native result or from data saved
before the masking. They are masked now. A digest (a long run of lowercase hexadecimal digits, as a document's id or a
source's sha256) is never taken for a card number, so masking cannot change an identity. Synthetic values: the
published Visa test number, built at run time."""

import io
import json

from openpyxl import load_workbook

from jav import corrections, export
from jav.documents import native_results
from jav.core import card_numbers, store
from tests.test_native_review_api_109 import published_run

CARD = " ".join(["4111"] + ["1111"] * 3)  # the schemes' published test number
FULL, MASKED = "4111 1111", "411111xxxxxx1111"


def _run_with_a_card_number(tmp_path):
    rid, item, publication, _cite = published_run(tmp_path)
    fact = native_results.machine_facts(publication)[0]
    corrections.save_native(rid, item["item_id"], values={fact.fact_id: f"card {CARD}"}, native_sources={},
                            expected_revision=0, expected_result_version=publication.result_version, actor="reviewer")
    return rid


def test_the_native_facts_csv_and_excel_sheet_mask_a_full_card_number(tmp_path):
    with store.use_store(tmp_path / "card.sqlite"):
        rid = _run_with_a_card_number(tmp_path)
        csv = export.render(rid, "csv", "native_facts")[0].decode("utf-8-sig")
        book = load_workbook(io.BytesIO(export.xlsx_bytes(rid)))
        cells = [str(c.value) for row in book["Native facts"].iter_rows() for c in row if c.value is not None]
    assert FULL not in csv and MASKED in csv
    assert not any(FULL in c for c in cells) and any(MASKED in c for c in cells)


def test_the_json_download_masks_a_full_card_number(tmp_path):
    with store.use_store(tmp_path / "card.sqlite"):
        rid = _run_with_a_card_number(tmp_path)
        body = export.json_bytes(rid).decode("utf-8")
    assert FULL not in body and MASKED in body
    assert json.loads(body)["run_id"] == rid


def test_a_digest_is_never_taken_for_a_card_number():
    digest = "deadbeef" + "4111111111111111" + "c" * 40  # 64 hexadecimal digits with a Luhn-valid run of sixteen
    assert card_numbers.mask_text(digest) == digest
    assert card_numbers.mask_text(f"source {digest}, card 4111111111111111") == f"source {digest}, card {MASKED}"
