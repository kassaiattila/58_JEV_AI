"""Build real artificial files, independent expected labels and reader evidence.

Run from the repository: python -m tests.export_native_corpus is not required;
use python tests/export_native_corpus.py <new-output-directory>.
"""
from email.message import EmailMessage
from io import BytesIO
import json
from pathlib import Path
import sys
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from jav.readers.pipeline import read_files


def build(destination: Path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    destination.mkdir(parents=True, exist_ok=False)
    incoming = destination / "inputs"
    incoming.mkdir()
    expected = {}
    for index in range(6):
        code, quantity = f"SYN-{index:04d}", str(index + 1)
        values = {"order_code": code, "quantity": quantity}
        formats = {
            ".txt": f"Order code: {code}\nQuantity: {quantity}".encode(),
            ".csv": f"order_code,quantity\n{code},{quantity}\n".encode(),
            ".json": json.dumps(values).encode(),
            ".xml": f"<order><code>{code}</code><quantity>{quantity}</quantity></order>".encode(),
            ".html": f"<table><tr><th>Order code</th><th>Quantity</th></tr><tr><td>{code}</td><td>{quantity}</td></tr></table>".encode(),
        }
        document = Document()
        document.add_paragraph(f"Order code: {code}")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Quantity"
        table.cell(0, 1).text = quantity
        document.add_paragraph("Synthetic source ends here")
        stream = BytesIO()
        document.save(stream)
        formats[".docx"] = stream.getvalue()
        book = Workbook()
        sheet = book.active
        sheet.title = "Orders"
        sheet.append(["order_code", "quantity"])
        sheet.append([code, quantity])
        stream = BytesIO()
        book.save(stream)
        book.close()
        formats[".xlsx"] = stream.getvalue()
        stream = BytesIO()
        Image.new("RGB", (24 + index, 16), (index * 30, 80, 160)).save(stream, format="PNG")
        formats[".png"] = stream.getvalue()
        for extension, data in formats.items():
            name = f"sample-{index:02d}{extension}"
            (incoming / name).write_bytes(data)
            expected[name] = {"sentinel": None if extension == ".png" else code,
                              "facts": [] if extension == ".png" else [["order", "code", code], ["order", "quantity", quantity]],
                              "reading_gap_required": extension in {".png", ".json", ".xml", ".html"}}
    for index in range(12):
        message = EmailMessage()
        message["From"] = "sender@example.invalid"
        message["To"] = "reader@example.invalid"
        message["Subject"] = f"Synthetic package {index}"
        message.set_content(f"Please process synthetic package {index}.")
        source = incoming / f"sample-{index % 6:02d}"
        extension = ".xlsx" if index % 2 else ".docx"
        attachment = source.with_suffix(extension)
        message.add_attachment(attachment.read_bytes(), maintype="application", subtype="octet-stream", filename=attachment.name)
        message.add_attachment(b"unsupported synthetic binary", maintype="application", subtype="octet-stream", filename="unknown.bin")
        if index % 3 == 0:
            image = incoming / f"sample-{index % 6:02d}.png"
            message.add_attachment(image.read_bytes(), maintype="image", subtype="png", filename=image.name, disposition="inline")
        name = f"package-{index:02d}.eml"
        (incoming / name).write_bytes(message.as_bytes())
        expected[name] = {"attachments": 3 if index % 3 == 0 else 2}
    started = time.perf_counter()
    delivery = read_files([incoming])
    delivery.save(destination / "delivery")
    roots = {o.original_name: o for o in delivery.bundle.manifest.occurrences if o.parent_id is None}
    readings = {r.attempt.occurrence_id: r for r in delivery.bundle.results}
    inventories = {i.occurrence_id: i for i in delivery.bundle.manifest.inventories}
    checks = []
    for name, gold in expected.items():
        occurrence = roots[name]
        result = readings[occurrence.occurrence_id]
        passed = result.status in {"complete", "partial"}
        if "attachments" in gold:
            passed = passed and inventories[occurrence.occurrence_id].expected_children == gold["attachments"]
        else:
            if gold["sentinel"]:
                passed = passed and any(gold["sentinel"] in (e.text or "") for e in result.elements)
            if gold["reading_gap_required"]:
                passed = passed and result.status == "partial" and bool(result.issues)
        checks.append({"name": name, "passed": bool(passed), "status": result.status})
    report = {"scope": "native-reader-synthetic-not-AI-accuracy", "files": 48, "email_packages": 12,
              "occurrences": len(delivery.bundle.manifest.occurrences), "objects": len(delivery.bundle.manifest.objects),
              "seconds": round(time.perf_counter() - started, 3), "checks": checks,
              "passed": sum(row["passed"] for row in checks), "provider_calls": 0,
              "remaining_acceptance": "This corpus does not cover all formats and failure cases required by plan 101."}
    (destination / "expected.json").write_text(json.dumps(expected, indent=2), encoding="utf-8")
    (destination / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "checks"}))
    if report["passed"] != len(checks):
        raise SystemExit(1)


if __name__ == "__main__":
    build(Path(sys.argv[1]))
