"""231 (stage D, D-f7): a rated offer as a Word document, and the priced document that stores it with its version.

A document whose domain does not store it at the approval (the insurance offer's) is stored by the priced document with
the version it printed, once, with the case's chosen template (the built-in one here): the result names the file and its
sha256, and a second run stores nothing new. The built-in template shows how the total follows from the modules (their
subtotal, the factors on the total, the total rounded), leaves out a line or the factors' table the data has none of,
and marks a draft. A decision that names no premium the case kept is no offer. Invented data and a stand-in builder; no paid call.
"""

from __future__ import annotations

import hashlib
import io

import pytest

from jav.app import package_release, rated_offer, registrations
from jav.app.outputs import builtin, contract, parts
from jav.core import identity, store
from jav.core.process import kinds
from jav.core.process.definition import ProcessDef, StepDef
from jav.core.process.engine import StepContext

CASE = "wp-00000000d0f7"


def _table(key: str, columns: tuple[str, ...], rows: list[list[str]]) -> contract.Table:
    heads = rated_offer.settings()["heads"][key]
    spec = contract.spec(rated_offer.OUTPUT)["tables"][key]["columns"]
    return contract.table(key, key, columns, tuple(heads[c] for c in columns), rows, tuple(spec[c] for c in columns))


def _offer(state: str | None = None, minimum: str | None = None, factors: bool = True) -> contract.OutputData:
    """An offer of two modules (472,825 HUF); with `factors`, a payment frequency factor of 1.03 and the total rounded."""
    labels = rated_offer.settings()["labels"]
    values = {"package": "Minta Kft.", "version": None if state else 2, "state": state, "offer_date": "2026-10-26",
              "approved_by": "Minta Bela", "policyholder": "Minta Kft.", "tax_number": "12345676-2-41",
              "contact_name": "Minta Anna", "contact_email": "anna@minta.hu",
              "product": "Kitalált vállalkozói vagyonbiztosítás", "cover_start": "2026-11-01", "verdict": "elfogadható",
              "currency": "HUF", "subtotal": "472825", "total": "487000" if factors else "472825",
              "rounded": "kerekítve" if factors else None, "policy_minimum": minimum, "conditions": "Általános feltételek"}
    modules = _table("modules", rated_offer.COLUMNS, [["Vagyonbiztosítás", "400000"], ["Túlfeszültség", "72825"]])
    adjusted = _table("adjustments", rated_offer.ADJUSTMENT_COLUMNS, [["Díjfizetési gyakoriság", "1.03"]] if factors else [])
    return contract.OutputData(rated_offer.OUTPUT, tuple(contract.Value(k, labels[k], v) for k, v in values.items()),
                               (modules, adjusted))


def _text(file: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(file))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([*(p.text for p in doc.paragraphs), *cells, *(p.text for p in doc.sections[0].footer.paragraphs)])


def _ctx() -> StepContext:
    premium = StepDef(id="premium", kind="cap.decision.v1")
    document = StepDef(id="document", kind="cap.priced_document.v1", gates=("premium",),
                       params={"output": rated_offer.OUTPUT, "from_step": "premium"})
    definition = ProcessDef(id="invented", version="1", title="Invented", steps=(premium, document))
    ctx = StepContext(case_id=CASE, step=document, definition=definition, data={"container": CASE, "name": "Minta Kft."},
                      actor=identity.machine("test"))
    return ctx


@pytest.fixture()
def decided(tmp_path, monkeypatch):
    """A premium decision's approved version in a store of its own, and a stand-in builder for the offer."""
    registrations.load_all()  # the real builders first, so the stand-in replaces them for the test
    seen: list[dict] = []

    def build(content, *, case_id, package, version, approved_by, approved_at, output):
        seen.append({"case_id": case_id, "version": version, "approved_by": approved_by})
        return _offer()

    monkeypatch.setattr(parts, "_VERSIONS", {rated_offer.OUTPUT: parts.VersionDocument(
        rated_offer.OUTPUT, "tests.synthetic", build, reads_case=True, name=lambda content, package, n: f"ajanlat-v{n}.docx")})
    with store.use_store(tmp_path / "offer.sqlite"):
        with store.connect() as c:
            store.begin_immediate(c)
            n = package_release.record_own(c, CASE, content={"choice": "accept"}, fingerprint="f", locks=[],
                                           actor=identity.person("Minta Bela"), step="premium")
        ctx = _ctx()
        object.__setattr__(ctx, "result_of", lambda step_id: (1, {"version": n, "choice": "accept", "about_version": 1}))
        yield ctx, n, seen


def test_the_priced_document_stores_the_offer_with_the_version_once(decided) -> None:
    ctx, n, seen = decided
    first = kinds.of("cap.priced_document.v1").run(ctx)
    stored = package_release.own_file(CASE, n, package_release.APPROVED)
    assert stored is not None and stored["name"] == f"ajanlat-v{n}.docx" and stored["template"] == "builtin"
    assert first == kinds.Done({"output": rated_offer.OUTPUT, "version": n, "values": len(_offer().values),
                                "file": {"name": stored["name"], "sha256": hashlib.sha256(stored["content"]).hexdigest()}})
    assert seen == [{"case_id": CASE, "version": n, "approved_by": "Minta Bela"}]  # the builder reads the case
    assert kinds.of("cap.priced_document.v1").run(ctx) == first  # a second run stores nothing new
    assert len(package_release.own_files(CASE, n)) == 1
    text = _text(stored["content"])
    assert "Szerződő: Minta Kft." in text and "Díj (HUF)" in text
    assert "Túlfeszültség\n72\u00a0825\n" in text  # each module's premium in its row
    assert "Részösszeg: 472\u00a0825 HUF" in text and "Díjfizetési gyakoriság\n1,03" in text  # the factor on the total
    assert "Éves díj: 487\u00a0000 HUF (kerekítve)" in text


def test_the_built_in_template_leaves_out_an_empty_line_and_marks_a_draft() -> None:
    final = _text(builtin.render(_offer(minimum="a szerződés minimális díja érvényes")))
    assert "Biztosítási ajánlat" in final and "Tervezet" not in final and "(a szerződés minimális díja érvényes)" in final
    assert "2. változat" in final and "Kockázatviselés kezdete: 2026. 11. 01." in final
    draft = _text(builtin.render(_offer(state="Tervezet")))
    assert "Biztosítási ajánlat – Tervezet" in draft and "minimális díja" not in draft and "változat" not in draft
    plain = _text(builtin.render(_offer(factors=False)))  # no factor on the total: the modules add up to it
    assert "Részösszeg" not in plain and "Díjmódosító tényező" not in plain and "Éves díj: 472\u00a0825 HUF\n" in plain


def test_a_decision_that_names_no_kept_premium_is_no_offer(tmp_path) -> None:
    with store.use_store(tmp_path / "empty.sqlite"):
        for content in ({"about": "rate", "result": {"version": 1}}, {"about": "rate"}, {}):
            with pytest.raises(rated_offer.RatedOfferError, match="names no premium the case kept"):
                rated_offer.data(content, case_id=CASE, package="Minta Kft.")


def test_the_offer_s_contract_and_labels_agree() -> None:
    contract.check(_offer())
    assert set(rated_offer.sample_values()) == set(contract.spec(rated_offer.OUTPUT)["values"])
    assert rated_offer.sample_heads() == {"modules": {"name": "Fedezet", "premium": "Díj"},
                                          "adjustments": {"name": "Díjmódosító tényező", "factor": "Szorzó"}}
    assert parts.version_document(rated_offer.OUTPUT).reads_case
