"""Native result tables and exports preserve values, source identity and machine claims."""
import io
import json

import pytest
from openpyxl import load_workbook

from jav import corrections, datasets, export, native_results, store, work, work_views
from jav.readers import native, pipeline
from jav.tablequery import Query
from tests.test_native_review_api_109 import published_run


def test_native_dataset_cache_version_and_exports(tmp_path):
    with store.use_store(tmp_path / "export.sqlite"):
        rid, item, publication, _cite = published_run(tmp_path)
        assert work_views.result_tables(rid) == ["native_facts"]
        before = datasets.query("native_facts", {"run_id": rid}, Query())
        fact = native_results.machine_facts(publication)[0]
        corrections.save_native(rid, item["item_id"], values={fact.fact_id: "=00123"}, native_sources={},
            expected_revision=0, expected_result_version=publication.result_version, actor="reviewer", confirm=[fact.fact_id])
        after = datasets.query("native_facts", {"run_id": rid}, Query())
        assert before["review_version"] != after["review_version"]
        assert after["rows"][0]["effective_value"] == "=00123"
        body = json.loads(export.json_bytes(rid))
        actual = body["native_items"][0]["native_facts"][0]
        assert actual["effective_value"] == "=00123" and actual["proposal"]["value"] == "00123"
        assert body["documents"] == []
        assert body["native_items"][0]["interpretation"]["correctness"] == "not_established"
        assert body["review_version"] == after["review_version"]
        book = load_workbook(io.BytesIO(export.xlsx_bytes(rid)), data_only=False)
        sheet = book["Native facts"]
        value_column = export.NATIVE_COLUMNS.index("effective_value") + 1
        assert sheet.cell(2, value_column).value == "=00123"
        assert sheet.cell(2, value_column).data_type == "s"
        content, _mime, _name, count = export.render(rid, "csv", "native_facts")
        assert count == 1 and "'=00123" in content.decode("utf-8-sig")


def test_null_and_leading_zero_remain_distinct_json_values(tmp_path):
    with store.use_store(tmp_path / "values.sqlite"):
        rid, item, publication, _cite = published_run(tmp_path)
        assert json.loads(export.json_bytes(rid))["native_items"][0]["native_facts"][0]["effective_value"] == "00123"
        fact = native_results.machine_facts(publication)[0]
        corrections.save_native(rid, item["item_id"], values={fact.fact_id: None}, native_sources={},
            expected_revision=0, expected_result_version=publication.result_version, actor="reviewer")
        assert json.loads(export.json_bytes(rid))["native_items"][0]["native_facts"][0]["effective_value"] is None


@pytest.mark.parametrize("entrypoint", ["json_bytes", "render"])
@pytest.mark.parametrize("edit_after", ["document_rows", "native_rows"])
def test_json_export_interleaved_correction_cannot_approve_unseen_values(tmp_path, monkeypatch, entrypoint, edit_after):
    monkeypatch.setattr(pipeline, "run", native.read)
    with store.use_store(tmp_path / "interleaved.sqlite"):
        rid, item, publication, _cite = published_run(tmp_path)
        fact = native_results.machine_facts(publication)[0]
        reviewed_version = corrections.review_version(rid)
        owner, loader = (datasets, "run_records") if edit_after == "document_rows" else (export, "native_records")
        original = getattr(owner, loader)
        edits = []

        def collect_then_edit(run_id):
            records = original(run_id)
            if not edits:
                edits.append(True)
                corrections.save_native(rid, item["item_id"], values={fact.fact_id: "00456"}, native_sources={},
                    expected_revision=0, expected_result_version=publication.result_version, actor="second-reviewer")
            return records

        monkeypatch.setattr(owner, loader, collect_then_edit)
        content = export.json_bytes(rid) if entrypoint == "json_bytes" else export.render(rid, "json")[0]
        body = json.loads(content)
        current = corrections.item_result(rid, item["item_id"])
        assert edits == [True]
        assert body["review_version"] == reviewed_version != current["review_version"]
        expected = "00123" if edit_after == "native_rows" else "00456"
        assert body["native_items"][0]["native_facts"][0]["effective_value"] == expected
        assert current["native_facts"][0]["effective_value"] == "00456"
        with pytest.raises(work.RevisionConflict):
            work.approve_run(rid, actor="reader-of-export", review_version=body["review_version"])
        assert work.get_run(rid)["approval"] is None
        with store.connect() as connection:
            assert connection.execute("SELECT COUNT(*) FROM invocations").fetchone()[0] == 0


def test_native_json_export_does_not_bind_supplied_unversioned_rows_to_current_version(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "run", native.read)
    with store.use_store(tmp_path / "supplied-rows.sqlite"):
        rid, _item, _publication, _cite = published_run(tmp_path)
        body = json.loads(export.json_bytes(rid, records=[{"stale": "unversioned caller data"}]))
        assert body["documents"] == []
        assert body["review_version"] == corrections.review_version(rid)
